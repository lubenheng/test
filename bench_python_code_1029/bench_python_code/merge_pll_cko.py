"""将 Save_Results/OCD5010/PLL_CKO_vs_Vin_1~5.xlsx 合并到一个 Excel 文件中。
每个原文件的数据表与折线图分别以 PLL_CKO_数据_N / PLL_CKO_折线图_N 命名放入同一工作簿。
"""
import os
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

SAVE_DIR = os.path.join(os.path.dirname(__file__), "Save_Results", "OCD5010")
SRC_FILES = [os.path.join(SAVE_DIR, f"PLL_CKO_vs_Vin_{i}.xlsx") for i in range(1, 6)]
OUT_FILE = os.path.join(SAVE_DIR, "PLL_CKO_vs_Vin_merged.xlsx")

RW_PLL_VCO_COLORS = ["1F77B4", "FF7F0E", "2CA02C", "D62728",
                      "9467BD", "8C564B", "E377C2", "7F7F7F"]

header_font = Font(bold=True)
header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
header = ["RW_PLL_VCO", "序号", "设置电压(mV)", "实测电压(mV)", "频率(Hz)", "周期(us)", "占空比"]
col_widths = [12, 8, 16, 16, 16, 14, 12]

wb_out = openpyxl.Workbook()
wb_out.remove(wb_out.active)

for i, src_path in enumerate(SRC_FILES, start=1):
    if not os.path.exists(src_path):
        print(f"[WARN] file not found, skip: {src_path}")
        continue

    wb_src = openpyxl.load_workbook(src_path, data_only=True)
    ws_src = wb_src["PLL_CKO_数据"]

    ws_data = wb_out.create_sheet(f"PLL_CKO_数据_{i}")
    ws_data.append(header)
    for col in range(1, len(header) + 1):
        cell = ws_data.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    vco_row_ranges = {}
    current_vco = None
    first_row = None
    for row in ws_src.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        ws_data.append(list(row))
        vco = row[0]
        if vco != current_vco:
            if current_vco is not None:
                vco_row_ranges[current_vco] = (first_row, ws_data.max_row - 1)
            current_vco = vco
            first_row = ws_data.max_row
    if current_vco is not None:
        vco_row_ranges[current_vco] = (first_row, ws_data.max_row)

    for c, w in enumerate(col_widths, start=1):
        ws_data.column_dimensions[get_column_letter(c)].width = w
    ws_data.freeze_panes = "A2"

    ws_chart = wb_out.create_sheet(f"PLL_CKO_折线图_{i}")
    chart = ScatterChart()
    chart.title = f"PLL_CKO 频率 vs 输入电压 (按 RW_PLL_VCO 分组) - 第{i}次测试"
    chart.x_axis.title = "实测电压(mV)"
    chart.y_axis.title = "频率(Hz)"
    chart.x_axis.axPos = "b"
    chart.y_axis.axPos = "l"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.style = 13
    chart.width, chart.height = 42, 26

    vco_order = [f"{v:03b}" for v in range(0b111, -1, -1)]
    for idx, vco_code in enumerate(vco_order):
        if vco_code not in vco_row_ranges:
            continue
        first_row, last_row = vco_row_ranges[vco_code]
        vol_ref = Reference(ws_data, min_col=4, min_row=first_row, max_row=last_row)
        freq_ref = Reference(ws_data, min_col=5, min_row=first_row, max_row=last_row)

        series = Series(freq_ref, vol_ref, title=f"RW_PLL_VCO={vco_code}")
        series.marker = Marker(symbol="circle", size=6)
        color = RW_PLL_VCO_COLORS[idx % len(RW_PLL_VCO_COLORS)]
        series.graphicalProperties.line = LineProperties(w=20000, solidFill=color)
        series.marker.graphicalProperties.solidFill = color
        series.smooth = False
        chart.series.append(series)

    ws_chart.add_chart(chart, "B2")
    print(f"[OK] merged: {src_path} -> PLL_CKO_数据_{i} / PLL_CKO_折线图_{i}")

wb_out.save(OUT_FILE)
print(f"\n[OK] merge complete, saved to: {OUT_FILE}")
