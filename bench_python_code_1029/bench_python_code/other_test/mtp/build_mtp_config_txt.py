# -*- coding: utf-8 -*-
"""
把 MTP 配置寄存器表导出为纯文本（.txt），便于在编辑器里直接查看、搜索和 diff。

数据来源：chips/OCD2361/mtp_config_table.py（该文件可由 xlsx 生成或手工修改）
输出    ：Save_Results/MTP_Table/mtp_config_table.txt

运行：
    python other_test/mtp/build_mtp_config_txt.py
    python other_test/mtp/build_mtp_config_txt.py --from-xlsx   # 改从原始 xlsx 读取
"""

import os
import sys
from pathlib import Path

if os.name == 'nt':
    os.system('chcp 65001 > nul')
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[2]
DST = ROOT / "Save_Results" / "MTP_Table" / "mtp_config_table.txt"

sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "instrument" / "IIC"))

CRC_ADDR_LIMIT = 416
CRC_RSV = (507, 508, 509, 510)


def load_rows(from_xlsx=False):
    """返回 (rows, others)。

    rows:   [(addr, index, b31, b23, b15, b7, module)]
    others: [(src_row, operation, [7 个单元格])]
    """
    if not from_xlsx:
        from chips.OCD2361 import mtp_config_table as cfg
        rows = [(a, i, h, k, l, m, mod) for a, i, h, k, l, m, mod in cfg.MTP_CONFIG_ROWS]
        others = [(r, op, list(vals)) for r, op, *vals in getattr(cfg, "OTHER_ROWS", [])]
        return rows, others

    import openpyxl
    from iic_ops import IICLogic
    src = ROOT / "MTP_配置寄存器（测试版本2）.xlsx"
    ws = openpyxl.load_workbook(src, data_only=True)["MTP批量模板"]
    rows, others = [], []
    for r in range(3, ws.max_row + 1):
        cells = [ws.cell(r, c).value for c in range(1, 11)]
        if all(v is None or str(v).strip() == "" for v in cells[:3]):
            continue
        op = str(cells[0]).strip().upper() if cells[0] else ""
        if op != "PROGRAM":
            others.append((r, op, cells[:7]))
            continue
        rows.append((int(cells[1]), int(cells[2]),
                     *(IICLogic.parse_hex_byte(cells[c], f"第{r}行") for c in range(3, 7)),
                     cells[9]))
    return rows, others


def build_text(rows, others, origin):
    from crc import calculate_crc32

    byte_map = {}
    for addr, index, b31, b23, b15, b7, _ in rows:
        phys = addr + index * 4
        byte_map[phys] = b7
        byte_map[phys + 1] = b15
        byte_map[phys + 2] = b23
        byte_map[phys + 3] = b31

    crc_addrs = sorted(a for a in byte_map if a < CRC_ADDR_LIMIT)
    crc = calculate_crc32(bytes(byte_map[a] for a in crc_addrs))
    phys_all = sorted(byte_map)

    out = []
    w = out.append
    w("=" * 96)
    w("MTP 配置寄存器表（测试版本2）")
    w("=" * 96)
    w(f"数据来源    : {origin}")
    w(f"数据格式    : 每行 (ADDR, INDEX, [31:24], [23:16], [15:8], [7:0], 模块号)")
    w(f"地址映射    : phys = ADDR + INDEX*4")
    w(f"PROGRAM 行  : {len(rows)}")
    w(f"烧录字节数  : {len(phys_all)}（地址 {phys_all[0]} ~ {phys_all[-1]}）")
    w(f"CRC 覆盖    : {len(crc_addrs)} 字节（地址 {crc_addrs[0]} ~ {crc_addrs[-1]}，phys < {CRC_ADDR_LIMIT}）")
    w(f"CRC32       : 0x{crc:08X}（crc.py，LSB 优先）")
    w(f"CRC 存放    : 地址 {' / '.join(map(str, CRC_RSV))}，小端 "
      f"{' '.join(f'{(crc >> (8*i)) & 0xFF:02X}' for i in range(4))}")
    w("=" * 96)
    w("")

    header = (f"{'#':>3}  {'ADDR':>5}  {'INDEX':>5}  "
              f"{'[31:24]':>7}  {'[23:16]':>7}  {'[15:8]':>7}  {'[7:0]':>6}  "
              f"{'PHYS':>5}  {'CRC?':>4}  {'模块':>4}")
    w(header)
    w("-" * 96)

    prev_addr = None
    for n, (addr, index, b31, b23, b15, b7, mod) in enumerate(rows, start=1):
        if prev_addr is not None and addr != prev_addr:
            w("")
        prev_addr = addr
        phys = addr + index * 4
        incrc = "Y" if phys < CRC_ADDR_LIMIT else "-"
        mod_txt = "" if mod is None else str(mod)
        w(f"{n:>3}  0x{addr:02X}  {index:>5}  "
          f"0x{b31:02X}   0x{b23:02X}   0x{b15:02X}   0x{b7:02X}   "
          f"{phys:>5}  {incrc:>4}  {mod_txt:>4}")

    w("")
    w("-" * 96)
    w("图例：PHYS = 物理字节地址；CRC? = 是否计入 CRC（Y=计入，-=只烧录不进 CRC）")
    w("")

    if others:
        w("=" * 96)
        w("源表中未烧录的行（非 PROGRAM，保留供参考）")
        w("=" * 96)
        w(f"{'源行':>4}  {'OPERATION':<16}  {'[31:24]':>8}  {'[23:16]':>8}  "
          f"{'[15:8]':>8}  {'[7:0]':>8}")
        for r, op, cells in others:
            vals = cells[3:7] if len(cells) >= 7 else cells[-4:]
            w(f"{r:>4}  {op:<16}  " + "  ".join(f"{str(v) if v is not None else '':>8}" for v in vals))
        w("")
        w("说明：这些行不属于 MTP 烧录数据，未纳入 MTP_CONFIG_ROWS。")
        w("")

    w("=" * 96)
    w("修改指引")
    w("=" * 96)
    w("1. 直接编辑 chips/OCD2361/mtp_config_table.py 中的 MTP_CONFIG_ROWS")
    w("2. python chips/OCD2361/mtp_config_table.py          # 打印并自检 CRC")
    w("3. python other_test/mtp/build_mtp_config_txt.py     # 重新生成本文本")
    w("4. python other_test/mtp/build_mtp_template.py --from-module   # 生成 Excel 烧录模板")
    w("   或 main.py 菜单 [e] 选择本 txt / --from-txt 本文件，生成 <txt名>_烧录模板.xlsx")
    w("5. main.py 菜单 [32] → 选择 _烧录模板.xlsx 或本文本 → 全片擦除+烧录+校验+自检")
    w("   本文件输出在 Save_Results/MTP_Table/ 下，菜单 [32] 的弹窗默认打开该目录")
    return "\n".join(out) + "\n", crc


def main():
    from_xlsx = "--from-xlsx" in sys.argv
    rows, others = load_rows(from_xlsx)
    origin = "MTP_配置寄存器（测试版本2）.xlsx" if from_xlsx else "chips/OCD2361/mtp_config_table.py"
    text, crc = build_text(rows, others, origin)
    DST.parent.mkdir(parents=True, exist_ok=True)
    DST.write_text(text, encoding='utf-8')
    print(f"PROGRAM 行 {len(rows)}，CRC32 = 0x{crc:08X}")
    print(f"✅ 已生成：{DST}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
