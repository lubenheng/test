import sys
import os

# 添加项目根目录到 Python 路径（相对导入方式，支持路径转移）
current_dir = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.dirname(os.path.dirname(current_dir))
if project_root not in sys.path:
    sys.path.insert(0, project_root)

# 串口通信模块
from instrument.platform import serial_comm as _serial
# 万用表模块（RIGOL DM3068，函数名与 Keysight 兼容）
from instrument.multimeter import RIGOL_DM3068 as _dmm

INSTR_AVAILABLE = True
# 兼容旧变量名
KEYSIGHT_AVAILABLE = True
SERIAL_AVAILABLE = True

# 示例使用函数
def test_imports():
    """测试导入是否成功"""
    if KEYSIGHT_AVAILABLE:
        print("RIGOL DM3068 模块可用")
        # 这里可以添加测试代码
    else:
        print("RIGOL DM3068 模块不可用")

    if SERIAL_AVAILABLE:
        print("串口通信模块可用")
        # 这里可以添加测试代码
    else:
        print("串口通信模块不可用")

# DAC测试示例函数
def dac_output_test_example():
    """
    DAC输出测试
    使用串口控制DAC和万用表测量
    """
    print("=== DAC输出测试示例 ===")

    # 1. 初始化串口（假设DAC通过串口控制）
    if SERIAL_AVAILABLE:
        print("1. 初始化串口...")
        # 根据实际串口号修改，例如 "COM3"
        if _serial.serial_init(port="COM5", baudrate=115200, timeout=2):
            print("   串口初始化成功")
        else:
            print("   串口初始化失败，请检查连接")
            return
    else:
        print("串口模块不可用，跳过串口测试")
        return

    # 2. 初始化万用表
    if KEYSIGHT_AVAILABLE:
        print("2. 初始化RIGOL DM3068万用表...")
        if _dmm.init_dmm():
            print("   万用表初始化成功")
        else:
            print("   万用表初始化失败，请检查连接")
            _serial.serial_close()
            return
    else:
        print("万用表模块不可用，跳过电压测量")
        # 继续串口测试但不测量

    try:
        # 3. 通过串口设置DAC输出电压（示例命令）
        print("3. 设置DAC输出电压...")
        # 假设DAC命令格式：SET_VOLTAGE <电压值>
        dac_voltage = 2.5  # 设置2.5V
        command = f"SET_VOLTAGE {dac_voltage}\r\n"
        bytes_sent = _serial.serial_send(command)
        print(f"   发送命令: {command.strip()}，发送字节数: {bytes_sent}")

        # 等待DAC稳定
        import time
        time.sleep(0.2)

        # 4. 读取DAC响应（如果有）
        response = _serial.serial_receive(1024)
        if response:
            print(f"   DAC响应: {response.decode('utf-8', errors='ignore').strip()}")

        # 5. 使用万用表测量实际输出电压
        if KEYSIGHT_AVAILABLE:
            print("4. 使用万用表测量输出电压...")
            # 读取单次DC电压，10V量程
            measured_voltage_mv = _dmm.read_single_voltage(
                voltage_type="DC",
                range_val="10V",
                resolution=0.00001
            )

            if measured_voltage_mv is not None:
                measured_voltage_v = measured_voltage_mv / 1000.0  # 转换回V
                print(f"   测量电压: {measured_voltage_v:.3f} V")
                print(f"   设置电压: {dac_voltage:.3f} V")
                print(f"   误差: {abs(measured_voltage_v - dac_voltage):.3f} V")

                # 简单判断是否在允许误差范围内
                allowed_error = 0.01  # 10mV误差
                if abs(measured_voltage_v - dac_voltage) <= allowed_error:
                    print("   ✓ DAC输出测试通过")
                else:
                    print("   ✗ DAC输出测试失败，误差过大")
            else:
                print("   电压测量失败")

        # 6. 测试多个电压点
        if KEYSIGHT_AVAILABLE:
            print("5. 测试多个电压点...")
            test_voltages = [1.0, 2.0, 3.0, 4.0, 5.0]
            results = []

            for voltage in test_voltages:
                # 设置DAC电压
                command = f"SET_VOLTAGE {voltage}\r\n"
                _serial.serial_send(command)
                time.sleep(0.1)  # 等待稳定

                # 测量电压
                measured_mv = _dmm.read_single_voltage(voltage_type="DC", range_val="10V")
                if measured_mv is not None:
                    measured_v = measured_mv / 1000.0
                    error = abs(measured_v - voltage)
                    results.append((voltage, measured_v, error))
                    print(f"   设置: {voltage:.3f}V, 测量: {measured_v:.3f}V, 误差: {error:.3f}V")
                else:
                    print(f"   设置: {voltage:.3f}V, 测量失败")

            # 输出测试总结
            if results:
                print("\n测试总结:")
                for set_v, meas_v, err in results:
                    print(f"   {set_v:.3f}V -> {meas_v:.3f}V (误差: {err:.3f}V)")

    finally:
        # 7. 清理资源
        print("6. 清理资源...")
        if KEYSIGHT_AVAILABLE:
            _dmm.release_dmm()
            print("   万用表资源已释放")

        if SERIAL_AVAILABLE:
            _serial.serial_close()
            print("   串口已关闭")

        print("=== 测试完成 ===")

def generate_values(start=0, end=65535, num_points=32):
    """
    生成从start到end的num_points个等间隔值（包含两端）

    Args:
        start: 起始值
        end: 结束值
        num_points: 点数

    Returns:
        list: 生成的值列表
    """
    if num_points <= 1:
        return [start]

    step = (end - start) / (num_points - 1)
    values = [start + i * step for i in range(num_points)]
    # 四舍五入为整数
    values = [round(v) for v in values]
    # 确保第一个和最后一个值精确
    values[0] = start
    values[-1] = end
    return values

# ──────────────────────────────────────────────
# Excel 格式化辅助函数
# ──────────────────────────────────────────────
def _apply_excel_styling(ws, data_start_row=2, header_row=None):
    """
    对 Excel 工作表应用统一样式美化。

    包括：表头蓝底白字、交替行色、细边框、数字列格式、自动列宽、冻结窗格、自动筛选。

    Args:
        ws: openpyxl Worksheet 对象
        data_start_row: 数据起始行号（默认2，即第一行为表头，数据从第2行开始）
        header_row: 表头行号（默认 None = data_start_row - 1）
    """
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if header_row is None:
        header_row = data_start_row - 1
    if header_row < 1:
        header_row = 1
    if data_start_row <= header_row:
        data_start_row = header_row + 1

    # ===== 样式定义 =====
    header_font = Font(name='Microsoft YaHei', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF'),
    )

    even_fill = PatternFill(start_color='EBF1FA', end_color='EBF1FA', fill_type='solid')
    odd_fill  = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    # 成功/失败 条件颜色
    green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    red_fill   = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    red_font   = Font(color='9C0006')
    green_font = Font(color='006100')

    # 识别数字列（列名含 mV、百分比、误差、电压 等关键字）
    num_keywords = ['mv', '电压', '误差', '百分比', 'error', 'percent']
    num_format_3 = '#,##0.000'  # 3位小数

    # ===== 1. 表头样式（使用 header_row 定位） =====
    if ws.max_row >= header_row:
        for cell in ws[header_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

    # ===== 2. 数据区域样式 =====
    header_texts = [str(c.value or '') for c in ws[header_row]]
    total_cols = len(header_texts)

    if ws.max_row >= data_start_row:
        for row in ws.iter_rows(min_row=data_start_row, max_row=ws.max_row,
                                min_col=1, max_col=total_cols):
            for col_idx, cell in enumerate(row):
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')
                # 交替行色
                cell.fill = even_fill if cell.row % 2 == 0 else odd_fill

                # 数字列格式
                if col_idx < len(header_texts):
                    header_lower = header_texts[col_idx].lower()
                    if any(kw in header_lower for kw in num_keywords):
                        if '%' in header_texts[col_idx] or '百分比' in header_texts[col_idx]:
                            cell.number_format = '0.00"%"'
                        else:
                            cell.number_format = num_format_3

                # 状态列条件着色
                if col_idx < len(header_texts) and cell.value and '状态' in header_texts[col_idx]:
                    val_str = str(cell.value)
                    if '成功' in val_str:
                        cell.fill = green_fill
                        cell.font = green_font
                    elif '失败' in val_str or '读取失败' in val_str:
                        cell.fill = red_fill
                        cell.font = red_font

    # ===== 3. 自动列宽（考虑中文字符双倍宽） =====
    for col_cells in ws.columns:
        col_letter = col_cells[0].column_letter
        max_len = 0
        for cell in col_cells:
            try:
                val_str = str(cell.value) if cell.value is not None else ''
                str_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
                max_len = max(max_len, str_len)
            except Exception:
                pass
        # 至少保证表头完整显示
        first_cell = str(col_cells[0].value)
        header_len = sum(2 if ord(c) > 127 else 1 for c in first_cell)
        max_len = max(max_len, header_len)
        ws.column_dimensions[col_letter].width = min(max_len + 4, 30)

    # ===== 4. 冻结窗格（冻结到 header_row 之后一行，即 data_start_row） =====
    freeze_pane_row = data_start_row
    if ws.max_row >= freeze_pane_row:
        ws.freeze_panes = ws.cell(row=freeze_pane_row, column=1).coordinate

    # ===== 5. 自动筛选 =====
    if total_cols > 0 and ws.max_row >= data_start_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(total_cols)}{ws.max_row}"


def _store_numeric(value):
    """将数值转换为 float 或 None（避免 Excel 存为文本）"""
    if value is None:
        return None
    try:
        return round(float(value), 6)
    except (ValueError, TypeError):
        return None


def _save_csv_simple(csv_path, results, vref_value=None):
    """
    CSV 简易保存（当 pandas 不可用时的降级方案）

    Args:
        csv_path: 保存路径
        results: 结果字典列表
        vref_value: 基准电压值（可选）
    """
    import csv
    fieldnames = ['序号', '参数值', '基准电压(mV)', '理论电压(mV)',
                  'DAC电压(mV)', '运放电压(mV)', '误差(mV)', '误差百分比(%)', '状态']
    with open(csv_path, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for i, r in enumerate(results):
            writer.writerow({
                '序号': i + 1,
                '参数值': r['param'],
                '基准电压(mV)': _store_numeric(vref_value),
                '理论电压(mV)': _store_numeric(r.get('theoretical_mv')),
                'DAC电压(mV)': _store_numeric(r.get('dac_voltage_mv')),
                '运放电压(mV)': _store_numeric(r.get('actual_mv')),
                '误差(mV)': _store_numeric(r.get('error_mv')),
                '误差百分比(%)': _store_numeric(r.get('error_percent')),
                '状态': r['status']
            })
    print(f"   CSV 已保存: {csv_path}")


def spi_voltage_test():
    """
    SPI电压测试
    发送spi_cmd_set_val Ori {value}命令，读取万用表电压值
    测试32个从0到65535的等间隔值
    """
    import time
    import datetime
    import os

    # 基准电压值（将在测量后设置）
    vref_value = None
    vref_source = None  # 基准电压来源：万用表测量/手动输入

    print("=== SPI电压测试（增强版） ===")
    print("命令格式: spi_cmd_set_val Ori {value}")
    print("参数范围: 0-65535，测试32个等间隔点")
    print("")
    print("测试前准备:")
    print("1. 确保设备已正确连接")
    print("2. 确保串口连接正常")
    print("3. 确保万用表连接正常")
    print("")

    # 1. 提示烧录代码
    print("1. 烧录代码提示")
    print("   请确保设备已烧录正确的固件/代码")
    input("   按回车键继续...")

    # 2. 基准电压设置
    print("\n2. 基准电压设置")
    print("   请输入基准电压值（参考电压），默认 4095 mV")
    vref_input = input("   基准电压(mV) [默认 4095]: ").strip()
    if vref_input:
        try:
            vref_value = float(vref_input)
            if vref_value < 100:
                print(f"   警告: 输入值 {vref_value} mV 偏低（<100mV），请确认单位是否为mV")
            vref_source = "手动输入"
        except ValueError:
            print("   输入无效，使用默认值 4095 mV")
            vref_value = 4095.0
            vref_source = "默认值"
    else:
        vref_value = 4095.0
        vref_source = "默认值"
    print(f"   基准电压: {vref_value:.3f} mV ({vref_source})")

    # 默认终止符
    terminator = '\r\n'  # Windows默认

    # 4. 初始化串口
    if SERIAL_AVAILABLE:
        print("4. 初始化串口...")

        # 显示可用串口
        available_ports = _serial.SerialComm.list_serial_ports()
        print(f"   可用串口: {available_ports}")

        # 询问用户串口端口
        port_input = input("   请输入串口号 (如 COM3) 或按回车自动检测: ").strip()
        port = None

        if port_input:
            # 检查输入是否像COM端口（不区分大小写）
            if port_input.upper().startswith('COM'):
                port = port_input
                print(f"   使用指定端口: {port}")
            else:
                print(f"   输入 '{port_input}' 不是有效的COM端口格式，将使用自动检测")
        else:
            print("   使用自动检测")

        # 自动检测串口，波特率115200
        ser = _serial.serial_init(port=port, baudrate=115200, timeout=2)
        if ser:
            print(f"   串口初始化成功: {ser.port}")

            # 选择命令终止符
            print("   选择命令终止符...")
            print("   1. \\r\\n (Windows默认)")
            print("   2. \\n (Linux/Unix)")
            print("   3. \\r (Mac)")
            print("   4. 无终止符")
            term_choice = input("   请选择终止符类型 (1-4, 默认1): ").strip()

            if term_choice == '2':
                terminator = '\n'
                print("   使用终止符: \\n")
            elif term_choice == '3':
                terminator = '\r'
                print("   使用终止符: \\r")
            elif term_choice == '4':
                terminator = ''
                print("   使用终止符: 无")
            else:
                terminator = '\r\n'
                print("   使用终止符: \\r\\n")

            # 询问是否先测试串口通信
            test_com = input("   是否先测试串口通信? (y/n, 推荐y): ").strip().lower()
            if test_com == 'y':
                print("   串口通信测试...")
                test_command = f"TEST{terminator}"  # 使用选择的终止符
                print(f"   发送测试命令: {test_command.strip()}")

                # 显示原始字节（十六进制）用于调试
                cmd_bytes = test_command.encode('utf-8')
                hex_str = ' '.join([f'{b:02X}' for b in cmd_bytes])
                print(f"   原始字节(HEX): {hex_str}")

                sent = _serial.serial_send(test_command)
                print(f"   发送字节数: {sent}")

                # 尝试接收响应
                import time
                time.sleep(0.1)
                response = _serial.serial_receive(1024)
                if response:
                    response_hex = ' '.join([f'{b:02X}' for b in response])
                    response_str = response.decode('utf-8', errors='ignore').strip()
                    print(f"   收到响应: {response_str}")
                    print(f"   响应字节(HEX): {response_hex}")
                else:
                    print("   未收到响应")

                # 测试实际命令
                test_dac = input("   是否测试DAC命令? (y/n): ").strip().lower()
                if test_dac == 'y':
                    test_val = input("   输入测试参数值 (0-65535): ").strip()
                    try:
                        test_val = int(test_val)
                        if 0 <= test_val <= 65535:
                            test_cmd = f"spi_cmd_set_val Ori {test_val}{terminator}"
                            print(f"   发送DAC命令: {test_cmd.strip()}")

                            # 显示原始字节（十六进制）用于调试
                            cmd_bytes = test_cmd.encode('utf-8')
                            hex_str = ' '.join([f'{b:02X}' for b in cmd_bytes])
                            print(f"   原始字节(HEX): {hex_str}")

                            sent = _serial.serial_send(test_cmd)
                            print(f"   发送字节数: {sent}")
                            time.sleep(0.1)
                            response = _serial.serial_receive(1024)
                            if response:
                                response_hex = ' '.join([f'{b:02X}' for b in response])
                                response_str = response.decode('utf-8', errors='ignore').strip()
                                print(f"   收到响应: {response_str}")
                                print(f"   响应字节(HEX): {response_hex}")
                            else:
                                print("   无响应")
                        else:
                            print("   参数值超出范围")
                    except ValueError:
                        print("   无效的参数值")

                continue_test = input("   继续执行SPI电压测试? (y/n): ").strip().lower()
                if continue_test != 'y':
                    print("   测试中止")
                    _serial.serial_close()
                    return
        else:
            print("   串口初始化失败，请检查连接")
            return
    else:
        print("串口模块不可用，无法进行测试")
        return

    # 5. 初始化万用表（用于DAC测试）
    if KEYSIGHT_AVAILABLE:
        print("5. 初始化RIGOL DM3068万用表...")
        if _dmm.init_dmm():
            print("   万用表初始化成功")
        else:
            print("   万用表初始化失败，请检查连接")
            _serial.serial_close()
            return
    else:
        print("万用表模块不可用，无法进行测试")
        _serial.serial_close()
        return

    try:
        # 6. 生成32个测试值
        print("6. 生成测试参数...")
        values = generate_values(0, 65535, 32)
        print(f"   生成{len(values)}个测试值: {values}")

        # 7. 准备存储结果
        results = []

        # 8. 循环测试
        print("8. 开始循环测试...")
        for i, val in enumerate(values):
            print(f"   测试 {i+1}/{len(values)}: 参数值={val}")

            # 计算理论电压值
            if vref_value is not None:
                theoretical_voltage = (val / 65535.0) * vref_value  # vref_value单位为mV
                print(f"   理论电压: {theoretical_voltage:.3f} mV")
            else:
                theoretical_voltage = None
                print("   警告: 基准电压未设置，无法计算理论值")

            # 发送SPI命令
            cmd = f"spi_cmd_set_val Ori {val}{terminator}"  # 使用选择的终止符
            print(f"   发送命令: {cmd.strip()}")

            # 显示原始字节（十六进制）用于调试
            cmd_bytes = cmd.encode('utf-8')
            hex_str = ' '.join([f'{b:02X}' for b in cmd_bytes])
            print(f"   原始字节(HEX): {hex_str}")

            sent = _serial.serial_send(cmd)
            print(f"   发送字节数: {sent}")

            # 尝试接收响应
            time.sleep(0.05)  # 短暂等待响应
            response = _serial.serial_receive(1024)
            if response:
                response_hex = ' '.join([f'{b:02X}' for b in response])
                response_str = response.decode('utf-8', errors='ignore').strip()
                print(f"   收到响应: {response_str}")
                print(f"   响应字节(HEX): {response_hex}")

            if sent <= 0:
                print(f"   发送失败，跳过该点")
                results.append({
                    'param': val,
                    'theoretical_mv': theoretical_voltage,
                    'actual_mv': None,
                    'error_mv': None,
                    'error_percent': None,
                    'status': '发送失败'
                })
                continue

            # 延迟0.2秒
            time.sleep(0.2)

            # 测量DAC芯片输出电压
            print(f"   第{i+1}/{len(values)}次循环：请将万用表连接到DAC芯片输出测试点")
            input("   按回车键开始测量DAC输出电压...")
            dac_voltage = _dmm.read_single_voltage(voltage_type="DC", range_val="10V")
            if dac_voltage is None:
                print(f"   DAC电压读取失败")
                dac_voltage = None
            else:
                print(f"   DAC输出电压: {dac_voltage:.3f} mV")

            # 测量运放输出电压
            print(f"   第{i+1}/{len(values)}次循环：请将万用表连接到运放输出测试点")
            input("   按回车键开始测量运放输出电压...")
            amp_voltage = _dmm.read_single_voltage(voltage_type="DC", range_val="10V")
            if amp_voltage is None:
                print(f"   运放电压读取失败")
                amp_voltage = None
            else:
                print(f"   运放输出电压: {amp_voltage:.3f} mV")

            # 确定状态和误差计算
            if dac_voltage is None or amp_voltage is None:
                # 任一测量失败
                status = '读取失败'
                error_mv = None
                error_percent = None
            else:
                # 两个测量都成功
                status = '成功'
                # 计算误差（基于理论电压和运放输出电压？或者DAC输出电压？）
                # 根据原有逻辑，误差是基于理论电压和运放输出电压（actual_mv）
                if theoretical_voltage is not None and theoretical_voltage != 0:
                    error_mv = amp_voltage - theoretical_voltage
                    error_percent = (error_mv / theoretical_voltage) * 100
                    print(f"   误差: {error_mv:.3f} mV ({error_percent:.3f}%)")
                else:
                    error_mv = None
                    error_percent = None
                    status = '成功(无理论值)'

            # 存储结果
            results.append({
                'param': val,
                'theoretical_mv': theoretical_voltage,
                'dac_voltage_mv': dac_voltage,
                'actual_mv': amp_voltage,  # 运放输出电压，保持向后兼容
                'error_mv': error_mv,
                'error_percent': error_percent,
                'status': status
            })

        # 9. 分析结果
        successful_results = [r for r in results if r['actual_mv'] is not None]
        failed_results = [r for r in results if r['actual_mv'] is None]

        # 初始化统计变量（确保在所有代码路径中都有定义）
        actual_voltages = []
        errors = []
        avg_error = None
        max_error = None

        # 10. 打印详细结果表格
        print("\n" + "="*80)
        print("SPI电压测试结果")
        print("="*80)
        print(f"测试点数: {len(results)}")
        print(f"成功: {len(successful_results)} 点")
        print(f"失败: {len(failed_results)} 点")

        if vref_value is not None:
            print(f"基准电压: {vref_value:.3f} mV")
            if vref_source is not None:
                print(f"基准电压来源: {vref_source}")

        if successful_results:
            actual_voltages = [r['actual_mv'] for r in successful_results if r['actual_mv'] is not None]
            if actual_voltages:
                print(f"实际电压范围: {min(actual_voltages):.3f} ~ {max(actual_voltages):.3f} mV")
                print(f"实际平均电压: {sum(actual_voltages)/len(actual_voltages):.3f} mV")

            # 计算平均误差
            errors = [r['error_mv'] for r in successful_results if r['error_mv'] is not None]
            if errors:
                avg_error = sum(errors) / len(errors)
                max_error = max(errors, key=abs)
                print(f"平均误差: {avg_error:.3f} mV")
                print(f"最大误差: {max_error:.3f} mV")

        print("\n详细数据:")
        print("-"*80)
        print(f"{'序号':<4} {'参数值':<8} {'理论值(mV)':<15} {'DAC电压(mV)':<15} {'运放电压(mV)':<15} {'误差(mV)':<15} {'状态':<12}")
        print("-"*80)

        for i, r in enumerate(results):
            param = r['param']
            theoretical = r['theoretical_mv']
            dac_voltage = r['dac_voltage_mv']
            amp_voltage = r['actual_mv']  # 运放电压
            error = r['error_mv']
            status = r['status']

            theoretical_str = f"{theoretical:.3f}" if theoretical is not None else "N/A"
            dac_str = f"{dac_voltage:.3f}" if dac_voltage is not None else "N/A"
            amp_str = f"{amp_voltage:.3f}" if amp_voltage is not None else "N/A"
            error_str = f"{error:.3f}" if error is not None else "N/A"

            print(f"{i+1:<4} {param:<8} {theoretical_str:<15} {dac_str:<15} {amp_str:<15} {error_str:<15} {status:<12}")

        print("-"*80)

        # 11. 自动保存结果到Excel文件（dac_test文件夹）
        print("\n11. 保存结果到Excel文件...")
        current_dir = os.path.dirname(os.path.abspath(__file__))
        timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        excel_filename = f"spi_voltage_test_{timestamp}.xlsx"
        csv_filename = f"spi_voltage_test_{timestamp}.csv"
        excel_filepath = os.path.join(current_dir, excel_filename)
        csv_filepath = os.path.join(current_dir, csv_filename)

        try:
            import pandas as pd

            # ---- 详细数据表（存为数值，便于Excel计算） ----
            df_data = []
            for i, r in enumerate(results):
                df_data.append({
                    '序号': i + 1,
                    '参数值': r['param'],
                    '基准电压(mV)': _store_numeric(vref_value),
                    '理论电压(mV)': _store_numeric(r['theoretical_mv']),
                    'DAC电压(mV)': _store_numeric(r['dac_voltage_mv']),
                    '运放电压(mV)': _store_numeric(r['actual_mv']),
                    '误差(mV)': _store_numeric(r['error_mv']),
                    '误差百分比(%)': _store_numeric(r['error_percent']),
                    '状态': r['status']
                })

            df = pd.DataFrame(df_data)

            # ---- 汇总信息 ----
            summary_data = {
                '项目': [
                    '测试时间',
                    '命令格式',
                    '参数范围',
                    '测试点数',
                    '成功点数',
                    '失败点数',
                    '基准电压(mV)',
                    '基准电压来源',
                    '理论电压计算公式',
                    '实际电压范围(mV)',
                    '实际平均电压(mV)',
                    '平均误差(mV)',
                    '最大误差(mV)'
                ],
                '值': [
                    datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                    'spi_cmd_set_val Ori {value}',
                    '0-65535',
                    len(results),
                    len(successful_results),
                    len(failed_results),
                    _store_numeric(vref_value),
                    vref_source if vref_source is not None else "N/A",
                    '理论电压(mV) = (参数值 / 65535) × 基准电压(mV)',
                    (f"{min(actual_voltages):.3f} ~ {max(actual_voltages):.3f}"
                     if successful_results and actual_voltages else "N/A"),
                    (f"{sum(actual_voltages)/len(actual_voltages):.3f}"
                     if successful_results and actual_voltages else "N/A"),
                    _store_numeric(avg_error),
                    _store_numeric(max_error)
                ]
            }
            df_summary = pd.DataFrame(summary_data)

            # ---- 写出 Excel 并应用样式 ----
            with pd.ExcelWriter(excel_filepath, engine='openpyxl') as writer:
                # 写入数据
                df.to_excel(writer, sheet_name='详细数据', index=False)
                df_summary.to_excel(writer, sheet_name='测试汇总', index=False)

                # 获取 worksheet
                ws_detail = writer.sheets['详细数据']
                ws_summary = writer.sheets['测试汇总']

                # 应用统一样式
                _apply_excel_styling(ws_detail)
                _apply_excel_styling(ws_summary)

            print(f"   ✓ Excel 文件已保存: {excel_filepath}")

            # ---- CSV 备份（同样存数值，减少引号干扰） ----
            df.to_csv(csv_filepath, index=False, encoding='utf-8-sig', float_format='%.3f')
            print(f"   ✓ CSV 备份文件已保存: {csv_filepath}")

            # 文件大小信息
            print(f"   Excel: {os.path.getsize(excel_filepath):,} 字节")
            print(f"   CSV:   {os.path.getsize(csv_filepath):,} 字节")

        except ImportError as e:
            # pandas/openpyxl 缺失 → 纯 CSV 保存
            print(f"   警告: 无法导入 pandas/openpyxl，使用 CSV 格式保存")
            print(f"   错误: {e}")
            _save_csv_simple(csv_filepath, results, vref_value)

        except Exception as e:
            print(f"   保存 Excel 失败: {e}")
            _save_csv_simple(csv_filepath, results, vref_value)

    finally:
        # 12. 清理资源
        print("12. 清理资源...")
        if KEYSIGHT_AVAILABLE:
            _dmm.release_dmm()
            print("   万用表资源已释放")

        if SERIAL_AVAILABLE:
            _serial.serial_close()
            print("   串口已关闭")

        print("=== SPI电压测试完成 ===")

def dac_two_channel_595d_test():
    """
    两路DAC输出测试（使用GPIO 595D控制多路复用切换）

    测试流程：
    1. 发送 cmd_gpio_set_595d_test reset（初始复位）
    2. 第一路测试（通道6）：
       a. 发送 cmd_gpio_set_595d_test set 6 33
       b. 遍历测试参数值，发送 spi_cmd_set_val Ori {value}，测量电压
       c. 发送 cmd_gpio_set_595d_test reset（复位）
    3. 确认复位后开始第二路测试（通道7）：
       a. 发送 cmd_gpio_set_595d_test set 7 33
       b. 遍历测试参数值，发送 spi_cmd_set_val Ori {value}，测量电压
       c. 发送 cmd_gpio_set_595d_test reset（复位）
    4. 保存结果到Excel/CSV，清理资源
    """
    import time
    import datetime

    print("=== 两路DAC输出测试（GPIO 595D控制） ===")

    # ========== 1. 基准电压 ==========
    print("\n1. 请输入基准电压")
    vref_value = None
    vref_source = None
    vref_input = input("   基准电压值(mV) [默认 4095]: ").strip()
    if vref_input:
        try:
            vref_value = float(vref_input)
            vref_source = "手动输入"
        except ValueError:
            print("   输入无效，使用默认值 4095 mV")
    if vref_value is None:
        vref_value = 4095.0
        vref_source = "默认值"
    print(f"   基准电压: {vref_value:.3f} mV ({vref_source})")

    # ========== 2. 串口初始化（自动检测） ==========
    print("\n2. 初始化串口...")
    ports = _serial.SerialComm.list_serial_ports()
    if not ports:
        print("   未检测到可用串口")
        return
    print(f"   检测到串口: {ports}，正在连接...")
    ser_ok = False
    for p in ports:
        if _serial.serial_init(port=p, baudrate=115200, timeout=2):
            print(f"   串口初始化成功: {p}")
            ser_ok = True
            break
    if not ser_ok:
        print("   串口初始化失败")
        return

    # ========== 3. 万用表初始化 ==========
    print("\n3. 初始化万用表...")
    if not _dmm.init_dmm():
        print("   万用表初始化失败")
        _serial.serial_close()
        return
    print("   万用表初始化成功")

    terminator = '\r\n'

    # 声明变量供 finally 块使用
    test_values = []
    all_results = {}
    excel_path = None
    csv_path = None

    try:
        # ========== 4. 测试数据准备 ==========
        test_values = generate_values(0, 65535, 11)
        print(f"\n4. 测试参数准备: {len(test_values)}个测试点")
        print(f"   参数值: {test_values}")
        if vref_value:
            print(f"   基准电压: {vref_value:.3f} mV ({vref_source})")
        else:
            print("   警告: 无基准电压，将仅记录测量值")

        # ========== 5. 初始复位 ==========
        print("\n5. 发送初始复位指令...")
        _serial.serial_send(f"cmd_gpio_set_595d_test reset{terminator}")
        time.sleep(0.1)
        resp = _serial.serial_receive(1024)
        if resp:
            print(f"   响应: {resp.decode('utf-8', errors='ignore').strip()}")
        print("   ✓ 初始复位完成")

        # ========== 6. 两路通道测试 ==========
        channels = [
            {'name': 'CH1', 'channel': 6, 'value': 33, 'label': '第一路', 'gain': 1},
            {'name': 'CH2', 'channel': 7, 'value': 33, 'label': '第二路(2x增益)', 'gain': 2, 'max_voltage_mv': 5000}
        ]

        for ch_idx, ch in enumerate(channels):
            print(f"\n" + "="*60)
            print(f"{ch['label']}测试 (GPIO 595D 通道{ch['channel']}, set {ch['channel']} {ch['value']})")
            print("="*60)

            if ch_idx > 0:
                print(f"\n   切换到{ch['label']}...")
                time.sleep(0.5)  # 短暂等待稳定

            # 发送通道选择指令: cmd_gpio_set_595d_test set <ch> <val>
            set_cmd = f"cmd_gpio_set_595d_test set {ch['channel']} {ch['value']}{terminator}"
            print(f"\n   发送通道选择: {set_cmd.strip()}")
            _serial.serial_send(set_cmd)
            time.sleep(0.1)
            resp = _serial.serial_receive(1024)
            if resp:
                print(f"   响应: {resp.decode('utf-8', errors='ignore').strip()}")

            ch_results = []
            for i, val in enumerate(test_values):
                print(f"\n   测试点 {i+1}/{len(test_values)}: 参数值={val}")

                # 计算理论电压（乘以通道增益）
                gain = ch['gain']
                theoretical_mv = (val / 65535.0) * vref_value * gain if vref_value else None
                if theoretical_mv is not None:
                    # 考虑增益芯片电压上限（如第二路增益芯片最高输出约5V）
                    max_v = ch.get('max_voltage_mv')
                    if max_v is not None and theoretical_mv > max_v:
                        print(f"   理论电压: {theoretical_mv:.3f} mV (超增益芯片上限 {max_v} mV，已钳位)")
                        theoretical_mv = max_v
                    else:
                        gain_info = f" (增益{gain}x)" if gain != 1 else ""
                        print(f"   理论电压: {theoretical_mv:.3f} mV{gain_info}")

                # 发送SPI命令设置DAC值
                spi_cmd = f"spi_cmd_set_val Ori {val}{terminator}"
                print(f"   发送: {spi_cmd.strip()}")
                _serial.serial_send(spi_cmd)
                time.sleep(0.05)
                resp = _serial.serial_receive(1024)
                if resp:
                    resp_str = resp.decode('utf-8', errors='ignore').strip()
                    print(f"   DAC响应: {resp_str}")

                time.sleep(0.2)

                # 使用万用表测量实际输出电压
                measured = _dmm.read_single_voltage(voltage_type="DC", range_val="10V")
                if measured is not None:
                    print(f"   测量电压: {measured:.3f} mV")
                    if theoretical_mv is not None and theoretical_mv != 0:
                        error_mv = measured - theoretical_mv
                        error_pct = (error_mv / theoretical_mv) * 100
                        print(f"   误差: {error_mv:.3f} mV ({error_pct:.3f}%)")
                    else:
                        error_mv = 0.0 if theoretical_mv == 0 else None
                        error_pct = 0.0 if theoretical_mv == 0 else None
                        if theoretical_mv == 0:
                            print(f"   误差: 0.000 mV (0.000%)")

                    ch_results.append({
                        'param': val,
                        'theoretical_mv': theoretical_mv,
                        'measured_mv': measured,
                        'error_mv': error_mv,
                        'error_percent': error_pct,
                        'status': '成功'
                    })
                else:
                    print(f"   测量失败")
                    ch_results.append({
                        'param': val,
                        'theoretical_mv': theoretical_mv,
                        'measured_mv': None,
                        'error_mv': None,
                        'error_percent': None,
                        'status': '测量失败'
                    })

            all_results[ch['name']] = ch_results

            # 每路测试结束后发送reset指令
            print(f"\n   {ch['label']}测试完成，发送复位指令...")
            reset_cmd = f"cmd_gpio_set_595d_test reset{terminator}"
            print(f"   发送: {reset_cmd.strip()}")
            _serial.serial_send(reset_cmd)
            time.sleep(0.1)
            resp = _serial.serial_receive(1024)
            if resp:
                print(f"   响应: {resp.decode('utf-8', errors='ignore').strip()}")
            print(f"   ✓ {ch['label']}复位完成")

        # ========== 7. 结果汇总展示 ==========
        print("\n" + "="*80)
        print("两路DAC测试结果汇总")
        print("="*80)

        for ch_name, ch_results in all_results.items():
            # 查找该通道的增益信息
            ch_info = next((c for c in channels if c['name'] == ch_name), None)
            gain = ch_info['gain'] if ch_info else 1
            gain_label = f" (增益{gain}x)" if gain != 1 else ""
            successful = [r for r in ch_results if r['measured_mv'] is not None]
            failed = [r for r in ch_results if r['measured_mv'] is None]
            print(f"\n--- {ch_name}{gain_label} ---")
            print(f"   成功: {len(successful)} 点, 失败: {len(failed)} 点")

            if successful:
                voltages = [r['measured_mv'] for r in successful]
                print(f"   电压范围: {min(voltages):.3f} ~ {max(voltages):.3f} mV")
                errors = [r['error_mv'] for r in successful if r['error_mv'] is not None]
                if errors:
                    avg_err = sum(errors) / len(errors)
                    max_err = max(errors, key=abs)
                    print(f"   平均误差: {avg_err:.3f} mV")
                    print(f"   最大误差: {max_err:.3f} mV")

            # 打印详细表格
            print(f"   {'序号':<4} {'参数值':<8} {'理论值(mV)':<14} {'测量值(mV)':<14} {'误差(mV)':<14} {'状态':<10}")
            for i, r in enumerate(ch_results):
                t = f"{r['theoretical_mv']:.3f}" if r['theoretical_mv'] is not None else 'N/A'
                m = f"{r['measured_mv']:.3f}" if r['measured_mv'] is not None else 'N/A'
                e = f"{r['error_mv']:.3f}" if r['error_mv'] is not None else 'N/A'
                print(f"   {i+1:<4} {r['param']:<8} {t:<14} {m:<14} {e:<14} {r['status']:<10}")

        # ========== 8. 保存结果 ==========
        print("\n8. 保存测试结果...")
        current_dir = os.path.dirname(os.path.abspath(__file__))
        timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        excel_path = os.path.join(current_dir, f"dac_595d_two_channel_{timestamp}.xlsx")
        csv_path = os.path.join(current_dir, f"dac_595d_two_channel_{timestamp}.csv")

        try:
            import pandas as pd
            with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
                # ---- 各通道单独 sheet ----
                for ch_info in channels:
                    ch_name = ch_info['name']
                    ch_results = all_results[ch_name]
                    gain = ch_info['gain']
                    gain_label = f" (增益{gain}x)" if gain != 1 else ""
                    rows = []
                    for i, r in enumerate(ch_results):
                        rows.append({
                            '序号': i + 1,
                            '参数值': r['param'],
                            '理论电压(mV)': _store_numeric(r['theoretical_mv']),
                            '测量电压(mV)': _store_numeric(r['measured_mv']),
                            '误差(mV)': _store_numeric(r['error_mv']),
                            '误差百分比(%)': _store_numeric(r['error_percent']),
                            '状态': r['status']
                        })
                    df = pd.DataFrame(rows)
                    df.to_excel(writer, sheet_name=ch_name, index=False)

                    # 先美化（表头在行1），再插入顶部信息行
                    ws = writer.sheets[ch_name]
                    _apply_excel_styling(ws)
                    ws.insert_rows(1, 2)
                    ws['A1'] = f"{ch_info['label']}{gain_label}"
                    ws['A2'] = f"基准电压: {vref_value:.3f} mV ({vref_source})" if vref_value else "基准电压: N/A"
                    # 重新冻结：前3行固定（信息+表头）
                    ws.freeze_panes = 'A4'

                # ---- 合并对比 sheet ----
                combined_rows = []
                for ch_info in channels:
                    ch_name = ch_info['name']
                    ch_results = all_results[ch_name]
                    gain = ch_info['gain']
                    for i, r in enumerate(ch_results):
                        combined_rows.append({
                            '通道': f"{ch_name}({ch_info['label']})",
                            '增益': f"{gain}x",
                            '序号': i + 1,
                            '参数值': r['param'],
                            '理论电压(mV)': _store_numeric(r['theoretical_mv']),
                            '测量电压(mV)': _store_numeric(r['measured_mv']),
                            '误差(mV)': _store_numeric(r['error_mv']),
                            '误差百分比(%)': _store_numeric(r['error_percent']),
                            '状态': r['status']
                        })
                df_combined = pd.DataFrame(combined_rows)
                df_combined.to_excel(writer, sheet_name='两路对比汇总', index=False)
                _apply_excel_styling(writer.sheets['两路对比汇总'])

            print(f"   ✓ Excel 已保存: {excel_path}（各通道单独 sheet + 两路对比汇总）")

            # CSV 备份（用 pandas 的 to_csv，存数值）
            df_combined.to_csv(csv_path, index=False, encoding='utf-8-sig', float_format='%.3f')
            print(f"   ✓ CSV 备份已保存: {csv_path}")

        except ImportError:
            print("   警告: 无法导入 pandas/openpyxl，仅保存 CSV 格式")
            with open(csv_path, 'w', encoding='utf-8-sig', newline='') as f:
                import csv
                fnames = ['通道', '增益', '序号', '参数值', '理论电压(mV)',
                          '测量电压(mV)', '误差(mV)', '误差百分比(%)', '状态']
                writer = csv.DictWriter(f, fieldnames=fnames)
                writer.writeheader()
                for ch_info in channels:
                    ch_name = ch_info['name']
                    ch_results = all_results[ch_name]
                    gain = ch_info['gain']
                    for i, r in enumerate(ch_results):
                        writer.writerow({
                            '通道': ch_name,
                            '增益': f"{gain}x",
                            '序号': i + 1,
                            '参数值': r['param'],
                            '理论电压(mV)': _store_numeric(r.get('theoretical_mv')),
                            '测量电压(mV)': _store_numeric(r.get('measured_mv')),
                            '误差(mV)': _store_numeric(r.get('error_mv')),
                            '误差百分比(%)': _store_numeric(r.get('error_percent')),
                            '状态': r['status']
                        })
            print(f"   CSV 已保存: {csv_path}")
        except Exception as e:
            print(f"   保存失败: {e}")

    finally:
        # ========== 9. 清理资源 ==========
        print("\n9. 清理资源...")
        _dmm.release_dmm()
        print("   万用表资源已释放")
        _serial.serial_close()
        print("   串口已关闭")
        print("=== 两路DAC 595D测试完成 ===")


# 主函数
def main():
    """主函数：可以选择运行测试"""
    print("DAC输出测试模块")
    print("1. 测试导入")
    print("2. 运行DAC测试示例")
    print("3. 运行SPI电压测试")
    print("4. 运行两路DAC 595D测试")
    print("5. 退出")

    try:
        choice = input("请选择 (1-5): ").strip()
        if choice == "1":
            test_imports()
        elif choice == "2":
            dac_output_test_example()
        elif choice == "3":
            spi_voltage_test()
        elif choice == "4":
            dac_two_channel_595d_test()
        elif choice == "5":
            print("退出")
        else:
            print("无效选择")
    except KeyboardInterrupt:
        print("\n用户中断")
    except Exception as e:
        print(f"发生错误: {e}")

if __name__ == "__main__":
    # 直接运行测试示例，或调用main()进行交互
    # test_imports()  # 简单测试导入
    # dac_output_test_example()  # 运行完整示例
    main()  # 交互式选择
