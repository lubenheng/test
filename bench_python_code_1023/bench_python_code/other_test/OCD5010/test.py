import sys
import os
import time
import random
import openpyxl
from pathlib import Path

# 确保 project root 在 sys.path 中（独立运行时需要）
_project_root = str(Path(__file__).parents[2])
if _project_root not in sys.path:
    sys.path.insert(0, _project_root)

from instrument import instrument_set as instr
from instrument.platform.serial_init import serial_find_and_init
from instrument.platform.init_helper import check_init
from chips.OCD5010 import chip
from chips.OCD5010.chip import SERIAL_KEYWORD, SERIAL_BAUDRATE
from chips.OCD5010.gpio_config import init_pa5_pa6_altfunc14

def pll_atp_output():
    init_pa5_pa6_altfunc14()
    time.sleep(1)
    instr.swd_modify_register(0x4800003C, 0x0E, 24, 27)
    instr.swd_modify_register(0x40007C00, 0x01, 26, 26)
    instr.swd_modify_register(0x40007C1C, 0b10000, 8, 12)
    instr.swd_modify_register(0x40007C1C, 0x01, 13, 13)
    instr.swd_modify_register(0x40007C0C, 0x01, 17, 17)
    time.sleep(1)    


def otp_read_test(addr):
    """输入值为0-31, 对应OTP地址为0x8000-0x801F"""
    print("开始OTP回读测试")
    time.sleep(1)
    instr.swd_modify_register(0x4002204C, 0xA1B2C3D4, 0, 31)
    time.sleep(0.01)
    instr.swd_modify_register(0x4002204C, 0x5D6C7B8A, 0, 31)
    time.sleep(1)
    read_back = instr.swd_read_register(0x40022040)
    read_back = ((read_back >> 16) & 0x01)
    if read_back == 1:
        print("解锁失败, exit!!")
        return
    print("解锁成功，写入读取地址")
    time.sleep(1)
    addr = (addr & 0x1F)
    instr.swd_modify_register(0x40022044, addr, 0, 6)
    read_addr_data = instr.swd_read_register(0x40022048)
    print(f"读取出来的地址对应的值为：{read_addr_data:04X}")
    instr.swd_modify_register(0x4002204C, 0, 0, 31)

def reg_0x40012404_bit0_11_sample(duration_s=5):
    """
    每隔 1ms 读取一次寄存器 0x40012404 的值, 只取 bit0~bit11, 持续采样 duration_s 秒。
    采样结束后将 (时间, 寄存器值) 保存到 Excel, 并以时间为横坐标、寄存器值为纵坐标画折线图。

    参数:
        duration_s: float - 采样总时长(秒), 默认 5s
    """
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    REG_ADDR = 0x40012404
    INTERVAL_S = 0.001  # 1ms

    print(f"开始采样寄存器 0x{REG_ADDR:08X}[bit0:11], 持续 {duration_s}s, 采样间隔 {INTERVAL_S * 1000:.0f}ms...")

    samples = []  # [(elapsed_s, reg_val), ...]
    start_time = time.perf_counter()
    next_sample_time = start_time
    while True:
        now = time.perf_counter()
        elapsed = now - start_time
        if elapsed >= duration_s:
            break
        raw_val = instr.swd_read_register(REG_ADDR)
        val = raw_val & 0xFFF if raw_val is not None else None  # 只取 bit0~bit11
        samples.append((elapsed, val))

        next_sample_time += INTERVAL_S
        sleep_time = next_sample_time - time.perf_counter()
        if sleep_time > 0:
            time.sleep(sleep_time)

    print(f"采样完成, 共 {len(samples)} 个点")

    if not samples:
        print("⚠️ 无有效采样数据，跳过保存")
        return

    # ---------- 保存到 Excel: 数据表 + 折线图 ----------
    save_dir = os.path.join(_project_root, "Save_Results", "OCD5010")
    os.makedirs(save_dir, exist_ok=True)
    excel_path = os.path.join(save_dir, "reg_0x40012404_sample.xlsx")

    wb = openpyxl.Workbook()
    ws_data = wb.active
    ws_data.title = "采样数据"

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    header = ["时间(s)", "寄存器值(bit0-11)"]
    ws_data.append(header)
    for col in range(1, len(header) + 1):
        cell = ws_data.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for elapsed, val in samples:
        ws_data.append([round(elapsed, 4), val])

    col_widths = [14, 20]
    for i, w in enumerate(col_widths, start=1):
        ws_data.column_dimensions[get_column_letter(i)].width = w
    ws_data.freeze_panes = "A2"

    ws_chart = wb.create_sheet("折线图")
    chart = LineChart()
    chart.title = "寄存器 0x40012404[bit0:11] 随时间变化"
    chart.x_axis.title = "时间(s)"
    chart.y_axis.title = "寄存器值(bit0-11)"
    chart.style = 13
    chart.width, chart.height = 30, 15

    last_row = ws_data.max_row
    values_ref = Reference(ws_data, min_col=2, min_row=1, max_row=last_row)
    cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(values_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    for series in chart.series:
        series.smooth = False
        series.marker.symbol = "none"

    ws_chart.add_chart(chart, "B2")

    wb.save(excel_path)
    print(f"✅ 采样数据已保存至: {excel_path}")


def adc_3v_ibias_sample(duration_s=18 * 60):
    """
    3V量程下依次切换 8 个电流档, 每档按 reg_0x40012404_bit0_11_sample 的方式采样:
    每隔 1ms 读取一次 0x40012404 的 bit0~bit11, 持续 duration_s 秒。
    ADC 配置(含 3V 量程)由 C 程序完成, 本函数切换电流档、产生 START 信号并记录数据, 不画图。

    START 信号: 每 1ms 对 0x40007C18 bit0 (RW_ADC_START) 写 1, 半周期后写 0 (1KHz, 占空比50%),
    拉低后读取 DATA(0x40012404) bit0~11。写 START 用缓存的整字直接写入, 不做读-改-写,
    因此 C 程序在采样期间不能再写 0x40007C18, 否则会被这里的缓存值覆盖。

    电流档 (均在 0x40007C18):
        RW_ADC_IBIAS       bit11    : 0/1
        RW_ADC_VREFH_IBIAS bit16:17 : 00/01/10/11
        共 2 x 4 = 8 档

    每档采完立即单独保存一个 Excel (adc_3v_ibias_{时间戳}_IB{x}_VI{y}.xlsx), 中途出错不丢已完成档位。
    Excel 单个 sheet 最多 1048576 行(含表头), 超出时自动续写到 采样数据_2、_3 ... sheet。

    参数:
        duration_s: float - 每个电流档的采样时长(秒), 默认 18 分钟
    """
    REG_ADDR = 0x40012404
    IBIAS_REG = 0x40007C18
    INTERVAL_S = 0.001  # 1ms
    MAX_DATA_ROWS = 1048576 - 1  # Excel 单 sheet 行数上限扣除表头

    save_dir = os.path.join(_project_root, "Save_Results", "OCD5010")
    os.makedirs(save_dir, exist_ok=True)
    run_tag = time.strftime('%Y%m%d_%H%M%S')

    def _new_sheet(wb, title):
        ws = wb.create_sheet(title)
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 20
        ws.freeze_panes = "A2"
        ws.append(["时间(s)", "寄存器值(bit0-11)"])
        return ws

    SWITCH_DELAY_S = 3      # 电流档切换前后各延时
    SET_RETRY = 3           # 回读不匹配时重试次数

    def _set_bits(value, start_bit, end_bit, label):
        mask = (1 << (end_bit - start_bit + 1)) - 1
        for attempt in range(1, SET_RETRY + 1):
            instr.swd_modify_register(IBIAS_REG, value, start_bit, end_bit)
            time.sleep(0.5)
            reg_val = instr.swd_read_register(IBIAS_REG)
            if reg_val is None:
                print(f"  ⚠️ {label} 第{attempt}次回读失败")
                continue
            actual = (reg_val >> start_bit) & mask
            if actual == value:
                return True
            print(f"  ⚠️ {label} 第{attempt}次 写入={value}, 回读={actual}, 不匹配 "
                  f"(0x{IBIAS_REG:08X}=0x{reg_val:08X})")
            time.sleep(1)
        return False

    print(f"每档采样 {duration_s / 60:.1f} 分钟, 共 8 档, 预计总耗时约 {duration_s * 8 / 3600:.1f} 小时")
    saved_paths = []
    for ibias_bit in (0, 1):
        for vi_code in (0b00, 0b01, 0b10, 0b11):
            label = f"IB{ibias_bit}_VI{vi_code}"
            print(f"\n-- 电流档 {label} (RW_ADC_IBIAS bit11={ibias_bit}, RW_ADC_VREFH_IBIAS bit16:17={vi_code:02b}) "
                  f"开始于 {time.strftime('%H:%M:%S')} --")
            time.sleep(SWITCH_DELAY_S)  # 切换前延时
            ok = _set_bits(ibias_bit, 11, 11, "RW_ADC_IBIAS[11]")
            ok = _set_bits(vi_code, 16, 17, "RW_ADC_VREFH_IBIAS[16:17]") and ok
            if not ok:
                print(f"  ❌ {label} 电流档切换失败(已重试{SET_RETRY}次), 跳过该档")
                continue
            time.sleep(SWITCH_DELAY_S)  # 切换后等待稳定

            # 缓存当前寄存器值(已含本档电流配置), 之后整字直接写 START=1/0, 不再读-改-写
            base_val = instr.swd_read_register(IBIAS_REG)
            if base_val is None:
                print(f"  ❌ {label} 读取 0x{IBIAS_REG:08X} 失败, 跳过该档")
                continue
            start_hi = base_val | 0x1
            start_lo = base_val & ~0x1
            instr.swd_write_register(IBIAS_REG, start_lo)  # 确保从 0 开始, 下一次 0→1 才是上升沿

            # write_only 模式逐行写入, 数据量大时内存占用低
            wb = openpyxl.Workbook(write_only=True)
            ws = _new_sheet(wb, "采样数据")
            part, rows_in_sheet, total, fail_cnt, overrun_cnt = 1, 0, 0, 0, 0
            start_time = time.perf_counter()
            next_sample_time = start_time
            next_report = 60.0
            while True:
                # 等到本周期起点 (sleep 精度不够, 用忙等)
                while time.perf_counter() < next_sample_time:
                    pass
                elapsed = time.perf_counter() - start_time
                if elapsed >= duration_s:
                    break

                instr.swd_write_register(IBIAS_REG, start_hi)  # START 上升沿触发转换
                half_time = next_sample_time + INTERVAL_S / 2
                while time.perf_counter() < half_time:
                    pass
                instr.swd_write_register(IBIAS_REG, start_lo)  # 半周期后拉低, 占空比 50%
                raw_val = instr.swd_read_register(REG_ADDR)
                if raw_val is None:
                    fail_cnt += 1
                val = raw_val & 0xFFF if raw_val is not None else None  # 只取 bit0~bit11

                if rows_in_sheet >= MAX_DATA_ROWS:
                    part += 1
                    ws = _new_sheet(wb, f"采样数据_{part}")
                    rows_in_sheet = 0
                ws.append([round(elapsed, 4), val])
                rows_in_sheet += 1
                total += 1

                if elapsed >= next_report:
                    print(f"  [{label}] 已采样 {elapsed / 60:.0f}/{duration_s / 60:.0f} 分钟, {total} 个点")
                    next_report += 60.0

                next_sample_time += INTERVAL_S
                now = time.perf_counter()
                if now > next_sample_time:
                    # SWD 写1/写0/读 三次操作超出 1ms 周期: 从当前时刻重新对齐, 不追赶
                    overrun_cnt += 1
                    next_sample_time = now

            instr.swd_write_register(IBIAS_REG, start_lo)
            excel_path = os.path.join(save_dir, f"adc_3v_ibias_{run_tag}_{label}.xlsx")
            wb.save(excel_path)
            saved_paths.append(excel_path)
            print(f"  采样完成, 共 {total} 个点 (读取失败 {fail_cnt} 个, 超出1ms周期 {overrun_cnt} 次)"
                  + (f", 分 {part} 个sheet" if part > 1 else ""))
            print(f"  ✅ 已保存: {excel_path}")

    print(f"\n全部电流档完成, 共生成 {len(saved_paths)} 个文件")


def test_tj():
    init_pa5_pa6_altfunc14()
    time.sleep(1)
    instr.swd_modify_register(0x40007C00, 0x01, 15, 15)
    time.sleep(1)
    instr.swd_modify_register(0x40007C1C, 0x00, 6, 6)
    instr.swd_modify_register(0x40007C1C, 0x01, 7, 7)
    instr.swd_modify_register(0x40007C0C, 0b10011, 0, 4)
    instr.swd_modify_register(0x40007C1C, 0x01, 5, 5)
    print("tj test")

def test_pll():
    init_pa5_pa6_altfunc14()
    #time.sleep(1)
    #instr.swd_modify_register(0x40007C00, 0x01, 15, 15)
    time.sleep(1)
    instr.swd_modify_register(0x40007C1C, 0x00, 6, 6)
    instr.swd_modify_register(0x40007C1C, 0x01, 7, 7)
    instr.swd_modify_register(0x40007C0C, 0b10001, 0, 4)
    instr.swd_modify_register(0x40007C1C, 0x01, 5, 5)
    print("1")
    #instr.swd_modify_register(0x40005008, 0x03, 19, 22)
    time.sleep(1)
    data = 0x03 << 19
    instr.swd_write_register(0x40005008, data)
    #read_0x40005008 = instr.swd_read_register(0x40005008)
    #print(f"0x40005008 = {read_0x40005008:08X}")
    print("pll test")

def ADC_DNL_Test():

    """new add"""
    PA_ANA_EN = 0x48000050
    RW_ANA_REF_EN = 0x40007C00
    RW_ADC_VREFH_ATP = 0x40007C18
    RW_ATP_MUX = 0x40007C1C
    RW_ATP_MUX_EN = 0x40007C1C
    RW_ADC_VREFH = 0x40007C18
    RW_ANA_VREFH_EN = 0x40007C18
    RW_ADC_VREFH_BYPS = 0x40007C18
    RW_ANA_ADC_EN = 0x40007C00
    RW_ADC_BUF_TP = 0x40007C18
    RW_ADC_MUX = 0x40007C18
    
    # 测试前置: 配置 PA5/PA6 复用 ALT_FUNC_14 (开 GPIO 时钟 + 输出方向 + AF14)
    init_pa5_pa6_altfunc14()
    time.sleep(0.5)

    #使能相关模块
    instr.swd_modify_register(RW_ANA_REF_EN, 0x01, 24, 24)
    instr.swd_modify_register(RW_ANA_ADC_EN, 0x01, 19, 19)
    time.sleep(0.1)

    instr.swd_modify_register(PA_ANA_EN, 0x20, 0, 15)
    instr.swd_modify_register(RW_ADC_VREFH_ATP, 0x01, 18, 18)
    instr.swd_modify_register(RW_ADC_MUX, 0B10010, 3, 7)
    instr.swd_modify_register(RW_ADC_MUX, 0x01, 2, 2)
    time.sleep(0.1)
    instr.swd_modify_register(RW_ATP_MUX_EN, 0x00, 6, 6)
    instr.swd_modify_register(RW_ATP_MUX_EN, 0x01, 7, 7)
    instr.swd_modify_register(RW_ATP_MUX, 0b11111, 0, 4)
    instr.swd_modify_register(RW_ATP_MUX_EN, 0x01, 5, 5)
    instr.swd_modify_register(RW_ADC_VREFH, 0x02, 13, 14)
    instr.swd_modify_register(RW_ADC_VREFH, 0x01, 16, 17)
    instr.swd_modify_register(RW_ANA_VREFH_EN, 0x01, 18, 18)
    instr.swd_modify_register(RW_ADC_BUF_TP, 0x00, 12, 12)
    time.sleep(0.1)
    instr.swd_modify_register(RW_ADC_VREFH_BYPS, 0x01, 15, 15)
    time.sleep(1)
    vdd_vol = instr.dmm_measure_single_voltage()
    time.sleep(1)



def ADC_0V_Measure_Test(adc_channel=18):
    """ADC采集0V电压测试
    流程: 1. 串口切换继电器为 38 39, 用万用表测出实际电压(应接近0V)
          2. 串口切换继电器为 39 5, 切到ADC采集通路
          3. 触发ADC采集并读取转换值
    """
    RCC_ANACKE  = 0x4002005C   # 模拟时钟使能: bit0 ADCCKE=1（必须）
    ADC_CR      = 0x40012460   # 控制: [0]ADC_EN [16]VREFH_EN [21:20]VREFH [5:4]CLK
    ADC_OTR     = 0x40012464   # 软件触发: 写 0x5AA5 产生 ADC_SOC
    ADC_TGSR    = 0x40012448   # 触发源: [19:16]TRIG_EN [3:0]TRIG_SEL0(=8 软件)
    ADC_TGNR    = 0x4001244C   # 组0通道数: [7:0]TGN0
    ADC_SEQR0   = 0x40012450   # 通道选择: [4:0]SEQ0
    ADC_TCFG    = 0x40012468   # 时序: 复位默认(INIT=0x40 不能为0)
    ADC_IR      = 0x40012474   # 状态(只读): [31]RDY [0]S0CDIF
    ADC_DATA0   = 0x40012404   # 序列0转换结果: [11:0]

    PA_ANA_EN = 0x48000050
    RW_ANA_REF_EN = 0x40007C00
    RW_ADC_VREFH_ATP = 0x40007C18
    RW_ATP_MUX = 0x40007C1C
    RW_ATP_MUX_EN = 0x40007C1C
    RW_ADC_VREFH = 0x40007C18
    RW_ANA_VREFH_EN = 0x40007C18
    RW_ADC_VREFH_BYPS = 0x40007C18
    RW_ANA_ADC_EN = 0x40007C00

    VREFH_mV = 3000
    vrefh_code, vrefh_byps = 0b11, 0

    # 测试前置: 配置 PA5/PA6 复用 ALT_FUNC_14
    init_pa5_pa6_altfunc14()
    time.sleep(0.5)

    # 使能相关模块
    instr.swd_modify_register(RW_ANA_REF_EN, 0x01, 24, 24)
    instr.swd_modify_register(RW_ANA_ADC_EN, 0x01, 19, 19)
    time.sleep(0.1)

    instr.swd_modify_register(PA_ANA_EN, 0x20, 0, 15)
    instr.swd_modify_register(RW_ADC_VREFH_ATP, 0x01, 18, 18)
    time.sleep(0.1)
    instr.swd_modify_register(RW_ATP_MUX_EN, 0x00, 6, 6)
    instr.swd_modify_register(RW_ATP_MUX_EN, 0x01, 7, 7)
    instr.swd_modify_register(RW_ATP_MUX, 0x06, 0, 4)
    instr.swd_modify_register(RW_ATP_MUX_EN, 0x01, 5, 5)
    instr.swd_modify_register(RW_ADC_VREFH, vrefh_code, 13, 14)
    instr.swd_modify_register(RW_ADC_VREFH, 0x01, 16, 17)
    instr.swd_modify_register(RW_ANA_VREFH_EN, 0x01, 18, 18)
    time.sleep(0.1)

    # ---------- 1. 初始化 ADC 通路 ----------
    instr.swd_modify_register(RCC_ANACKE, 0x01, 0, 0)
    instr.swd_modify_register(ADC_SEQR0, adc_channel, 0, 4)
    instr.swd_modify_register(ADC_TGNR,   1, 0, 7)
    instr.swd_modify_register(ADC_TGSR,   0b0001, 16, 19)
    instr.swd_modify_register(ADC_TGSR,   0b1000,  0,  3)
    instr.swd_write_register(ADC_TCFG,    0x0A4020F4)
    CR_VAL = 0x00000031 | (vrefh_code << 20) | 0x00010000
    if vrefh_byps:
        CR_VAL |= 0x00020000
    instr.swd_write_register(ADC_CR, CR_VAL)
    time.sleep(0.1)

    # ---------- 2. 切换继电器 38 39, 用万用表实测电压 ----------
    instr.serial_send("cmd_gpio_set_595d_test set 38 39\r\n")
    time.sleep(2)
    real_vol = instr.dmm_measure_single_voltage()
    print(f"实测电压(38 39通路) = {real_vol} mV")
    time.sleep(1)
    instr.serial_send("cmd_gpio_set_595d_test reset\r\n")
    time.sleep(1)

    # ---------- 3. 切换继电器 39 5, 用ADC采集 ----------
    instr.serial_send("cmd_gpio_set_595d_test set 39 5\r\n")
    time.sleep(1)

    instr.swd_write_register(ADC_OTR, 0x5AA5)      # 软件触发一次
    time.sleep(1)
    done = False
    for _ in range(20):
        ir_val = instr.swd_read_register(ADC_IR)
        if ir_val is not None and (ir_val & 0x1):
            done = True
            break
        time.sleep(0.05)

    conv_mv = None
    raw = None
    if not done:
        print("⚠️ ADC触发超时")
    else:
        raw = instr.swd_read_register(ADC_DATA0)
        if raw is None:
            print("⚠️ 读取 ADC_DATA0 失败 (SWD 异常)")
        else:
            raw &= 0xFFF
            conv_mv = raw / 4095.0 * VREFH_mV
            print(f"ADC采集结果: raw={raw}, 换算电压={conv_mv:.3f} mV")

    # ---------- 4. 清理 ----------
    instr.swd_modify_register(ADC_CR, 0x00, 0, 0)
    instr.serial_send("cmd_gpio_set_595d_test reset\r\n")

    return real_vol, raw, conv_mv


def VREF_Voltage_Measure():
    """输出VREF的电压, 配置参考chips/OCD5010/tests.py中的REF_VTRM函数"""
    PA_ANA_EN = 0x48000050
    RW_ANA_REF_EN = 0x40007C00
    PA6_AFS14_EN = 0x4800003C
    RW_ATP_MUX = 0x40007C1C
    RW_ATP_MUX_EN = 0x40007C1C

    # 测试前置: 配置 PA5/PA6 复用 ALT_FUNC_14 (开 GPIO 时钟 + 输出方向 + AF14)
    init_pa5_pa6_altfunc14()
    time.sleep(0.1)

    instr.swd_modify_register(PA_ANA_EN, 0x20, 0, 15)
    time.sleep(0.1)
    instr.swd_modify_register(RW_ANA_REF_EN, 0x01, 24, 24)
    time.sleep(0.1)
    instr.swd_modify_register(PA6_AFS14_EN, 0x0E, 24, 27)
    time.sleep(0.1)

    # 确认 ATP_MUX 选中 VREF
    instr.swd_modify_register(RW_ATP_MUX_EN, 0x00, 6, 6)
    time.sleep(0.1)
    instr.swd_modify_register(RW_ATP_MUX_EN, 0x01, 7, 7)
    time.sleep(0.1)
    instr.swd_modify_register(RW_ATP_MUX_EN, 0x01, 5, 5)
    time.sleep(0.1)
    instr.swd_modify_register(RW_ATP_MUX, 0x02, 0, 4)
    time.sleep(0.5)

    v_vref = instr.dmm_measure_single_voltage()
    print(f"VREF实测电压 = {v_vref:.3f} mV")

    instr.serial_send_and_receive("cmd_gpio_set_595d_test reset\r\n")
    time.sleep(0.5)
    return v_vref


def VREF_TG_ATP_Buf_Diff_Test():
    """多温度点测试: 切换TG/ATP_BUF测量VREF电压并做差
    每个温度点流程:
      1. 切换TG (RW_ATP_MUX_EN bit6=1, bit7=0), 延时1s, 测量电压
      2. 切换ATP_BUF (RW_ATP_MUX_EN bit6=0, bit7=1), 延时1s, 测量电压
      3. 计算差值 = ATP_BUF - TG
    所有温度点完成后在终端汇总显示
    """
    PA_ANA_EN = 0x48000050
    RW_ANA_REF_EN = 0x40007C00
    PA6_AFS14_EN = 0x4800003C
    RW_ATP_MUX = 0x40007C1C
    RW_ATP_MUX_EN = 0x40007C1C

    try:
        temp_input = input("请输入温度点(多个温度点用逗号分隔, 例如: 25,-40,85): ").strip()
        temp_list = [float(t.strip()) for t in temp_input.split(",") if t.strip() != ""]
    except ValueError:
        print("❌ 温度点输入无效，测试终止")
        return
    if not temp_list:
        print("❌ 未输入有效温度点，测试终止")
        return

    # 前置: 配置 PA5/PA6 复用 ALT_FUNC_14 (开 GPIO 时钟 + 输出方向 + AF14)
    init_pa5_pa6_altfunc14()
    time.sleep(0.1)

    instr.swd_modify_register(PA_ANA_EN, 0x20, 0, 15)
    time.sleep(0.1)
    instr.swd_modify_register(RW_ANA_REF_EN, 0x01, 24, 24)
    time.sleep(0.1)
    instr.swd_modify_register(PA6_AFS14_EN, 0x0E, 24, 27)
    time.sleep(0.1)
    instr.swd_modify_register(RW_ATP_MUX, 0x02, 0, 4)
    time.sleep(0.1)

    results = []
    for temp in temp_list:
        input(f"请将温箱设置为 {temp}℃ 并等待温度稳定后按回车键继续...")

        # 切换TG: bit6=1, bit7=0
        instr.swd_modify_register(RW_ATP_MUX_EN, 0x01, 6, 6)
        instr.swd_modify_register(RW_ATP_MUX_EN, 0x00, 7, 7)
        time.sleep(1)
        v_tg = instr.dmm_measure_single_voltage()

        # 切换ATP_BUF: bit6=0, bit7=1
        instr.swd_modify_register(RW_ATP_MUX_EN, 0x00, 6, 6)
        instr.swd_modify_register(RW_ATP_MUX_EN, 0x01, 7, 7)
        time.sleep(1)
        v_atp_buf = instr.dmm_measure_single_voltage()

        diff = v_atp_buf - v_tg
        print(f"温度={temp}℃, TG电压={v_tg:.3f} mV, ATP_BUF电压={v_atp_buf:.3f} mV, 差值={diff:.3f} mV")
        results.append((temp, v_tg, v_atp_buf, diff))

    instr.serial_send_and_receive("cmd_gpio_set_595d_test reset\r\n")

    print("\n========== 各温度点测试结果汇总 ==========")
    print(f"{'温度(℃)':>10} {'TG电压(mV)':>14} {'ATP_BUF电压(mV)':>18} {'差值(mV)':>12}")
    for temp, v_tg, v_atp_buf, diff in results:
        print(f"{temp:>10} {v_tg:>14.3f} {v_atp_buf:>18.3f} {diff:>12.3f}")

    return results


def PLL_CKO_TRIM_Test():
    """PLL VCO 测试"""
    init_pa5_pa6_altfunc14()
    instr.serial_send_and_receive("cmd_gpio_set_595d_test reset\r\n")
    time.sleep(1)
    instr.swd_modify_register(0x40020058, 0x01, 16, 16)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C00, 0x01, 21, 21)
    time.sleep(0.1)
    instr.swd_modify_register(0x48000050, 0x20, 0, 15)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0x01, 5, 5)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0x01, 6, 6)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0b10001, 0, 4)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C0C, 0x00, 19, 19)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0x01, 13, 13)     #使能DTP_MUX
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0b10000, 8, 12)   #选择PLL_CKO_DTP
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C0C, 0x01, 17, 17)
    time.sleep(0.1)

    instr.swd_modify_register(0x48000818, 0b1, 0, 0)   # OEC_STR (0x18), bit0 = 1
    instr.swd_modify_register(0x48000814, 0b1, 0, 0)   # OES_STR (0x14), bit0 = 1
    time.sleep(0.1)
    instr.swd_modify_register(0x4800084C, 0b1, 0, 0)   # AFC_STR (0x4C), bit0 = 1 （先清）
    instr.swd_modify_register(0x4800083C, 0b0001, 0, 3) # AF0_STR (0x3C), pin0字段bits[3:0] = 1
    time.sleep(0.1)
    instr.swd_modify_register(0x4800085C, 0b1, 0, 0)   # PUC_STR (0x5C), bit0 = 1
    instr.swd_modify_register(0x48000864, 0b1, 0, 0)   # PDC_STR (0x64), bit0 = 1
    time.sleep(0.1)
    instr.swd_modify_register(0x48000878, 0b1, 0, 0)   # DRVC_STR (0x78), bit0 = 1 （先清）
    instr.swd_modify_register(0x48000868, 0b001, 0, 2) # DRVS0_STR (0x68), pin0字段bits[2:0] = 1
    time.sleep(0.1)
    instr.swd_modify_register(0x40020054, 0b011, 0, 3)
    time.sleep(0.1)
    """
    instr.swd_modify_register(0x40007C0C, 0x03, 3, 6)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C0C, 0x02, 15, 16)
    time.sleep(0.1)
    """




def PLL_CKO_VS_VIN_Test_Simple():
    """简易版 PLL_CKO 频率随输入电压变化测试
    用户输入电压范围/步进值/RW_PLL_VCO取值，每次步进由用户按回车确认后再测量（便于查看波形），不保存Excel"""
    init_pa5_pa6_altfunc14()
    instr.serial_send_and_receive("cmd_gpio_set_595d_test reset\r\n")
    time.sleep(1)
    instr.swd_modify_register(0x40020058, 0x01, 16, 16)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C00, 0x01, 21, 21)
    time.sleep(0.1)
    instr.swd_modify_register(0x48000050, 0x20, 0, 15)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0x01, 5, 5)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0x01, 6, 6)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0b10001, 0, 4)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C0C, 0x01, 19, 19)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0x01, 13, 13)     # 使能DTP_MUX
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C1C, 0b10000, 8, 12)   # 选择PLL_CKO_DTP
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C0C, 0x01, 17, 17)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C0C, 0x01, 19, 19)
    time.sleep(0.1)
    instr.swd_modify_register(0x40007C0C, 0x03, 3, 6)
    time.sleep(0.1)

    time.sleep(1)

    try:
        v_start = float(input("请输入起始电压(mV): ").strip())
        v_end = float(input("请输入终止电压(mV): ").strip())
        v_step = float(input("请输入步进值(mV): ").strip())
        vco_code = int(input("请输入RW_PLL_VCO的值(0~7): ").strip())
    except ValueError:
        print("❌ 输入无效，测试终止")
        return
    if v_step == 0:
        print("❌ 步进值不能为0，测试终止")
        return
    if vco_code < 0 or vco_code > 7:
        print("❌ RW_PLL_VCO 取值必须在0~7之间，测试终止")
        return
    if (v_end - v_start) * v_step < 0:
        v_step = -v_step

    vin_list = []
    v = v_start
    if v_step > 0:
        while v <= v_end + 1e-9:
            vin_list.append(round(v, 4))
            v += v_step
    else:
        while v >= v_end - 1e-9:
            vin_list.append(round(v, 4))
            v += v_step

    print(f"共 {len(vin_list)} 个测试点: {vin_list}")

    instr.swd_modify_register(0x40007C0C, vco_code, 12, 14)
    time.sleep(0.5)
    print(f"已设置 RW_PLL_VCO = {vco_code:03b}")

    for idx, set_vol in enumerate(vin_list, start=1):
        instr.serial_send(f"spi_cmd_set_val mV {set_vol}\r\n")
        time.sleep(0.5)

        instr.serial_send(f"cmd_gpio_set_595d_test set 38 6\r\n")
        time.sleep(2)
        real_vol = instr.dmm_measure_single_voltage()
        print(f"[{idx}/{len(vin_list)}] 设置电压={set_vol} mV, 实测电压={real_vol:.3f} mV")

        time.sleep(1)
        instr.serial_send(f"cmd_gpio_set_595d_test set 15 6\r\n")
        time.sleep(1)

        input("按回车键继续查看波形并测量频率...")

    instr.serial_send("cmd_gpio_set_595d_test reset\r\n")


def serial_receive_to_excel(port="COM1", baudrate=2000000, save_every=1000):
    """打开指定串口, 按行持续接收数据存入 Excel, Ctrl+C 停止并保存"""
    import serial
    from datetime import datetime

    save_dir = os.path.join(_project_root, "Save_Results", "OCD5010")
    os.makedirs(save_dir, exist_ok=True)
    save_path = os.path.join(save_dir, f"serial_rx_{datetime.now():%Y%m%d_%H%M%S}.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "serial_rx"
    ws.append(["序号", "时间", "数据"])

    ser = serial.Serial(port, baudrate, timeout=1)
    print(f"✅ {port} 已打开, 波特率 {baudrate}, 开始接收 (Ctrl+C 停止)")
    print(f"数据保存到: {save_path}")
    count = 0
    try:
        while True:
            line = ser.readline()
            if not line:
                continue
            text = line.decode("utf-8", errors="replace").strip()
            if not text:
                continue
            count += 1
            ws.append([count, datetime.now().strftime("%H:%M:%S.%f")[:-3], text])
            print(f"[{count}] {text}")
            if count % save_every == 0:
                wb.save(save_path)
    except KeyboardInterrupt:
        print("\n停止接收")
    finally:
        ser.close()
        wb.save(save_path)
        print(f"✅ 共接收 {count} 行, 已保存: {save_path}")


def serial_receive_by_current_gear(gears, port="COM1", baudrate=2000000, duration_s=6 * 60):
    """按电流档位分段接收串口数据, 每档接收 duration_s 秒, 每档单独存一个 CSV 文件

    gears: [(RW_ADC_IBIAS bit11, RW_ADC_VREFH_IBIAS bit16:17), ...], 仅提示用户手动切换, 不自动写寄存器

    数据量大 (2Mbps x 18min 可达数百MB, 远超 Excel 104万行上限), 故按档位流式写入 CSV:
        序号, 相对时间(s), 数据
    CSV 可直接用 Excel / pandas 打开。
    """
    import csv
    import serial
    from datetime import datetime

    session_dir = os.path.join(_project_root, "Save_Results", "OCD5010",
                               f"serial_rx_{datetime.now():%Y%m%d_%H%M%S}")
    os.makedirs(session_dir, exist_ok=True)
    print(f"数据保存目录: {session_dir}")

    ser = serial.Serial(port, baudrate, timeout=0.5)
    print(f"✅ {port} 已打开, 波特率 {baudrate}")
    try:
        for gi, (ibias_bit, vi_code) in enumerate(gears, 1):
            gear = f"IB{ibias_bit}_VI{vi_code}"
            print(f"\n[{gi}/{len(gears)}] 请切换电流档位到【{gear}】:")
            print(f"    0x40007C18 RW_ADC_IBIAS       bit11    = {ibias_bit}")
            print(f"    0x40007C18 RW_ADC_VREFH_IBIAS bit16:17 = {vi_code:02b}")
            input("  切换完成后按回车开始接收...")
            save_path = os.path.join(session_dir, f"{gi:02d}_{gear}.csv")
            ser.reset_input_buffer()    # 丢弃切档期间的残留数据
            count = 0
            bad = 0
            buf = b""
            t0 = time.time()
            next_report = 60
            with open(save_path, "w", newline="", encoding="utf-8-sig", buffering=1024 * 1024) as f:
                writer = csv.writer(f)
                writer.writerow(["序号", "相对时间(s)", "ADC值"])
                try:
                    while True:
                        elapsed = time.time() - t0
                        if elapsed >= duration_s:
                            break
                        if elapsed >= next_report:
                            print(f"  [{gear}] 已接收 {elapsed / 60:.0f}/{duration_s / 60:.0f} 分钟, "
                                  f"{count} 个数据, 无效行 {bad}")
                            f.flush()
                            next_report += 60
                        # 按块读取再拆行, 避免 readline 逐字节读取跟不上 2Mbps
                        chunk = ser.read(ser.in_waiting or 1)
                        if not chunk:
                            continue
                        buf += chunk
                        *lines, buf = buf.split(b"\n")
                        ts = f"{elapsed:.4f}"
                        for line in lines:
                            # 下位机格式: INFO("%d,0,0,0\r\n", ADC_ptr->DATA[0]), 只取第一个字段
                            field = line.split(b",", 1)[0].strip()
                            try:
                                val = int(field)
                            except ValueError:
                                if field:
                                    bad += 1
                                continue
                            count += 1
                            writer.writerow([count, ts, val])
                except KeyboardInterrupt:
                    print(f"\n  [{gear}] 手动中断当前档位")
            print(f"✅ 档位【{gear}】完成, 共 {count} 个数据, 无效行 {bad}, 已保存: {save_path}")
    except KeyboardInterrupt:
        print("\n已停止")
    finally:
        ser.close()
        print(f"串口已关闭, 全部数据在: {session_dir}")


# 电流档位列表, 同 adc_3v_ibias_sample: 0x40007C18 RW_ADC_IBIAS bit11 x RW_ADC_VREFH_IBIAS bit16:17, 共 8 档
CURRENT_GEARS = [(ib, vi) for ib in (0, 1) for vi in (0b00, 0b01, 0b10, 0b11)]


if __name__ == "__main__":
    serial_receive_by_current_gear(CURRENT_GEARS, port="COM1", baudrate=2000000, duration_s=6 * 60)
    # serial_receive_to_excel(port="COM1", baudrate=2000000)

    # check_init("SWD", chip.init_swd(instr))
    # check_init("串口", serial_find_and_init(keyword=SERIAL_KEYWORD, baudrate=SERIAL_BAUDRATE))
    # check_init("万用表", instr.dmm_init())
    # time.sleep(1)
    # print("初始化完毕，可以开始测试......")
    # #reg_0x40012404_bit0_11_sample()
    # adc_3v_ibias_sample()
    
    