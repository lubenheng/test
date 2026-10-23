# -*- coding: utf-8 -*-
"""
把工程内的 MTP 配置表输出成 Excel 表格，
供程序「菜单 [32] MTP 全片擦除+烧录+校验」直接加载，也便于在 Excel 中查看/对照。

数据来源（默认）：chips/OCD2361/mtp_config_table.py
    在该文件中查看/修改 MTP_CONFIG_ROWS，运行本脚本即可重新生成表格。
    python other_test/mtp/build_mtp_template.py
    python other_test/mtp/build_mtp_template.py --from-source   # 改从原始 xlsx 读取
    python other_test/mtp/build_mtp_template.py --from-txt Save_Results/MTP_Table/mtp_config_table.txt
                                    # 从导出的 .txt 配置表生成（菜单 e 也走这条路径）

输入格式：MTP_CONFIG_ROWS = [(ADDR, INDEX, B31, B23, B15, B7, 模块号), ...]

输出：Save_Results/MTP_Table/MTP_配置寄存器（测试版本2）_烧录模板.xlsx
      （--from-txt / --out 时按指定路径，否则该 txt 同目录的 <txt名>_烧录模板.xlsx）
      · 工作表 "MTP批量模板"：程序可直接读取（表头 2 行 + PROGRAM 数据）
        列：OPERATION / ADDR / INDEX / [31:24] / [23:16] / [15:8] / [7:0]
            / 物理地址 / 寄存器ADDR / 计入CRC
      · 工作表 "说明"：CRC、地址分布、被排除的非 PROGRAM 行
"""

import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 与 main.py 一致：Windows 控制台默认 GBK，直接 print ✅/❌ 会抛 UnicodeEncodeError
if os.name == 'nt':
    os.system('chcp 65001 > nul')
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "MTP_配置寄存器（测试版本2）.xlsx"
# 直接从配置模块生成时也落到 Save_Results/MTP_Table，与 txt 配置表放在一起
DST = ROOT / "Save_Results" / "MTP_Table" / "MTP_配置寄存器（测试版本2）_烧录模板.xlsx"

sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "instrument" / "IIC"))
from iic_ops import IICLogic                      # noqa: E402
from crc import calculate_crc32                   # noqa: E402
from mtp_batch_ops import MTPBatchLogic           # noqa: E402

SHEET_MAIN = "MTP批量模板"
CRC_ADDR_LIMIT = 416          # CRC 覆盖 phys < 416
CRC_RSV_LO, CRC_RSV_HI = 507, 510

HEADERS = ["OPERATION", "ADDR(基地址)", "INDEX(Word偏移)",
           "[31:24]", "[23:16]", "[15:8]", "[7:0]", "物理地址(计算)", "寄存器ADDR(原表)",
           "计入CRC"]
NOTES = ["READ/PROGRAM/ERASE_SECTOR", "如 0~511 (MTP基地址)", "Word偏移(1代表偏4字节)",
         "对应物理地址 ADDR+INDEX*4 +3 (HEX)", "对应物理地址 ADDR+INDEX*4 +2 (HEX)",
         "对应物理地址 ADDR+INDEX*4 +1 (HEX)", "对应物理地址 ADDR+INDEX*4 +0 (HEX)",
         "自动公式: =B列+C列*4", "寄存器模块地址(仅供参考)",
         "Y=计入 CRC，空=只烧录（phys ≥ 416）"]

# 表格样式
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
NO_CRC_FILL = PatternFill("solid", fgColor="FFF2CC")   # phys ≥ 416：只烧录不进 CRC
MONO_FONT = Font(name="Consolas")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
COL_WIDTHS = (18, 14, 18, 10, 10, 10, 10, 16, 18, 10)


def parse_source():
    """从原始 xlsx 读取，返回 (program_rows, other_rows)。

    program_rows: [(addr, index, [b31, b23, b15, b7], mod_addr)]
    other_rows:   [(row_no, 原始单元格列表)]  —— 非 PROGRAM 行（如 WRITETHENREAD）
    """
    ws = openpyxl.load_workbook(SRC, data_only=True)[SHEET_MAIN]
    program, other = [], []
    for r in range(3, ws.max_row + 1):
        cells = [ws.cell(r, c).value for c in range(1, 11)]   # A~J，J 为「寄存器ADDR」
        if all(v is None or str(v).strip() == "" for v in cells[:3]):
            continue                      # 空行：跳过
        op = str(cells[0]).strip().upper() if cells[0] else ""
        if op != "PROGRAM":
            other.append((r, cells))
            continue
        addr = int(cells[1])
        index = int(cells[2])
        data = [IICLogic.parse_hex_byte(cells[c], f"第{r}行第{c+1}列")
                for c in range(3, 7)]
        program.append((addr, index, data, cells[9]))
    return program, other


def parse_from_module():
    """从 chips/OCD2361/mtp_config_table.py 读取数据。

    返回与 parse_source() 相同的结构。OTHER_ROWS 直接沿用。
    """
    sys.path.insert(0, str(ROOT))
    from chips.OCD2361 import mtp_config_table as cfg

    program = [(addr, index, [b31, b23, b15, b7], mod)
               for addr, index, b31, b23, b15, b7, mod in cfg.MTP_CONFIG_ROWS]
    other = [(r, [op, addr, index, *data])
             for r, op, addr, index, *data in cfg.OTHER_ROWS] if getattr(cfg, "OTHER_ROWS", None) else []
    return program, other


def _txt_int(token):
    """txt 里的整数：0x 前缀按 HEX，其余按十进制。"""
    s = str(token).strip()
    return int(s, 16) if s.lower().startswith("0x") else int(s, 10)


def parse_from_txt(path):
    """从 mtp_config_table.txt（build_mtp_config_txt.py 的产物）读取数据。

    数据行格式：
        行号  ADDR  INDEX  [31:24]  [23:16]  [15:8]  [7:0]  PHYS  CRC?  模块
    「源表中未烧录的行」小节（OPERATION 不是 PROGRAM）收进 other，供说明页展示。

    返回与 parse_source() 相同的结构。
    """
    program, other = [], []
    for lineno, line in enumerate(Path(path).read_text(encoding="utf-8", errors="replace").splitlines(), 1):
        parts = line.split()
        if len(parts) < 7 or not parts[0].isdigit():
            continue                       # 说明行 / 表头 / 图例
        if not parts[1].lower().startswith("0x"):
            # 非 PROGRAM 参考行：行号 OPERATION [31:24] [23:16] [15:8] [7:0]
            try:
                data = [_txt_int(p) for p in parts[2:6]]
            except ValueError:
                continue
            if len(data) == 4:
                other.append((int(parts[0]), [parts[1], "", "", *data]))
            continue
        try:
            addr = _txt_int(parts[1])
            index = _txt_int(parts[2])
            data = [_txt_int(p) for p in parts[3:7]]
        except ValueError as e:
            raise ValueError(f"第 {lineno} 行解析失败：{e}")
        if not all(0 <= b <= 0xFF for b in data):
            raise ValueError(f"第 {lineno} 行数据超出 1 字节：{parts[3:7]}")
        mod = parts[9] if len(parts) > 9 else None
        program.append((addr, index, data, mod))
    return program, other


def txt_header_crc(path):
    """取 txt 文件头里记录的 CRC32（没有则返回 None）。"""
    import re
    m = re.search(r"CRC32\s*:\s*0x([0-9A-Fa-f]+)",
                  Path(path).read_text(encoding="utf-8", errors="replace"))
    return int(m.group(1), 16) if m else None


def is_mtp_config_txt(path):
    """判断一个 .txt 是不是 MTP 配置表（而不是测试结果记录）。

    先看文件名，再看文件头/表头，避免误把测试结果 txt 当成配置表。
    """
    p = Path(path)
    if "mtp_config_table" in p.name.lower():
        return True
    try:
        head = p.read_text(encoding="utf-8", errors="replace")[:2000]
    except Exception:
        return False
    return "MTP 配置寄存器表" in head or ("PHYS" in head and "[31:24]" in head)


def compute_crc(program):
    """按烧录口径算 CRC：声明字节 ∩ phys<416，crc.py(LSB 优先)。"""
    buffer = bytearray(512)
    burn_addrs = set()
    for addr, index, data, _ in program:
        phys = addr + index * 4
        buffer[phys:phys + 4] = bytes([data[3], data[2], data[1], data[0]])
        burn_addrs.update(range(phys, phys + 4))
    burn_addrs = sorted(burn_addrs)
    crc_addrs = [a for a in burn_addrs if a < CRC_ADDR_LIMIT]
    final_crc = calculate_crc32(bytes(buffer[a] for a in crc_addrs))
    return burn_addrs, crc_addrs, final_crc


def _style_range(ws, row_from, row_to, col_from, col_to):
    """给区域加边框 + 居中。"""
    for r in range(row_from, row_to + 1):
        for c in range(col_from, col_to + 1):
            cell = ws.cell(r, c)
            cell.border = BORDER
            cell.alignment = CENTER


def build_main_sheet(wb, program, final_crc):
    """生成程序可读取的 "MTP批量模板" 工作表。"""
    ws = wb.active
    ws.title = SHEET_MAIN

    ws.append(HEADERS)
    ws.append(NOTES)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(1, c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        ws.cell(2, c).font = Font(italic=True, color="808080")
        ws.cell(2, c).alignment = CENTER
        ws.cell(2, c).border = BORDER

    row = 3
    prev_addr = None
    for addr, index, data, mod in program:
        phys = addr + index * 4
        in_crc = phys < CRC_ADDR_LIMIT
        ws.append(["PROGRAM", addr, index,
                   f"{data[0]:02X}", f"{data[1]:02X}", f"{data[2]:02X}", f"{data[3]:02X}",
                   phys, mod, "Y" if in_crc else ""])
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row, c)
            cell.border = BORDER
            cell.alignment = CENTER
            if c in (4, 5, 6, 7):
                cell.font = MONO_FONT
            if not in_crc:
                cell.fill = NO_CRC_FILL
        # 不同 ADDR 之间加一条上边框，视觉分块
        if prev_addr is not None and addr != prev_addr:
            for c in range(1, len(HEADERS) + 1):
                cur = ws.cell(row, c)
                cur.border = Border(left=THIN, right=THIN, bottom=THIN,
                                    top=Side(style="medium", color="000000"))
        prev_addr = addr
        row += 1

    # CRC 行：仅作记录，解析时因 OPERATION != PROGRAM 被忽略
    ws.append(["PROGRAM(CRC)", 507, 0,
               f"{(final_crc >> 24) & 0xFF:02X}", f"{(final_crc >> 16) & 0xFF:02X}",
               f"{(final_crc >> 8) & 0xFF:02X}", f"{final_crc & 0xFF:02X}",
               507, "自动生成", ""])
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row, c)
        cell.border = BORDER
        cell.alignment = CENTER
        cell.font = Font(bold=True, color="C00000")
    for c in (4, 5, 6, 7):
        ws.cell(row, c).font = Font(bold=True, color="C00000", name="Consolas")

    for i, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A3"                 # 固定表头两行
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{row - 1}"


def build_info_sheet(wb, program, other, burn_addrs, crc_addrs, final_crc, origin):
    """生成 "说明" 工作表。"""
    info = wb.create_sheet("说明")
    info.append(["项目", "内容"])
    info.cell(1, 1).font = info.cell(1, 2).font = HEADER_FONT
    info.cell(1, 1).fill = info.cell(1, 2).fill = HEADER_FILL
    for k, v in [
        ("数据来源", origin),
        ("PROGRAM 行数", len(program)),
        ("烧录字节数", f"{len(burn_addrs)}（地址 {burn_addrs[0]}~{burn_addrs[-1]}）"),
        ("CRC 覆盖字节数", f"{len(crc_addrs)}（地址 {crc_addrs[0]}~{crc_addrs[-1]}，phys<416）"),
        ("CRC32", f"0x{final_crc:08X}"),
        ("CRC 存放", f"地址 {CRC_RSV_LO}~{CRC_RSV_HI} 小端（由烧录流程自动写入）"),
        ("地址映射", "phys = ADDR + INDEX*4"),
        ("数值格式", "D~G 列统一为 2 位 HEX"),
        ("黄色行", "phys ≥ 416：只烧录，不计入 CRC"),
    ]:
        info.append([k, v])

    # 地址块分布
    info.append([])
    info.append(["地址块分布"])
    info.cell(info.max_row, 1).font = Font(bold=True)
    info.append(["ADDR", "INDEX 范围", "字数", "物理地址范围", "计入CRC"])
    blocks = {}
    for addr, index, _data, _mod in program:
        blocks.setdefault(addr, []).append(index)
    for addr in sorted(blocks):
        idxs = sorted(blocks[addr])
        start, end = addr + idxs[0] * 4, addr + idxs[-1] * 4
        info.append([f"0x{addr:02X}", f"{idxs[0]}~{idxs[-1]}", len(idxs),
                     f"{start}~{end}", "Y" if start < CRC_ADDR_LIMIT else ""])

    # 被排除的行
    if other:
        info.append([])
        info.append(["非 PROGRAM 行（未烧录，保留供参考）"])
        info.cell(info.max_row, 1).font = Font(bold=True)
        info.append(["源行号", "OPERATION", "ADDR", "INDEX", "[31:24]", "[23:16]", "[15:8]", "[7:0]"])
        for r, cells in other:
            info.append([r] + [cells[i] for i in range(7)])

    for i, width in enumerate((16, 22, 10, 18, 10), start=1):
        info.column_dimensions[get_column_letter(i)].width = width
    _style_range(info, 2, info.max_row, 1, 2)


def build_workbook(program, other, origin, expected_crc=None):
    """按 program/other 生成工作簿，返回 (wb, burn_addrs, crc_addrs, final_crc)。"""
    burn_addrs, crc_addrs, final_crc = compute_crc(program)
    if not burn_addrs:
        raise ValueError("没有可烧录的数据行")
    if expected_crc is not None and expected_crc != final_crc:
        print(f"⚠️ 文件头记录的 CRC32 0x{expected_crc:08X} 与按数据重算的 "
              f"0x{final_crc:08X} 不一致（配置表被改过但未重新生成？）")

    wb = openpyxl.Workbook()
    build_main_sheet(wb, program, final_crc)
    build_info_sheet(wb, program, other, burn_addrs, crc_addrs, final_crc, origin)
    return wb, burn_addrs, crc_addrs, final_crc


def txt_to_template(txt_path, xlsx_path=None):
    """把 MTP 配置表 txt 转成烧录模板 Excel。

    参数：
        txt_path:   str/Path - mtp_config_table.txt
        xlsx_path:  str/Path - 输出路径；None 则输出到 txt 同目录的
                    <txt名>_烧录模板.xlsx
    返回：
        (ok, 路径或错误信息)
    """
    txt_path = Path(txt_path)
    if xlsx_path is None:
        xlsx_path = txt_path.with_name(txt_path.stem + "_烧录模板.xlsx")
    xlsx_path = Path(xlsx_path)

    try:
        program, other = parse_from_txt(txt_path)
    except Exception as e:
        return False, f"解析配置表失败：{type(e).__name__}: {e}"
    if not program:
        return False, "文本中没有 PROGRAM 数据行（不是 MTP 配置表？）"

    try:
        wb, burn_addrs, crc_addrs, final_crc = build_workbook(
            program, other, str(txt_path), txt_header_crc(txt_path))
    except Exception as e:
        return False, f"生成表格失败：{type(e).__name__}: {e}"

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    replaced = xlsx_path.exists()
    try:
        wb.save(xlsx_path)
    except PermissionError:
        return False, f"无法写入（文件正被 Excel 等程序占用，请先关闭）：{xlsx_path.name}"
    if replaced:
        print(f"   （已覆盖同名文件 {xlsx_path.name}）")
    print(f"   PROGRAM 行 {len(program)}，烧录字节 {len(burn_addrs)}"
          f"（{burn_addrs[0]}~{burn_addrs[-1]}），CRC32 = 0x{final_crc:08X}")
    return True, str(xlsx_path)


def main():
    from_source = "--from-source" in sys.argv
    from_txt = None
    out_path = None
    argv = sys.argv[1:]
    for i, a in enumerate(argv):
        if a == "--from-txt" and i + 1 < len(argv):
            from_txt = argv[i + 1]
        elif a.startswith("--from-txt="):
            from_txt = a.split("=", 1)[1]
        elif a == "--out" and i + 1 < len(argv):
            out_path = argv[i + 1]
        elif a.startswith("--out="):
            out_path = a.split("=", 1)[1]

    if from_txt:
        ok, msg = txt_to_template(from_txt, out_path)
        print(("✅ 已生成：" if ok else "❌ ") + msg)
        return 0 if ok else 1

    program, other = parse_source() if from_source else parse_from_module()
    origin = SRC.name if from_source else "chips/OCD2361/mtp_config_table.py"
    print(f"数据来源：{origin}")
    print(f"PROGRAM 行：{len(program)}，其它行：{len(other)}")

    try:
        wb, burn_addrs, crc_addrs, final_crc = build_workbook(program, other, origin)
    except ValueError as e:
        print(f"❌ {e}")
        return 1
    print(f"烧录字节：{len(burn_addrs)}（{burn_addrs[0]}~{burn_addrs[-1]}）")
    print(f"CRC 覆盖：{len(crc_addrs)} 字节（{crc_addrs[0]}~{crc_addrs[-1]}），"
          f"CRC32 = 0x{final_crc:08X}")

    DST.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(DST)
    except PermissionError:
        print(f"❌ 无法写入（文件正被 Excel 等程序占用，请先关闭）：{DST.name}")
        return 1
    print(f"✅ 已生成：{DST}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
