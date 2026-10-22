# -*- coding: utf-8 -*-
import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox, filedialog
import serial
import time
import json
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import re

CONFIG_FILE = "config.json"
DEFAULT_CONFIG = {
    "port": "COM13",
    "addr": "1",
    "index": "1"
}

# 串口时序参数（按设备响应情况可继续微调）
READ_CMD_SETTLE_DELAY = 0.01
WRITE_THEN_READ_GAP = 0.01


def reverse_bits_32(x):
    x = ((x & 0x55555555) << 1) | ((x & 0xAAAAAAAA) >> 1)
    x = ((x & 0x33333333) << 2) | ((x & 0xCCCCCCCC) >> 2)
    x = ((x & 0x0F0F0F0F) << 4) | ((x & 0xF0F0F0F0) >> 4)
    x = ((x & 0x00FF00FF) << 8) | ((x & 0xFF00FF00) >> 8)
    x = ((x & 0x0000FFFF) << 16) | ((x & 0xFFFF0000) >> 16)
    return x


def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except:
            pass
    return DEFAULT_CONFIG.copy()


def save_config(port, addr, index):
    config = {
        "port": str(port).strip(),
        "addr": str(addr).strip(),
        "index": str(index).strip()
    }
    try:
        with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump(config, f, indent=4)
    except Exception as e:
        print(f"保存配置失败: {e}")


def reverse_bits(x, width):
    r = 0
    for _ in range(width):
        r = (r << 1) | (x & 1)
        x >>= 1
    return r


def reverse_byte(b):
    b = ((b & 0xF0) >> 4) | ((b & 0x0F) << 4)
    b = ((b & 0xCC) >> 2) | ((b & 0x33) << 2)
    b = ((b & 0xAA) >> 1) | ((b & 0x55) << 1)
    return b


def format_binary(value, width):
    bin_str = f"{value:0{width}b}"
    return ' '.join([bin_str[i:i+4] for i in range(0, len(bin_str), 4)])


def send_d9_spi(rw, addr, index, write_value=0x00000000, port="COM15", baudrate=921600, log_callback=None):
    def log(msg):
        if log_callback:
            log_callback(msg)
        else:
            print(msg)

    if not (0 <= addr <= 31):
        raise ValueError("addr 必须是 0~31")
    if not (0 <= index <= 63):
        raise ValueError("index 必须是 0~63")

    if rw == 'read':
        rwcode = 0b110
        recv_len = 4
        is_write = False
    elif rw == 'write':
        rwcode = 0b010
        recv_len = 0
        is_write = True
    else:
        raise ValueError("rw 必须是 'read' 或 'write'")

    addr_rev = reverse_bits(addr, 5)
    index_rev = reverse_bits(index, 6)
    header_logical = (rwcode << 13) | (addr_rev << 8) | (index_rev << 2)
    log(f"Logical Header (Binary): {format_binary(header_logical, 16)}")

    high_byte = (header_logical >> 8) & 0xFF
    low_byte = header_logical & 0xFF
    payload_bytes = [
        reverse_byte(high_byte),
        reverse_byte(low_byte)
    ]

    if is_write:
        physical_value = reverse_bits_32(write_value)
        data_bytes = [
            (physical_value >> 24) & 0xFF,
            (physical_value >> 16) & 0xFF,
            (physical_value >> 8) & 0xFF,
            physical_value & 0xFF
        ]
        reversed_data = [reverse_byte(b) for b in data_bytes]
        payload_bytes.extend(reversed_data)
        send_len = 6
    else:
        send_len = 2

    cmd = bytearray([
        0x07, 0x00, 0x0F,
        0x00, 0x00, 0x01, 0x00, 0x01,
        0x01, 0x00
    ])
    cmd.extend([(send_len >> 8) & 0xFF, send_len & 0xFF])
    cmd.extend([(recv_len >> 8) & 0xFF, recv_len & 0xFF])
    cmd.extend(payload_bytes)

    op_name = "WRITE" if is_write else "READ"
    log(f"[{port}] {op_name} | addr={addr}, index={index} | logical_header=0x{header_logical:04X}")
    if is_write:
        log(f"  -> 写入值: 0x{write_value:08X}")

    try:
        with serial.Serial(port, baudrate, timeout=1) as ser:
            # 先清空并短暂排空接收口，减少上一条命令延迟上报的残留干扰
            ser.reset_input_buffer()
            ser.reset_output_buffer()
            drain_deadline = time.time() + 0.03
            while time.time() < drain_deadline:
                pending = ser.in_waiting
                if pending > 0:
                    ser.read(pending)
                else:
                    time.sleep(0.002)

            if is_write:
                ser.write(cmd)
                return None

            def read_one_frame(enable_log=True):
                ser.write(cmd)
                time.sleep(READ_CMD_SETTLE_DELAY)

                recv_buf = bytearray()
                read_deadline = time.time() + 0.01
                while time.time() < read_deadline:
                    pending = ser.in_waiting
                    if pending > 0:
                        recv_buf.extend(ser.read(pending))
                        if len(recv_buf) >= 6:
                            time.sleep(0.005)
                    else:
                        time.sleep(0.002)

                response = bytes(recv_buf[-6:]) if len(recv_buf) >= 6 else bytes(recv_buf)
                if len(response) != 6:
                    if enable_log:
                        hex_resp = ' '.join(f"{b:02X}" for b in response) if response else "(无数据)"
                        log(f"  <- 接收失败 ({len(response)}字节，累计{len(recv_buf)}字节): {hex_resp}")
                    return None

                raw_bytes = response[2:6]
                if enable_log and len(recv_buf) > 6:
                    log(f"  <- 本次累计接收 {len(recv_buf)} 字节，已取最后6字节作为有效帧")
                if enable_log:
                    log(f"  <- 接收原始字节 (后4): {' '.join(f'{b:02X}' for b in raw_bytes)}")

                physical_bytes = [reverse_byte(b) for b in raw_bytes]
                physical_value = (
                    (physical_bytes[0] << 24) |
                    (physical_bytes[1] << 16) |
                    (physical_bytes[2] << 8) |
                    physical_bytes[3]
                )
                recv_bit_str = format_binary(physical_value, 32)
                if enable_log:
                    log(f"  <- 接收的物理 bit 流 (MSB→LSB): {recv_bit_str}")

                logical_value = reverse_bits_32(physical_value)
                bin_str = ' '.join(f"{logical_value:032b}"[i:i + 4] for i in range(0, 32, 4))
                if enable_log:
                    log(f"  <- 读回逻辑值二进制[31:0]: {bin_str}")
                    log(f"  <- 读回逻辑值 HEX: 0x{logical_value:08X}")
                return logical_value

            read_one_frame(enable_log=False)
            second_value = read_one_frame(enable_log=True)

            if second_value is None:
                log("  !! 读操作无有效返回")
                return None

            return second_value
    except Exception as e:
        log(f"  !! 串口错误: {e}")
        raise


def write_then_read(addr, index, write_value, port="COM15", baudrate=921600, log_callback=None):
    def log(msg):
        if log_callback:
            log_callback(msg)
        else:
            print(msg)
    log("=== 第1帧：写命令 (0b010) ===")
    send_d9_spi('write', addr, index, write_value, port, baudrate, log_callback)
    time.sleep(WRITE_THEN_READ_GAP)
    log("=== 第2帧：读命令 (0b110) ===")
    return send_d9_spi('read', addr, index, port=port, baudrate=baudrate, log_callback=log_callback)


# ========== GUI 界面 ==========
class D9SPIApp:
    def __init__(self, root):
        self.root = root
        root.title("2361寄存器调试工具")
        root.geometry("980x620")
        root.resizable(True, True)

        self.config = load_config()
        self.log_records = []  # 结构化日志
        self.create_widgets()

    def create_widgets(self):
        # 串口
        ttk.Label(self.root, text="串口号:").grid(row=0, column=0, sticky='w', padx=10, pady=5)
        self.port_var = tk.StringVar(value=self.config["port"])
        ttk.Entry(self.root, textvariable=self.port_var, width=15).grid(row=0, column=1, sticky='w', padx=5)

        # 地址 (5-bit)
        ttk.Label(self.root, text="5位地址 (0-31):").grid(row=1, column=0, sticky='w', padx=10, pady=5)
        self.addr_var = tk.StringVar(value=self.config["addr"])
        ttk.Entry(self.root, textvariable=self.addr_var, width=15).grid(row=1, column=1, sticky='w', padx=5)

        # 索引 (6-bit)
        ttk.Label(self.root, text="6位索引 (0-63):").grid(row=2, column=0, sticky='w', padx=10, pady=5)
        self.index_var = tk.StringVar(value=self.config["index"])
        ttk.Entry(self.root, textvariable=self.index_var, width=15).grid(row=2, column=1, sticky='w', padx=5)

        # 写入值
        ttk.Label(self.root, text="写入值 (HEX, 8位):").grid(row=3, column=0, sticky='w', padx=10, pady=5)
        self.write_val_var = tk.StringVar(value="AABBCCDD")
        ttk.Entry(self.root, textvariable=self.write_val_var, width=15).grid(row=3, column=1, sticky='w', padx=5)

        # 读导出地址
        ttk.Label(self.root, text="导出地址(0-31):").grid(row=4, column=0, sticky='w', padx=10, pady=5)
        self.export_addr_var = tk.StringVar(value=self.config["addr"])
        ttk.Entry(self.root, textvariable=self.export_addr_var, width=15).grid(row=4, column=1, sticky='w', padx=5)

        # 读导出索引范围
        ttk.Label(self.root, text="导出索引范围(0-63):").grid(row=5, column=0, sticky='w', padx=10, pady=5)
        self.export_index_range_var = tk.StringVar(value=f"{self.config['index']}-{self.config['index']}")
        ttk.Entry(self.root, textvariable=self.export_index_range_var, width=15).grid(row=5, column=1, sticky='w', padx=5)

        # 按钮
        btn_frame = ttk.Frame(self.root)
        btn_frame.grid(row=6, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="单独读", command=self.do_read, width=12).pack(side='left', padx=5)
        ttk.Button(btn_frame, text="单独写", command=self.do_write, width=12).pack(side='left', padx=5)
        ttk.Button(btn_frame, text="写后读", command=self.do_write_then_read, width=12).pack(side='left', padx=5)
        ttk.Button(btn_frame, text="导出日志", command=self.export_to_excel, width=12).pack(side='left', padx=5)
        ttk.Button(btn_frame, text="生成模板", command=self.generate_template, width=12).pack(side='left', padx=5)
        ttk.Button(btn_frame, text="批量执行", command=self.import_and_execute_batch, width=12).pack(side='left',
                                                                                                     padx=5)
        ttk.Button(btn_frame, text="读数据并导出", command=self.read_data_and_export_template, width=12).pack(side='left', padx=5)
        # 日志输出（列宽调整）
        ttk.Label(self.root, text="操作日志:").grid(row=7, column=0, sticky='w', padx=10, pady=(10, 0))
        self.log_text = scrolledtext.ScrolledText(self.root, width=75, height=18, state='disabled')
        self.log_text.grid(row=8, column=0, columnspan=3, padx=10, pady=5)

        # ========== 新增：生成Excel模板 ==========

    def generate_template(self):
        filename = filedialog.asksaveasfilename(
            title="保存批量操作模板",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            initialfile=f"批量操作模板_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if not filename:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "批量操作模板"

            # 表头
            headers = ["READ/WRITE/WRITETHENREAD", "ADDR", "INDEX", "[31:24]", "[23:16]", "[15:8]", "[7:0]"]
            ws.append(headers)

            # 说明行（灰色提示）
            note_row = [
                "填写操作类型（READ / WRITE / WRITETHENREAD）",
                "0~31",
                "0~63",
                "HEX (如 AA)",
                "HEX (如 BB)",
                "HEX (如 CC)",
                "HEX (如 DD)"
            ]
            ws.append(note_row)

            # 示例行
            example_row = ["WRITE", 1, 5, "AA", "BB", "CC", "DD"]
            ws.append(example_row)
            ws.append(["WRITETHENREAD", 2, 10, "12", "34", "56", "78"])
            ws.append(["READ", 3, 15, "", "", "", ""])

            # 样式
            header_font = Font(bold=True)
            for col in range(1, len(headers) + 1):
                ws.cell(row=1, column=col).font = header_font

            # 自动列宽
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 25)
                ws.column_dimensions[column].width = adjusted_width

            wb.save(filename)
            messagebox.showinfo("成功", f"模板已生成：\n{filename}")
        except Exception as e:
            messagebox.showerror("错误", f"生成模板失败：\n{str(e)}")

    def parse_hex_byte(self, s, col_name):
        if s is None or str(s).strip() == "":
            return 0
        s = str(s).strip().upper()
        if re.fullmatch(r'[0-9A-F]{1,2}', s):
            val = int(s, 16)
            if 0 <= val <= 255:
                return val
            else:
                raise ValueError(f"{col_name} 值 {s} 超出范围 (0~255)")
        else:
            raise ValueError(f"{col_name} 必须是 1-2 位 HEX（如 'A3' 或 '5'）")

    def parse_number_list_or_range(self, text, min_val, max_val, field_name):
        if text is None:
            raise ValueError(f"{field_name} 不能为空")

        raw = str(text).strip()
        if not raw:
            raise ValueError(f"{field_name} 不能为空")

        values = []
        for part in raw.replace('，', ',').split(','):
            item = part.strip()
            if not item:
                continue

            if '-' in item:
                range_parts = item.split('-', 1)
                if len(range_parts) != 2:
                    raise ValueError(f"{field_name} 范围格式错误: {item}")
                start = int(range_parts[0].strip())
                end = int(range_parts[1].strip())
                if start > end:
                    raise ValueError(f"{field_name} 范围起始不能大于结束: {item}")
                for v in range(start, end + 1):
                    if not (min_val <= v <= max_val):
                        raise ValueError(f"{field_name} 超出范围({min_val}~{max_val}): {v}")
                    values.append(v)
            else:
                v = int(item)
                if not (min_val <= v <= max_val):
                    raise ValueError(f"{field_name} 超出范围({min_val}~{max_val}): {v}")
                values.append(v)

        if not values:
            raise ValueError(f"{field_name} 未解析到有效值")

        return values

    def read_data_and_export_template(self):
        port = self.port_var.get().strip()
        if not port:
            messagebox.showerror("错误", "请先设置串口号！")
            return

        try:
            addr_values = self.parse_number_list_or_range(self.export_addr_var.get(), 0, 31, "ADDR")
            index_values = self.parse_number_list_or_range(self.export_index_range_var.get(), 0, 63, "INDEX")
        except Exception as e:
            messagebox.showerror("参数错误", str(e))
            return

        filename = filedialog.asksaveasfilename(
            title="保存读数据导出文件",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            initialfile=f"读数据导出_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if not filename:
            return

        headers = ["READ/WRITE/WRITETHENREAD", "ADDR", "INDEX", "[31:24]", "[23:16]", "[15:8]", "[7:0]"]

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "批量操作模板"
            ws.append(headers)

            header_font = Font(bold=True)
            for col in range(1, len(headers) + 1):
                ws.cell(row=1, column=col).font = header_font

            total_count = 0
            success_count = 0

            for addr in addr_values:
                for index in index_values:
                    total_count += 1
                    self.log(f"导出读取: addr={addr}, index={index}")
                    try:
                        read_val = send_d9_spi('read', addr, index, port=port, log_callback=self.log)
                        if read_val is None:
                            ws.append(["READ", addr, index, "", "", "", ""])
                            continue

                        b31 = (read_val >> 24) & 0xFF
                        b23 = (read_val >> 16) & 0xFF
                        b15 = (read_val >> 8) & 0xFF
                        b7 = read_val & 0xFF

                        ws.append(["READ", addr, index, f"{b31:02X}", f"{b23:02X}", f"{b15:02X}", f"{b7:02X}"])
                        success_count += 1

                        self.log_records.append({
                            'time': time.strftime('%Y-%m-%d %H:%M:%S'),
                            'operation': 'READ_EXPORT',
                            'addr': addr,
                            'index': index,
                            'write_value': '',
                            'read_value': f"0x{read_val:08X}",
                            'raw_message': ''
                        })
                        time.sleep(0.02)
                    except Exception as e:
                        self.log(f"!! 读取失败: addr={addr}, index={index}, err={e}")
                        ws.append(["READ", addr, index, "", "", "", ""])

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        cell_len = len(str(cell.value)) if cell.value is not None else 0
                        if cell_len > max_length:
                            max_length = cell_len
                    except Exception:
                        pass
                ws.column_dimensions[column].width = min(max_length + 2, 25)

            wb.save(filename)
            self.log(f"读数据并导出完成: 总计 {total_count} 条, 成功 {success_count} 条")
            messagebox.showinfo("完成", f"导出完成:\n{filename}\n总计 {total_count} 条，成功 {success_count} 条")
        except Exception as e:
            messagebox.showerror("导出失败", f"读数据并导出失败：\n{str(e)}")
            self.log(f"!! 读数据并导出异常: {e}")

    def import_and_execute_batch(self):
        filename = filedialog.askopenfilename(
            title="选择批量操作 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx")]
        )
        if not filename:
            return

        port = self.port_var.get().strip()
        if not port:
            messagebox.showerror("错误", "请先设置串口号！")
            return

        try:
            wb = load_workbook(filename, data_only=True)
            ws = wb.active

            rows = list(ws.iter_rows(min_row=2, values_only=True))  # 跳过表头
            if not rows:
                messagebox.showwarning("警告", "Excel 中无有效数据行（从第2行开始）")
                return

            success_count = 0
            total_count = 0

            for idx, row in enumerate(rows, start=2):
                # 跳过空行
                if all(cell is None or str(cell).strip() == "" for cell in row[:3]):
                    continue

                total_count += 1
                try:
                    op_str = str(row[0]).strip().upper() if row[0] else ""
                    if op_str not in ["READ", "WRITE", "WRITETHENREAD"]:
                        raise ValueError(f"操作类型必须是 READ / WRITE / WRITETHENREAD，当前值: '{op_str}'")

                    addr = int(row[1]) if row[1] is not None else 0
                    index = int(row[2]) if row[2] is not None else 0

                    if not (0 <= addr <= 31):
                        raise ValueError(f"ADDR 必须在 0~31，当前值: {addr}")
                    if not (0 <= index <= 63):
                        raise ValueError(f"INDEX 必须在 0~63，当前值: {index}")

                    write_val = 0
                    if op_str in ["WRITE", "WRITETHENREAD"]:
                        b31 = self.parse_hex_byte(row[3], "[31:24]")
                        b23 = self.parse_hex_byte(row[4], "[23:16]")
                        b15 = self.parse_hex_byte(row[5], "[15:8]")
                        b7 = self.parse_hex_byte(row[6], "[7:0]")
                        write_val = (b31 << 24) | (b23 << 16) | (b15 << 8) | b7

                    # 执行操作
                    self.log(f"--- 批量执行第 {total_count} 条 (Excel 第 {idx} 行) ---")
                    if op_str == "READ":
                        read_val = send_d9_spi('read', addr, index, port=port, log_callback=self.log)
                        if read_val is not None:
                            self.log_records.append({
                                'time': time.strftime('%Y-%m-%d %H:%M:%S'),
                                'operation': 'READ',
                                'addr': addr,
                                'index': index,
                                'write_value': '',
                                'read_value': f"0x{read_val:08X}",
                                'raw_message': ''
                            })
                            success_count += 1
                    elif op_str == "WRITE":
                        send_d9_spi('write', addr, index, write_value=write_val, port=port, log_callback=self.log)
                        self.log_records.append({
                            'time': time.strftime('%Y-%m-%d %H:%M:%S'),
                            'operation': 'WRITE',
                            'addr': addr,
                            'index': index,
                            'write_value': f"0x{write_val:08X}",
                            'read_value': '',
                            'raw_message': ''
                        })
                        success_count += 1
                    elif op_str == "WRITETHENREAD":
                        read_val = write_then_read(addr, index, write_val, port=port, log_callback=self.log)
                        if read_val is not None:
                            self.log_records.append({
                                'time': time.strftime('%Y-%m-%d %H:%M:%S'),
                                'operation': 'WRITETHENREAD',
                                'addr': addr,
                                'index': index,
                                'write_value': f"0x{write_val:08X}",
                                'read_value': f"0x{read_val:08X}",
                                'raw_message': ''
                            })
                            success_count += 1

                    time.sleep(0.05)  # 避免太快

                except Exception as e:
                    error_msg = f"第 {idx} 行解析/执行失败: {str(e)}"
                    self.log(f"!! {error_msg}")
                    messagebox.showwarning("行执行失败", error_msg)
                    continue

            self.log(f"=== 批量执行完成：共 {total_count} 条，成功 {success_count} 条 ===")
            messagebox.showinfo("批量执行完成", f"共处理 {total_count} 条命令\n成功 {success_count} 条\n详情见日志")

        except Exception as e:
            messagebox.showerror("批量执行错误", f"执行失败：\n{str(e)}")
            self.log(f"!! 批量执行异常: {e}")
    def log(self, msg, record=None):
        timestamp = time.strftime('%Y-%m-%d %H:%M:%S')
        full_msg = f"[{timestamp}] {msg}"
        self.log_text.configure(state='normal')
        self.log_text.insert(tk.END, full_msg + "\n")
        self.log_text.configure(state='disabled')
        self.log_text.yview(tk.END)

        # 记录结构化日志
        if record is None:
            self.log_records.append({
                'time': timestamp,
                'operation': '',
                'addr': '',
                'index': '',
                'write_value': '',
                'read_value': '',
                'raw_message': msg
            })
        else:
            self.log_records.append(record)

    def get_inputs(self):
        try:
            port = self.port_var.get().strip()
            if not port:
                raise ValueError("串口号不能为空")
            addr = int(self.addr_var.get())
            index = int(self.index_var.get())
            write_str = self.write_val_var.get().strip().replace("0x", "").replace(" ", "")
            write_val = int(write_str, 16) if write_str else 0
            return port, addr, index, write_val
        except ValueError as e:
            messagebox.showerror("输入错误", f"参数格式不正确:\n{e}")
            raise

    def do_read(self):
        try:
            port, addr, index, _ = self.get_inputs()
            save_config(port, addr, index)
            self.log(f"开始执行：单独读")
            read_val = send_d9_spi('read', addr, index, port=port, log_callback=self.log)
            if read_val is not None:
                self.log(
                    f"READ  | addr={addr}, index={index} | value=0x{read_val:08X}",
                    record={
                        'time': time.strftime('%Y-%m-%d %H:%M:%S'),
                        'operation': 'READ',
                        'addr': addr,
                        'index': index,
                        'write_value': '',
                        'read_value': f"0x{read_val:08X}",
                        'raw_message': ''
                    }
                )
        except Exception as e:
            messagebox.showerror("错误", str(e))

    def do_write(self):
        try:
            port, addr, index, write_val = self.get_inputs()
            save_config(port, addr, index)
            self.log(f"开始执行：单独写")
            send_d9_spi('write', addr, index, write_value=write_val, port=port, log_callback=self.log)
            self.log(
                f"WRITE | addr={addr}, index={index} | value=0x{write_val:08X}",
                record={
                    'time': time.strftime('%Y-%m-%d %H:%M:%S'),
                    'operation': 'WRITE',
                    'addr': addr,
                    'index': index,
                    'write_value': f"0x{write_val:08X}",
                    'read_value': '',
                    'raw_message': ''
                }
            )
        except Exception as e:
            messagebox.showerror("错误", str(e))

    def do_write_then_read(self):
        try:
            port, addr, index, write_val = self.get_inputs()
            save_config(port, addr, index)
            self.log(f"开始执行：写后读")
            read_val = write_then_read(addr, index, write_val, port=port, log_callback=self.log)
            if read_val is not None:
                self.log(
                    f"WRITE_THEN_READ | addr={addr}, index={index} | written=0x{write_val:08X}, read=0x{read_val:08X}",
                    record={
                        'time': time.strftime('%Y-%m-%d %H:%M:%S'),
                        'operation': 'WRITE_THEN_READ',
                        'addr': addr,
                        'index': index,
                        'write_value': f"0x{write_val:08X}",
                        'read_value': f"0x{read_val:08X}",
                        'raw_message': ''
                    }
                )
        except Exception as e:
            messagebox.showerror("错误", str(e))

    def export_to_excel(self):
        if not self.log_records:
            messagebox.showinfo("提示", "当前无日志记录可导出")
            return

        filename = filedialog.asksaveasfilename(
            title="保存日志",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            initialfile=f"SPI日志_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if not filename:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "SPI操作日志"

            headers = ['时间', '操作类型', '地址(addr)', '索引(index)', '写入值', '读取值', '原始消息']
            ws.append(headers)

            header_font = Font(bold=True)
            for col in range(1, len(headers) + 1):
                ws.cell(row=1, column=col).font = header_font

            for record in self.log_records:
                ws.append([
                    record['time'],
                    record['operation'],
                    record['addr'],
                    record['index'],
                    record['write_value'],
                    record['read_value'],
                    record['raw_message']
                ])

            # 自动列宽
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column].width = adjusted_width

            wb.save(filename)
            messagebox.showinfo("成功", f"日志已导出至：\n{filename}")
        except Exception as e:
            messagebox.showerror("导出失败", f"错误：{str(e)}")


# ========== 启动程序 ==========
if __name__ == "__main__":
    root = tk.Tk()
    app = D9SPIApp(root)
    root.mainloop()