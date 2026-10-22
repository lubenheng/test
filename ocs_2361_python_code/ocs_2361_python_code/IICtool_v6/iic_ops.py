import re
import os
import time
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def _reverse_bits(x, width):
    r = 0
    for _ in range(width):
        r = (r << 1) | (x & 1)
        x >>= 1
    return r


def _reverse_bits_32(x):
    x = ((x & 0x55555555) << 1) | ((x & 0xAAAAAAAA) >> 1)
    x = ((x & 0x33333333) << 2) | ((x & 0xCCCCCCCC) >> 2)
    x = ((x & 0x0F0F0F0F) << 4) | ((x & 0xF0F0F0F0) >> 4)
    x = ((x & 0x00FF00FF) << 8) | ((x & 0xFF00FF00) >> 8)
    x = ((x & 0x0000FFFF) << 16) | ((x & 0xFFFF0000) >> 16)
    return x


def build_iic_payload(rw, addr, index, write_data=None):
    """构建 2361 协议报文。

    对 ADDR 做 5 位 bit-reversal、INDEX 做 6 位 bit-reversal 编码报头；
    写数据做 32 位 bit-reversal。

    返回: list[int] — 报文字节列表（不含 0x25/0xA5 前缀）
    """
    if rw == 'read':
        rwcode = 0b110
    elif rw == 'write':
        rwcode = 0b010
    else:
        raise ValueError("rw must be 'read' or 'write'")

    addr_rev = _reverse_bits(addr, 5)
    index_rev = _reverse_bits(index, 6)
    header_logical = (rwcode << 13) | (addr_rev << 8) | (index_rev << 2)
    high_byte = (header_logical >> 8) & 0xFF
    low_byte = header_logical & 0xFF
    res = [high_byte, low_byte]
    if rw == 'write' and write_data:
        write_value = (write_data[0] << 24) | (write_data[1] << 16) | (write_data[2] << 8) | write_data[3]
        physical_value = _reverse_bits_32(write_value)
        data_bytes = [
            (physical_value >> 24) & 0xFF,
            (physical_value >> 16) & 0xFF,
            (physical_value >> 8) & 0xFF,
            physical_value & 0xFF
        ]
        res.extend(data_bytes)
    return res


def execute_write(i2c_ctrl, addr, index, write_data, log_cb=None, suppress_log=False):
    """通过 I2CController 执行 2361 协议写操作。

    返回: None
    """
    if not (0 <= addr <= 31):
        raise ValueError("ADDR 必须在 0~31 之间")
    if not (0 <= index <= 63):
        raise ValueError("INDEX 必须在 0~63 之间")
    if len(write_data) != 4:
        raise ValueError("写入数据必须为4字节")
    tx = [0x25] + build_iic_payload('write', addr, index, write_data)
    i2c_ctrl.send_custom_write_frame(tx, log_callback=None)
    if not suppress_log and log_cb:
        orig_data_hex = ' '.join(f'{b:02X}' for b in write_data)
        log_cb(
            f"WRITE | addr={addr}, index={index} | data={orig_data_hex}",
            record={
                'time': time.strftime('%Y-%m-%d %H:%M:%S'),
                'operation': 'WRITE',
                'addr': str(addr),
                'index': str(index),
                'write_value': orig_data_hex,
                'read_value': '',
                'raw_message': 'custom-write'
            }
        )


def execute_read(i2c_ctrl, addr, index, log_cb=None):
    """通过 I2CController 执行 2361 协议读操作。

    返回: list[int]（4 字节逻辑值）或 None
    """
    if not (0 <= addr <= 31):
        raise ValueError("ADDR 必须在 0~31 之间")
    if not (0 <= index <= 63):
        raise ValueError("INDEX 必须在 0~63 之间")

    cmd_tail = build_iic_payload('read', addr, index)
    tx1 = [0x25, cmd_tail[0], cmd_tail[1]]
    tx2 = [0xA5]
    rx_len = 4
    raw = i2c_ctrl.read_stop_then_start(tx1, tx2, rx_len, log_callback=None)
    if raw is None or len(raw) < 4:
        if log_cb:
            log_cb("READ | 无有效返回")
        return None
    physical_value = (raw[0] << 24) | (raw[1] << 16) | (raw[2] << 8) | raw[3]
    logical_value = _reverse_bits_32(physical_value)
    logical_bytes = [
        (logical_value >> 24) & 0xFF,
        (logical_value >> 16) & 0xFF,
        (logical_value >> 8) & 0xFF,
        logical_value & 0xFF
    ]
    data_hex = ' '.join(f'{b:02X}' for b in logical_bytes)
    if log_cb:
        log_cb(
            f"READ | addr={addr}, index={index} | rx={data_hex}",
            record={
                'time': time.strftime('%Y-%m-%d %H:%M:%S'),
                'operation': 'READ',
                'addr': str(addr),
                'index': str(index),
                'write_value': '',
                'read_value': data_hex,
                'raw_message': f"rx_len={rx_len}"
            }
        )
    return logical_bytes


class IICLogic:
    """
    负责普通IIC寄存器操作的业务与逻辑：
    字符串解析、16进制数据验证、Excel批量导入导出逻辑等。
    """
    def __init__(self, write_i2c_func, read_i2c_func, log_callback):
        self._write_i2c = write_i2c_func
        self._read_i2c = read_i2c_func
        self.log = log_callback

    # ----------- 数据解析辅助 -----------
    @staticmethod
    def parse_hex_bytes(s):
        text = str(s or "").strip()
        if not text:
            raise ValueError("原始TX不能为空")
        parts = re.split(r'[\s,]+', text)
        data = []
        for p in parts:
            if not p:
                continue
            p = p.strip().replace("0x", "").replace("0X", "")
            if not re.fullmatch(r'[0-9A-Fa-f]+', p):
                raise ValueError(f"非法HEX字节: {p}")
            if len(p) % 2 != 0:
                raise ValueError(f"HEX长度必须为偶数: {p}")
            for i in range(0, len(p), 2):
                data.append(int(p[i:i + 2], 16))
        if not data:
            raise ValueError("原始TX不能为空")
        return data

    @staticmethod
    def parse_hex_exact4(s):
        text = str(s or "").strip().replace("0x", "").replace("0X", "")
        compact = re.sub(r'[\s,]+', '', text)
        if re.fullmatch(r'[0-9A-Fa-f]{8}', compact):
            return [int(compact[i:i + 2], 16) for i in range(0, 8, 2)]
        data = IICLogic.parse_hex_bytes(s)
        if len(data) > 4:
            raise ValueError("最多只能输入4字节数据")
        elif len(data) < 4:
            data = [0] * (4 - len(data)) + data
        return data

    @staticmethod
    def parse_hex_byte(s, col_name):
        if s is None or str(s).strip() == "":
            return 0
        s = str(s).strip().upper()
        if re.fullmatch(r'[0-9A-F]{1,2}', s):
            val = int(s, 16)
            if 0 <= val <= 255:
                return val
            else:
                raise ValueError(f"{col_name} 值 {s} 超出范围 (0~255)")
        else:
            raise ValueError(f"{col_name} 必须是 1-2 位 HEX（如 'A3' 或 '5'）")

    @staticmethod
    def parse_number_list_or_range(text, min_val, max_val, field_name):
        if text is None:
            raise ValueError(f"{field_name} 不能为空")
        raw = str(text).strip()
        if not raw:
            raise ValueError(f"{field_name} 不能为空")
        values = []
        for part in raw.replace('，', ',').split(','):
            item = part.strip()
            if not item:
                continue
            if '-' in item:
                range_parts = item.split('-', 1)
                if len(range_parts) != 2:
                    raise ValueError(f"{field_name} 范围格式错误: {item}")
                start = int(range_parts[0].strip())
                end = int(range_parts[1].strip())
                if start > end:
                    raise ValueError(f"{field_name} 范围起始不能大于结束: {item}")
                for v in range(start, end + 1):
                    if not (min_val <= v <= max_val):
                        raise ValueError(f"{field_name} 超出范围({min_val}~{max_val}): {v}")
                    values.append(v)
            else:
                v = int(item)
                if not (min_val <= v <= max_val):
                    raise ValueError(f"{field_name} 超出范围({min_val}~{max_val}): {v}")
                values.append(v)
        if not values:
            raise ValueError(f"{field_name} 未解析到有效值")
        return values

    # ----------- 表格生成与常规批量业务逻辑 -----------
    @staticmethod
    def generate_template(filename):
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "批量操作模板"
            headers = ["READ/WRITE/WRITETHENREAD", "ADDR", "INDEX", "[31:24]", "[23:16]", "[15:8]", "[7:0]"]
            ws.append(headers)
            ws.append(["START", "", "", "", "", "", ""])
            ws.append(["# 示例：在 START 与 END 之间填写操作行，注释行以 # 开头将被忽略", "", "", "", "", "", ""])
            ws.append(["WRITE", 1, 5, "AA", "BB", "CC", "DD"])
            ws.append(["WRITETHENREAD", 2, 10, "12", "34", "56", "78"])
            ws.append(["READ", 3, 15, "", "", "", ""])
            ws.append(["END", "", "", "", "", "", ""])
            header_font = Font(bold=True)
            for col in range(1, len(headers) + 1):
                ws.cell(row=1, column=col).font = header_font
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                ws.column_dimensions[column].width = min(max_length + 2, 25)
            wb.save(filename)
            return True, None
        except Exception as e:
            return False, str(e)

    @staticmethod
    def read_data_and_export_template(export_addr_var, export_index_range_var, filename, execute_read_cb, log_cb):
        try:
            addr_values = IICLogic.parse_number_list_or_range(export_addr_var, 0, 31, "ADDR")
            index_values = IICLogic.parse_number_list_or_range(export_index_range_var, 0, 63, "INDEX")
        except Exception as e:
            return False, 0, 0, f"参数错误: {str(e)}"

        headers = ["READ/WRITE/WRITETHENREAD", "ADDR", "INDEX", "[31:24]", "[23:16]", "[15:8]", "[7:0]"]
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "批量操作模板"
            ws.append(headers)
            header_font = Font(bold=True)
            for col in range(1, len(headers) + 1):
                ws.cell(row=1, column=col).font = header_font
            total_count = 0
            success_count = 0
            for addr in addr_values:
                for index in index_values:
                    total_count += 1
                    log_cb(f"导出读取: addr={addr}, index={index}")
                    try:
                        raw = execute_read_cb(addr, index)
                        if not raw or len(raw) < 4:
                            ws.append(["READ", addr, index, "", "", "", ""])
                            continue
                        b31, b23, b15, b7 = raw[0], raw[1], raw[2], raw[3]
                        ws.append(["READ", addr, index, f"{b31:02X}", f"{b23:02X}", f"{b15:02X}", f"{b7:02X}"])
                        success_count += 1
                    except Exception as e:
                        log_cb(f"!! 读取失败: addr={addr}, index={index}, err={e}")
                        ws.append(["READ", addr, index, "", "", "", ""])
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        cell_len = len(str(cell.value)) if cell.value is not None else 0
                        if cell_len > max_length:
                            max_length = cell_len
                    except Exception:
                        pass
                ws.column_dimensions[column].width = min(max_length + 2, 25)
            wb.save(filename)
            return True, total_count, success_count, None
        except Exception as e:
            return False, 0, 0, str(e)

    @staticmethod
    def import_and_execute_batch(filename, execute_write_cb, execute_read_cb, log_cb):
        """处理 Excel 批量导入并使用回调与硬件通讯"""
        try:
            wb = load_workbook(filename, data_only=True)
            if len(wb.sheetnames) > 1 and "批量操作模板" in wb.sheetnames:
                ws = wb["批量操作模板"]
            else:
                ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            if not rows:
                return False, 0, 0, "Excel 中无有效数据行"

            start_row, end_row = 2, len(rows) + 1
            for i, row in enumerate(rows, start=2):
                first = str(row[0]).strip() if row[0] is not None else ""
                if first.upper() == 'START':
                    start_row = i + 1
                    break
            for i in range(start_row, len(rows) + 2):
                row = rows[i-2]
                first = str(row[0]).strip() if row[0] is not None else ""
                if first.upper() == 'END':
                    end_row = i - 1
                    break

            write_ops = []
            read_ops = []
            verify_ops = []

            for i in range(start_row, end_row + 1):
                row = rows[i-2]
                first = str(row[0]).strip() if row[0] is not None else ""
                if first == "" or first.startswith('#') or first.startswith('//') or first.startswith('注'):
                    continue
                try:
                    op_str = first.upper()
                    if op_str not in ["READ", "WRITE", "WRITETHENREAD"]:
                        raise ValueError(f"操作类型必须是 READ / WRITE / WRITETHENREAD，当前值: '{op_str}'")
                    def _parse_decimal(val):
                        if val is None: return 0
                        if isinstance(val, (int, float)): return int(val)
                        s = str(val).strip()
                        if s == '': return 0
                        return int(s, 10)
                    addr = _parse_decimal(row[1])
                    index = _parse_decimal(row[2])
                    if not (0 <= addr <= 31):
                        raise ValueError(f"ADDR 必须在 0~31，当前值: {addr}")
                    if not (0 <= index <= 63):
                        raise ValueError(f"INDEX 必须在 0~63，当前值: {index}")

                    op_data = {'addr': addr, 'index': index, 'row_idx': i}

                    if op_str in ["WRITE", "WRITETHENREAD"]:
                        b31 = IICLogic.parse_hex_byte(row[3], "[31:24]")
                        b23 = IICLogic.parse_hex_byte(row[4], "[23:16]")
                        b15 = IICLogic.parse_hex_byte(row[5], "[15:8]")
                        b7 = IICLogic.parse_hex_byte(row[6], "[7:0]")
                        op_data['data'] = [b31, b23, b15, b7]

                    if op_str == "WRITE":
                        write_ops.append(op_data)
                    elif op_str == "READ":
                        read_ops.append(op_data)
                    elif op_str == "WRITETHENREAD":
                        write_ops.append(op_data)
                        verify_ops.append(op_data)

                except Exception as e:
                    return False, 0, 0, f"第 {i} 行解析失败: {str(e)}"

            total_count = len(write_ops) + len(read_ops) + len(verify_ops)
            success_count = 0
            verify_errors = []

            log_cb(f"--- 开始执行 {len(write_ops)} 条写入操作 ---")
            for op in write_ops:
                try:
                    execute_write_cb(op['addr'], op['index'], op['data'])
                    success_count += 1
                except Exception as e:
                    log_cb(f"!! 写入失败 (Excel行 {op['row_idx']}): {e}")

            log_cb(f"--- 开始执行 {len(read_ops)} 条读取操作 ---")
            for op in read_ops:
                try:
                    execute_read_cb(op['addr'], op['index'])
                    success_count += 1
                except Exception as e:
                    log_cb(f"!! 读取失败 (Excel行 {op['row_idx']}): {e}")

            log_cb(f"--- 开始执行 {len(verify_ops)} 条验证操作 ---")
            for op in verify_ops:
                try:
                    read_val = execute_read_cb(op['addr'], op['index'])
                    written_data = op['data']
                    matched = read_val == written_data
                    if matched:
                        success_count += 1
                    else:
                        exp_str = " ".join([f"{x:02X}" for x in written_data])
                        act_str = " ".join([f"{x:02X}" for x in read_val]) if read_val else "N/A"
                        emsg = f"第 {op['row_idx']} 行(ADDR={op['addr']}, INDEX={op['index']})校验失败! 预期:{exp_str}, 实际:{act_str}"
                        verify_errors.append(emsg)
                        log_cb(f"!! {emsg}")
                except Exception as e:
                    emsg = f"第 {op['row_idx']} 行(ADDR={op['addr']}, INDEX={op['index']})验证读取失败: {e}"
                    verify_errors.append(emsg)
                    log_cb(f"!! {emsg}")

            if verify_errors:
                error_summary = "存在校验不通过项：\n" + "\n".join(verify_errors)
                return False, total_count, success_count, error_summary

            return True, total_count, success_count, None
        except Exception as e:
            return False, 0, 0, str(e)

    @staticmethod
    def import_and_update_reads(filename, execute_write_cb, execute_read_cb, log_cb, progress_cb=None):
        """导入 Excel 后逐行读回寄存器值并更新到工作簿。

        返回: (success, total_count, success_count, workbook_or_None, error_msg)
        """
        try:
            wb = load_workbook(filename, data_only=True)
            if len(wb.sheetnames) > 1 and "批量操作模板" in wb.sheetnames:
                ws = wb["批量操作模板"]
            else:
                ws = wb.active
            rows_data = list(ws.iter_rows(min_row=2, values_only=True))
            if not rows_data:
                return True, 0, 0, wb, "Excel 中无有效数据行"

            wb_out = load_workbook(filename)
            if len(wb_out.sheetnames) > 1 and "批量操作模板" in wb_out.sheetnames:
                ws_out = wb_out["批量操作模板"]
            else:
                ws_out = wb_out.active

            start_row, end_row = 2, len(rows_data) + 1
            for i, row in enumerate(rows_data, start=2):
                first = str(row[0]).strip() if row[0] is not None else ""
                if first.upper() == 'START':
                    start_row = i + 1
                    break
            for i in range(start_row, len(rows_data) + 2):
                row = rows_data[i - 2]
                first = str(row[0]).strip() if row[0] is not None else ""
                if first.upper() == 'END':
                    end_row = i - 1
                    break

            total_count = 0
            success_count = 0
            valid_indices = []
            for i in range(start_row, end_row + 1):
                row = rows_data[i - 2]
                first = str(row[0]).strip() if row[0] is not None else ""
                if first == "" or first.startswith('#') or first.startswith('//') or first.startswith('注'):
                    continue
                valid_indices.append(i)

            for count, i in enumerate(valid_indices, start=1):
                row = rows_data[i - 2]
                total_count += 1
                if progress_cb:
                    progress_cb(count, len(valid_indices))
                try:
                    def _parse_decimal(val):
                        if val is None: return 0
                        if isinstance(val, (int, float)): return int(val)
                        s = str(val).strip()
                        if s == '': return 0
                        return int(s, 10)

                    addr = _parse_decimal(row[1])
                    index = _parse_decimal(row[2])
                    if not (0 <= addr <= 31):
                        log_cb(f"!! 第 {i} 行 ADDR={addr} 超出范围，跳过")
                        continue
                    if not (0 <= index <= 63):
                        log_cb(f"!! 第 {i} 行 INDEX={index} 超出范围，跳过")
                        continue

                    read_val = execute_read_cb(addr, index)
                    if read_val and len(read_val) >= 4:
                        ws_out.cell(row=i, column=4).value = f"{read_val[0]:02X}"
                        ws_out.cell(row=i, column=5).value = f"{read_val[1]:02X}"
                        ws_out.cell(row=i, column=6).value = f"{read_val[2]:02X}"
                        ws_out.cell(row=i, column=7).value = f"{read_val[3]:02X}"
                        success_count += 1
                    else:
                        log_cb(f"!! 第 {i} 行读回无效数据 (addr={addr}, index={index})")
                except Exception as e:
                    log_cb(f"!! 第 {i} 行读回失败: {e}")

            return True, total_count, success_count, wb_out, None
        except Exception as e:
            return False, 0, 0, None, str(e)
