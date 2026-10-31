"""生成测试用的成对 docx 样本。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


def build_left(path: Path) -> None:
    doc = Document()
    doc.add_heading("项目合同", level=1)
    doc.add_paragraph("第一条 甲方应当在2024年1月1日前完成交付。")
    doc.add_paragraph("第二条 乙方负责验收，并在三日内反馈。")
    doc.add_paragraph("第三条 本合同一式两份，双方各执一份。")

    doc.add_paragraph("付款节点如下：")
    for item in ("首付款：合同签订后支付", "尾款：验收通过后支付"):
        doc.add_paragraph(item, style="List Number")

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "设备采购"
    table.cell(1, 1).text = "1000 元"

    doc.add_paragraph("第四条 争议由甲方所在地法院管辖。")
    doc.save(str(path))


def build_right(path: Path) -> None:
    doc = Document()
    doc.add_heading("项目合同", level=1)
    doc.add_paragraph("第一条 甲方应当在2025年6月30日前完成交付，并提交验收报告。")
    # 第二条整段删除
    p3 = doc.add_paragraph("第三条 本合同一式两份，双方各执一份。")
    for run in p3.runs:
        run.bold = True
    doc.add_paragraph("第五条 双方应对合同内容保密。")

    doc.add_paragraph("付款节点如下：")
    for item in ("首付款：合同签订后支付", "尾款：验收通过后支付", "质保金：满一年后支付"):
        doc.add_paragraph(item, style="List Number")

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "设备采购"
    table.cell(1, 1).text = "1200 元"

    doc.add_paragraph("第四条 争议由甲方所在地法院管辖。")
    doc.add_paragraph("第六条 本补充条款自签署之日起生效。").runs[0].font.size = Pt(12)
    doc.save(str(path))


def main(out_dir: str | Path = "tests/_tmp") -> tuple[Path, Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    left, right = out / "合同A.docx", out / "合同B.docx"
    build_left(left)
    build_right(right)
    print(f"已生成：{left} / {right}")
    return left, right


if __name__ == "__main__":
    main()
