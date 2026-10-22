# -*- coding: utf-8 -*-
import os, base64
from pathlib import Path
from docx import Document
from docx.shared import Inches

out = Path(os.environ.get("WM_TEST_DIR", r"C:\Users\user\Tools\WordMergeTool\tests\data"))
out.mkdir(parents=True, exist_ok=True)

# 1x1 red png
png = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
img = out / "pic.png"
img.write_bytes(png)

def build(name, h1, h2s, with_table=False, with_image=False):
    d = Document()
    d.add_heading(h1, level=1)
    d.add_paragraph(f"这是{h1}的正文段落。")
    for h in h2s:
        d.add_heading(h, level=2)
        d.add_paragraph(f"{h} 的内容。")
    if with_table:
        t = d.add_table(rows=2, cols=3)
        t.style = "Table Grid"
        for i, c in enumerate(["A", "B", "C"]):
            t.rows[0].cells[i].text = c
        for i, c in enumerate(["1", "2", "3"]):
            t.rows[1].cells[i].text = c
    if with_image:
        d.add_picture(str(img), width=Inches(0.5))
    path = out / name
    d.save(path)
    print("created", path)

build("第2章 系统时钟.docx", "第2章 系统时钟", ["2.1 时钟树", "2.2 寄存器"], with_table=True)
build("第1章 概述.docx", "第1章 概述", ["1.1 产品简介", "1.2 特性"], with_image=True)
build("第3章 附录.docx", "第3章 附录", ["3.1 封装尺寸"])
build("未编号的说明.docx", "修订记录", [])
