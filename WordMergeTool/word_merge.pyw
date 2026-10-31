# -*- coding: utf-8 -*-
"""Word 文档按章节合并工具

- 支持拖拽 / 添加多个 .docx/.doc/.docm/.rtf 文件
- 每份文档作为一章，按标题或文件名中的章号自动排序
- 合并后自动生成目录（Word COM 方式会立即更新页码）
- 默认只按章节拼接（docxcompose）：不修改源文件，不导入/删除样式，
  各章原有格式与样式库原样保留
- 支持把总文档按章节拆分为单独文档（格式保持不变）
- 样式模板、样式清理、样式统一均为手动勾选的可选操作，
  未安装 Word 时回退到 python-docx + docxcompose
"""

from __future__ import annotations

import os
import queue
import re
import shutil
import tempfile
import threading
import traceback
from dataclasses import dataclass
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

SUPPORTED_EXTENSIONS = {".docx", ".doc", ".docm", ".rtf"}
STYLE_SOURCE_EXTENSIONS = {".docx", ".doc", ".docm", ".dotx", ".dotm"}
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_CANCELLED = object()

try:  # 可选：支持拖拽
    from tkinterdnd2 import DND_FILES, TkinterDnD

    _BaseTk = TkinterDnD.Tk
    HAS_DND = True
except Exception:  # pragma: no cover - 依赖缺失时降级
    _BaseTk = tk.Tk
    HAS_DND = False


# --------------------------------------------------------------------------
# 章号识别
# --------------------------------------------------------------------------

_CN_DIGITS = {
    "零": 0, "〇": 0, "一": 1, "二": 2, "两": 2, "三": 3, "四": 4,
    "五": 5, "六": 6, "七": 7, "八": 8, "九": 9,
}
_CN_UNITS = {"十": 10, "百": 100, "千": 1000}

_RE_CHAPTER_CN = re.compile(
    r"第\s*([零〇一二三四五六七八九十百两\d]+(?:[.．]\d+)*)\s*[章篇部节]"
)
_RE_CHAPTER_EN = re.compile(r"^\s*(?:chapter|chap\.?|ch\.?)\s*(\d+)", re.I)
_RE_LEADING_NUM = re.compile(r"^\s*(\d+(?:[.．]\d+)*)(?=$|[\s、,，.．:：_\-—])")


def _cn_to_int(text: str) -> int | None:
    text = text.strip()
    if not text:
        return None
    if text.isdigit():
        return int(text)
    total = 0
    value = 0
    for ch in text:
        if ch in _CN_DIGITS:
            value = _CN_DIGITS[ch]
        elif ch in _CN_UNITS:
            value = (value or 1) * _CN_UNITS[ch]
            total += value
            value = 0
        else:
            return None
    return total + value


def _parse_number_token(token: str) -> tuple[int, ...] | None:
    token = token.strip().replace("．", ".")
    if not token:
        return None
    if re.fullmatch(r"\d+(?:\.\d+)*", token):
        return tuple(int(part) for part in token.split("."))
    number = _cn_to_int(token)
    return (number,) if number is not None else None


def parse_chapter_number(text: str | None) -> tuple[int, ...] | None:
    """从标题或文件名中解析章号，例如 第3章 -> (3,)，1.2 -> (1, 2)。"""
    if not text:
        return None
    match = _RE_CHAPTER_CN.search(text)
    if match:
        number = _parse_number_token(match.group(1))
        if number:
            return number
    match = _RE_CHAPTER_EN.search(text)
    if match:
        return (int(match.group(1)),)
    match = _RE_LEADING_NUM.search(text)
    if match:
        return _parse_number_token(match.group(1))
    return None


def _first_heading(path: Path) -> str | None:
    """读取 docx 的第一个标题（优先一级标题）。"""
    try:
        from docx import Document
    except Exception:
        return None
    try:
        document = Document(str(path))
    except Exception:
        return None

    fallback: str | None = None
    for index, paragraph in enumerate(document.paragraphs):
        if index > 500:
            break
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower().replace(" ", "")
        if style in {"heading1", "标题1"}:
            return text
        if fallback is None and style.startswith(("heading", "标题")):
            fallback = text
    return fallback


@dataclass
class DocumentItem:
    path: Path
    number: tuple[int, ...] | None
    heading: str

    @property
    def chapter_label(self) -> str:
        if self.number:
            return "第" + ".".join(str(n) for n in self.number) + "章"
        return "[未编号]"

    @property
    def label(self) -> str:
        text = f"{self.chapter_label}  {self.path.name}"
        heading = self.heading.strip()
        if heading and heading != self.path.stem:
            text += f"    （{heading}）"
        return text

    @property
    def sort_key(self):
        if self.number:
            # 只按章号排序；章号相同时保持原有的添加/插入顺序（稳定排序）
            return (0, self.number)
        # 未识别到章号的文件保持添加顺序
        return (1,)


def detect_document(path: Path) -> DocumentItem:
    heading = ""
    if path.suffix.lower() == ".docx":
        heading = _first_heading(path) or ""
    number = parse_chapter_number(heading)
    if number is None:
        number = parse_chapter_number(path.stem)
    return DocumentItem(path=path, number=number, heading=heading or path.stem)


# --------------------------------------------------------------------------
# 合并选项与实现
# --------------------------------------------------------------------------

@dataclass
class MergeOptions:
    make_toc: bool = True
    toc_title: str = "目录"
    page_break: bool = True
    toc_levels: int = 3
    style_source: Path | None = None
    prune_styles: bool = False
    unify_styles: bool = False
    save_source: bool = False

    def __post_init__(self):
        if not self.toc_levels:
            self.toc_levels = 3


# --------------------------------------------------------------------------
# 样式库导入与样式清理
# --------------------------------------------------------------------------

def _word_style_names(word_document) -> set[str]:
    """读取 Word 文档中全部样式的本地名称（小写）。"""
    names: set[str] = set()
    try:
        count = word_document.Styles.Count
    except Exception:
        return names
    for index in range(1, count + 1):
        try:
            name = str(word_document.Styles(index).NameLocal or "").strip()
        except Exception:
            continue
        if name:
            names.add(name.casefold())
    return names


def _import_styles_with_word(word, document, source: Path, log) -> set[str] | None:
    """用 Word 的 CopyStylesFromTemplate 把模板中的样式库导入目标文档。

    返回模板中的样式名集合；读取失败时返回 None（此时调用方应跳过样式清理，
    避免误删模板中本应保留的样式）。
    """
    names: set[str] | None = None
    source_document = None
    try:
        source_document = word.Documents.Open(
            FileName=str(source),
            ConfirmConversions=False,
            ReadOnly=True,
            AddToRecentFiles=False,
            Visible=False,
        )
        names = _word_style_names(source_document)
    except Exception as exc:
        log(f"读取样式模板失败：{exc}")
    finally:
        try:
            if source_document is not None:
                source_document.Close(SaveChanges=0)
        except Exception:
            pass

    try:
        document.CopyStylesFromTemplate(str(source))
    except Exception as exc:
        log(f"导入样式模板失败：{exc}")
        return None

    count_text = str(len(names)) if names is not None else "未知数量"
    log(f"已导入样式模板：{source.name}（{count_text} 个样式）")
    return names


def _scan_docx_style_refs(path: Path):
    """扫描 docx 包，返回（样式依赖表, 正文等处实际引用到的 styleId 集合）。"""
    import zipfile

    from lxml import etree

    ref_tags = ("pStyle", "rStyle", "tblStyle", "numStyleLink", "styleLink")
    style_deps: dict[str, tuple[str, list[str]]] = {}
    referenced: set[str] = set()

    with zipfile.ZipFile(str(path)) as archive:
        for entry in archive.namelist():
            if not entry.startswith("word/") or not entry.endswith(".xml"):
                continue
            try:
                root = etree.fromstring(archive.read(entry))
            except Exception:
                continue
            if entry == "word/styles.xml":
                for element in root.iter(f"{{{_W_NS}}}style"):
                    style_id = (
                        element.get(f"{{{_W_NS}}}styleId") or ""
                    ).strip().casefold()
                    if not style_id:
                        continue
                    name_element = element.find(f"{{{_W_NS}}}name")
                    name = (
                        (
                            name_element.get(f"{{{_W_NS}}}val")
                            if name_element is not None
                            else ""
                        )
                        or ""
                    ).strip()
                    deps = []
                    for tag in ("basedOn", "next", "link"):
                        node = element.find(f"{{{_W_NS}}}{tag}")
                        if node is not None:
                            value = (node.get(f"{{{_W_NS}}}val") or "").strip()
                            if value:
                                deps.append(value.casefold())
                    style_deps[style_id] = (name, deps)
                continue
            for tag in ref_tags:
                for element in root.iter(f"{{{_W_NS}}}{tag}"):
                    value = (element.get(f"{{{_W_NS}}}val") or "").strip()
                    if value:
                        referenced.add(value.casefold())
    return style_deps, referenced


def _style_names_from_refs(
    style_deps: dict, referenced: set[str], keep_names: set[str] | None
) -> set[str]:
    """把引用到的 styleId 展开成样式名称集合（含继承链与样式库依赖）。"""
    stack = list(referenced)
    if keep_names:
        for style_id, (name, _) in style_deps.items():
            if name and name.casefold() in keep_names:
                stack.append(style_id)

    used_names: set[str] = set()
    seen: set[str] = set()
    while stack:
        style_id = stack.pop()
        if style_id in seen:
            continue
        seen.add(style_id)
        info = style_deps.get(style_id)
        if info is None:
            continue
        name, deps = info
        if name:
            used_names.add(name.casefold())
        stack.extend(deps)
    return used_names


def _used_style_names_from_docx(path: Path, keep_names: set[str] | None) -> set[str]:
    """扫描 docx 包，返回正文/页眉页脚/编号等实际引用到的样式名称。

    除了直接引用的样式，还会沿 basedOn/next/link 继承链补全；样式库中的样式
    及其依赖也不会被判定为“多余”。
    """
    style_deps, referenced = _scan_docx_style_refs(path)
    return _style_names_from_refs(style_deps, referenced, keep_names)


def _referenced_style_ids_from_docx(path: Path) -> set[str]:
    """返回文档正文/页眉页脚/编号里真正引用到的 styleId（小写）。"""
    _, referenced = _scan_docx_style_refs(path)
    return referenced


def _prune_word_styles(
    document, package_path: Path | None, keep_names: set[str] | None, log
) -> int:
    """清理文档中的多余样式，返回删除的样式数量。

    指定样式库模板（keep_names）时，删除文档中与模板无关的样式（正文改用默认样式）；
    未指定时，只删除正文中未使用的自定义样式。

    Word 的 Style.InUse 对“已定义”的样式一律返回 True，不能用于判断引用情况，
    因此改为解析已保存文档包中的 OOXML，统计真正被引用的样式。
    """
    wd_style_type_list = 4
    template_mode = keep_names is not None

    used_names: set[str] = set()
    if not template_mode:
        try:
            used_names = _used_style_names_from_docx(package_path, keep_names)
        except Exception as exc:
            log(f"无法统计样式使用情况，已跳过样式清理：{exc}")
            return 0

    def _should_delete(style, name: str) -> bool:
        try:
            if style.BuiltIn or style.Type == wd_style_type_list:
                return False
        except Exception:
            return False
        if template_mode:
            # 与模板无关的样式一律删除（包括正文正在使用的）
            return name.casefold() not in keep_names
        return name.casefold() not in used_names

    removed = 0
    kept_in_use = 0
    # 多轮删除：子样式受 basedOn 引用保护时，先删子样式，下一轮父样式即可删除
    for _ in range(4):
        round_removed = 0
        for index in range(document.Styles.Count, 0, -1):
            try:
                style = document.Styles(index)
                name = str(style.NameLocal or "").strip()
            except Exception:
                continue
            if not name or not _should_delete(style, name):
                continue
            if not template_mode and name.casefold() in used_names:
                kept_in_use += 1
                continue
            try:
                style.Delete()
                round_removed += 1
            except Exception:
                # 被其他样式引用、被锁定等情况，Word 会拒绝删除，留待下一轮
                continue
        removed += round_removed
        if not round_removed:
            break

    if template_mode:
        message = f"样式清理：已删除 {removed} 个模板之外的样式"
    else:
        message = f"样式清理：已删除 {removed} 个未使用的多余样式"
        if kept_in_use:
            message += f"，保留 {kept_in_use} 个正文正在使用的样式"
    log(message + "。")

    if removed:
        try:
            document.Save()
        except Exception as exc:
            log(f"样式清理后的保存失败：{exc}")
    return removed


def _read_styles_xml(source: Path):
    """从 docx/dotx 等 Open XML 文件读取 styles.xml 根元素。"""
    import zipfile

    from lxml import etree

    try:
        with zipfile.ZipFile(str(source)) as archive:
            data = archive.read("word/styles.xml")
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise RuntimeError(f"无法读取样式模板 {source.name}：{exc}") from exc
    return etree.fromstring(data)


def _import_styles_python_docx(document, source: Path, log) -> set[str]:
    """无 Word 时的回退方案：把样式库文件中的样式合并进目标文档。"""
    from copy import deepcopy

    from docx.oxml.ns import qn

    source_root = _read_styles_xml(source)
    target_root = document.styles.element

    by_id: dict[str, object] = {}
    by_name: dict[str, object] = {}
    defaults: set[str] = set()
    for element in target_root.findall(qn("w:style")):
        style_id = (element.get(qn("w:styleId")) or "").strip()
        if style_id:
            by_id[style_id.casefold()] = element
        name_element = element.find(qn("w:name"))
        name = ((name_element.get(qn("w:val")) if name_element is not None else "") or "").strip()
        if name:
            by_name[name.casefold()] = element
        if element.get(qn("w:default")) in ("1", "true", "on"):
            defaults.add(element.get(qn("w:type")) or "paragraph")

    # 文档默认字体/字号等也属于样式库的一部分
    source_defaults = source_root.find(qn("w:docDefaults"))
    if source_defaults is not None:
        new_defaults = deepcopy(source_defaults)
        target_defaults = target_root.find(qn("w:docDefaults"))
        if target_defaults is not None:
            target_root.replace(target_defaults, new_defaults)
        else:
            target_root.insert(0, new_defaults)

    names: set[str] = set()
    imported = 0
    replaced = 0
    for element in source_root.findall(qn("w:style")):
        name_element = element.find(qn("w:name"))
        name = ((name_element.get(qn("w:val")) if name_element is not None else "") or "").strip()
        style_id = (element.get(qn("w:styleId")) or "").strip()
        if name:
            names.add(name.casefold())

        existing = by_id.get(style_id.casefold()) if style_id else None
        if existing is None and name:
            existing = by_name.get(name.casefold())

        new_element = deepcopy(element)
        if existing is not None:
            existing_id = (existing.get(qn("w:styleId")) or "").strip()
            # 同名不同 ID 时沿用目标文档的 ID，保证正文中的样式引用不失效
            if existing_id and style_id and existing_id != style_id:
                new_element.set(qn("w:styleId"), existing_id)
            target_root.replace(existing, new_element)
            if existing_id:
                by_id[existing_id.casefold()] = new_element
            if name:
                by_name[name.casefold()] = new_element
            replaced += 1
        else:
            style_type = new_element.get(qn("w:type")) or "paragraph"
            is_default = new_element.get(qn("w:default")) in ("1", "true", "on")
            if is_default and style_type in defaults:
                new_element.attrib.pop(qn("w:default"), None)
            target_root.append(new_element)
            if style_id:
                by_id[style_id.casefold()] = new_element
            if name:
                by_name[name.casefold()] = new_element
            if is_default:
                defaults.add(style_type)
            imported += 1

    log(
        f"已导入样式模板：{source.name}（新增 {imported} 个样式，覆盖 {replaced} 个样式）"
    )
    return names


def _prune_python_docx_styles(document, keep_names: set[str] | None, log) -> int:
    """回退方案：清理多余样式。

    指定样式库时删除与样式库无关的样式（被删除样式上的正文回落到默认样式）；
    未指定时只删除未使用的自定义样式。正文、页眉页脚、编号引用到的样式始终保留。
    """
    from docx.oxml.ns import qn

    styles_root = document.styles.element
    style_elements = list(styles_root.findall(qn("w:style")))
    protected_ids = {"normal", "defaultparagraphfont", "tablenormal", "nolist"}
    template_mode = keep_names is not None

    used_ids: set[str] = set()
    scanned_parts: set[int] = set()
    for part in document.part.package.parts:
        element = getattr(part, "element", None)
        if element is None or id(element) in scanned_parts:
            continue
        scanned_parts.add(id(element))
        for tag in ("w:pStyle", "w:rStyle", "w:tblStyle", "w:numStyleLink", "w:styleLink"):
            for node in element.iter(qn(tag)):
                value = (node.get(qn("w:val")) or "").strip()
                if value:
                    used_ids.add(value.casefold())

    def _style_info(element) -> tuple[str, str, bool]:
        style_id = (element.get(qn("w:styleId")) or "").strip()
        name_element = element.find(qn("w:name"))
        name = (
            (name_element.get(qn("w:val")) if name_element is not None else "") or ""
        ).strip()
        is_default = element.get(qn("w:default")) in ("1", "true", "on")
        return style_id, name, is_default

    candidates = []
    kept_in_use = 0
    for element in style_elements:
        style_id, name, is_default = _style_info(element)
        style_key = style_id.casefold()
        if is_default or style_key in protected_ids:
            continue
        if template_mode:
            if name.casefold() in keep_names:
                continue
        else:
            if style_key in used_ids:
                kept_in_use += 1
                continue
        # 编号定义直接引用样式时，删除会导致列表格式错乱，保留
        if style_key in used_ids and _is_numbering_reference(document, style_id):
            continue
        candidates.append(element)

    # 候选删除的样式如果被保留样式引用（basedOn/next/link），也必须保留
    candidate_markers = {id(element) for element in candidates}
    referenced_ids: set[str] = set()
    for element in style_elements:
        if id(element) in candidate_markers:
            continue
        for tag in ("w:basedOn", "w:next", "w:link"):
            node = element.find(qn(tag))
            if node is not None:
                value = (node.get(qn("w:val")) or "").strip()
                if value:
                    referenced_ids.add(value.casefold())

    removed = 0
    removed_ids: set[str] = set()
    for element in candidates:
        style_id, _, _ = _style_info(element)
        if style_id.casefold() in referenced_ids:
            continue
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
            removed += 1
            removed_ids.add(style_id.casefold())

    # 模板模式下会删除正文正在使用的样式，需要同时去掉对它们的引用，
    # 让这些段落/文字回落到默认样式（Word 中的效果等同于“删除样式”）
    if removed_ids:
        _drop_style_references(document, removed_ids)

    if template_mode:
        message = f"样式清理：已删除 {removed} 个模板之外的样式"
    else:
        message = f"样式清理：已删除 {removed} 个未使用的多余样式"
        if kept_in_use:
            message += f"，保留 {kept_in_use} 个正文正在使用的样式"
    log(message + "。")

    removed += _delete_duplicate_styles_python_docx(document, used_ids, log)
    return removed


def _delete_duplicate_styles_python_docx(document, used_ids: set[str], log) -> int:
    """删除“同名但未实际使用”的重复样式（python-docx 方式）。"""
    from docx.oxml.ns import qn

    styles_root = document.styles.element
    elements = list(styles_root.findall(qn("w:style")))

    groups: dict[tuple, list] = {}
    for element in elements:
        if element.get(qn("w:default")) in ("1", "true", "on"):
            continue
        # 只处理用户自定义样式，内置样式（同名重复）不动
        if element.get(qn("w:customStyle")) not in ("1", "true", "on"):
            continue
        style_id = (element.get(qn("w:styleId")) or "").strip()
        if not style_id:
            continue
        name_element = element.find(qn("w:name"))
        name = (
            (name_element.get(qn("w:val")) if name_element is not None else "") or ""
        ).strip()
        if not name:
            continue
        key = (element.get(qn("w:type")) or "paragraph", name.casefold())
        groups.setdefault(key, []).append((style_id, element))

    candidates: list[tuple[str, object]] = []
    for items in groups.values():
        if len(items) < 2:
            continue
        used = [item for item in items if item[0].casefold() in used_ids]
        keep = {id(element) for _sid, element in (used or items[:1])}
        for style_id, element in items:
            if id(element) in keep or style_id.casefold() in used_ids:
                continue
            candidates.append((style_id, element))

    if not candidates:
        return 0

    candidate_markers = {id(element) for _sid, element in candidates}
    referenced_ids: set[str] = set()
    for element in elements:
        if id(element) in candidate_markers:
            continue
        for tag in ("w:basedOn", "w:next", "w:link"):
            node = element.find(qn(tag))
            if node is not None:
                value = (node.get(qn("w:val")) or "").strip()
                if value:
                    referenced_ids.add(value.casefold())

    removed = 0
    for style_id, element in candidates:
        if style_id.casefold() in referenced_ids:
            continue
        if _is_numbering_reference(document, style_id):
            continue
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
            removed += 1
    if removed:
        log(f"重复样式清理：已删除 {removed} 个重复且未使用的样式。")
    return removed


def _dedupe_docx_styles(path: Path, log) -> int:
    """在已保存的 .docx 上清理重复且未使用的样式。

    Word 的对象模型拿不到 styleId，无法区分同名的两个样式；因此改为在
    Word 关闭文档后解析文档包，按正文实际引用的 styleId 判断保留哪一个。
    """
    from docx import Document

    if path.suffix.lower() != ".docx" or not path.is_file():
        return 0
    try:
        used_ids = _referenced_style_ids_from_docx(path)
    except Exception as exc:
        log(f"无法统计样式引用，跳过重复样式清理：{exc}")
        return 0
    try:
        document = Document(str(path))
    except Exception as exc:
        log(f"无法读取文档，跳过重复样式清理：{exc}")
        return 0
    removed = _delete_duplicate_styles_python_docx(document, used_ids, log)
    if removed:
        try:
            document.save(str(path))
        except Exception as exc:
            log(f"重复样式清理后的保存失败：{exc}")
            return 0
    return removed


def _drop_style_references(document, style_ids: set[str]) -> None:
    """去掉正文/页眉页脚中对已删除样式的引用，使其回落到默认样式。"""
    from docx.oxml.ns import qn

    scanned_parts: set[int] = set()
    for part in document.part.package.parts:
        element = getattr(part, "element", None)
        if element is None or id(element) in scanned_parts:
            continue
        scanned_parts.add(id(element))
        for tag in ("w:pStyle", "w:rStyle", "w:tblStyle"):
            for node in list(element.iter(qn(tag))):
                value = (node.get(qn("w:val")) or "").strip().casefold()
                if value and value in style_ids:
                    parent = node.getparent()
                    if parent is not None:
                        parent.remove(node)


def _is_numbering_reference(document, style_id: str) -> bool:
    """判断样式是否被编号定义引用（numStyleLink / pStyle 关联）。"""
    from docx.oxml.ns import qn

    try:
        numbering_root = document.part.numbering_part.element
    except Exception:
        return False
    key = style_id.casefold()
    for tag in ("w:pStyle", "w:numStyleLink", "w:styleLink"):
        for node in numbering_root.iter(qn(tag)):
            value = (node.get(qn("w:val")) or "").strip().casefold()
            if value and value == key:
                return True
    return False


# --------------------------------------------------------------------------
# 合并前：逐个文件统一到模板样式
# --------------------------------------------------------------------------

class WordUnavailableError(RuntimeError):
    """本机没有可用的 Microsoft Word。"""


def _make_backup_dir(parent: Path, prefix: str = "_样式统一备份") -> Path:
    """为源文件创建备份目录（每次合并独立，避免覆盖旧备份）。"""
    import time

    stamp = time.strftime("%Y%m%d_%H%M%S")
    candidate = parent / f"{prefix}_{stamp}"
    counter = 1
    while candidate.exists():
        candidate = parent / f"{prefix}_{stamp}_{counter}"
        counter += 1
    candidate.mkdir(parents=True)
    return candidate


def _paragraph_outline_level(paragraph) -> int | None:
    """段落的大纲级别：段落直接格式优先，其次样式链；正文段落返回 None。"""
    from docx.oxml.ns import qn

    def _level(ppr):
        if ppr is None:
            return None
        node = ppr.find(qn("w:outlineLvl"))
        if node is None:
            return None
        try:
            return int(node.get(qn("w:val")))
        except (TypeError, ValueError):
            return None

    level = _level(paragraph._p.pPr)
    if level is not None:
        return level
    style = paragraph.style
    seen: set[str] = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        level = _level(style.element.find(qn("w:pPr")))
        if level is not None:
            return level
        style = style.base_style
    return None


def _suppress_duplicate_heading_numbers_docx(document, log) -> int:
    """标题文字里已有手写编号时，取消样式带来的自动编号。

    导入带多级列表的模板后，样式会给标题加上自动编号，而标题文字本身
    已经写了「第1章」「1.1」，两者叠加就成了「1 第1章 概述」。
    这里按段落取消自动编号（numId=0），样式定义保持不变。
    只处理编号来自样式、段落自身没有声明编号的标题。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    fixed = 0
    for paragraph in document.paragraphs:
        if _paragraph_outline_level(paragraph) is None:
            continue
        text = paragraph.text.strip()
        if not text or not _RE_MANUAL_NUMBER.match(text):
            continue
        if _numbering_from_ppr(paragraph._p.pPr) is not None:
            continue  # 段落自己声明了编号，保持原样
        if _paragraph_numbering(paragraph) is None:
            continue  # 本来就没有自动编号
        ppr = paragraph._p.get_or_add_pPr()
        numpr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), "0")
        numpr.append(ilvl)
        numpr.append(num_id)
        _replace_ordered_child(ppr, numpr, _PPR_ORDER)
        fixed += 1
    if fixed:
        log(f"标题文字已带编号：已取消 {fixed} 个标题的自动编号，避免编号重复。")
    return fixed


def _suppress_duplicate_heading_numbers_word(document, log) -> int:
    """Word 版本：标题文字已带手写编号时，取消段落上的自动编号。"""
    fixed = 0
    try:
        total = int(document.Paragraphs.Count)
    except Exception:
        return 0
    for index in range(1, total + 1):
        try:
            paragraph = document.Paragraphs(index)
            outline = int(paragraph.OutlineLevel)
        except Exception:
            continue
        if not 1 <= outline <= 9:
            continue
        try:
            text = str(paragraph.Range.Text or "").strip()
            if not text or not _RE_MANUAL_NUMBER.match(text):
                continue
            list_format = paragraph.Range.ListFormat
            if not str(list_format.ListString or "").strip():
                continue
            list_format.RemoveNumbers()
            fixed += 1
        except Exception:
            continue
    if fixed:
        log(f"标题文字已带编号：已取消 {fixed} 个标题的自动编号，避免编号重复。")
    return fixed


def _process_styles_with_word(
    files: list[Path],
    template: Path | None,
    import_styles: bool,
    prune_styles: bool,
    backup_root: Path,
    temp_dir: Path,
    log,
    export_docx: bool = False,
    save_source: bool = True,
    output_dir: Path | None = None,
) -> tuple[list[Path], Path | None]:
    """用 Word 逐个文件导入模板样式 / 删除多余样式。

    save_source 为 True 时直接在源文件上修改（事先已备份）；为 False 时源文件保持
    不变，处理结果写到 output_dir（默认 temp_dir）下的 .docx 副本。
    export_docx 为 True 时，返回供合并使用的 .docx 路径列表。
    返回（供合并的 .docx 路径列表，模板基准文档路径）。
    """
    import shutil

    import pythoncom
    import win32com.client

    wd_format_document_default = 16
    merge_files: list[Path] = []
    base_docx: Path | None = None
    if import_styles and template is None:
        raise ValueError("导入样式模板前必须选择模板文件。")

    pythoncom.CoInitialize()
    word = None
    try:
        try:
            word = win32com.client.DispatchEx("Word.Application")
        except Exception as exc:
            raise WordUnavailableError(str(exc)) from exc

        word.Visible = False
        word.DisplayAlerts = 0
        try:
            word.ScreenUpdating = False
        except Exception:
            pass
        try:
            word.AutomationSecurity = 3  # 禁用宏
        except Exception:
            pass

        # 1) 读取模板样式名，并按需生成只保留样式的基准文档
        template_names: set[str] | None = None
        if import_styles:
            probe = None
            try:
                probe = word.Documents.Open(
                    FileName=str(template),
                    ConfirmConversions=False,
                    ReadOnly=True,
                    AddToRecentFiles=False,
                    Visible=False,
                )
                template_names = _word_style_names(probe)
            except Exception as exc:
                log(f"读取样式模板失败：{exc}")
            finally:
                try:
                    if probe is not None:
                        probe.Close(SaveChanges=0)
                except Exception:
                    pass

            if export_docx:
                try:
                    base = word.Documents.Open(
                        FileName=str(template),
                        ConfirmConversions=False,
                        ReadOnly=False,
                        AddToRecentFiles=False,
                        Visible=False,
                    )
                    try:
                        base.Content.Delete()  # 去掉模板正文，只保留样式与页面设置
                    except Exception:
                        pass
                    base_docx = temp_dir / "模板基准.docx"
                    base.SaveAs2(str(base_docx), wd_format_document_default)
                    base.Close(SaveChanges=0)
                except Exception as exc:
                    base_docx = None
                    log(f"生成模板基准文档失败（不影响样式导入）：{exc}")

        # 2) 逐个文件：导入模板样式 → 删除多余样式 → 保存
        for index, path in enumerate(files, start=1):
            if save_source:
                backup = backup_root / f"{index:03d}_{path.name}"
                shutil.copy2(path, backup)

            document = None
            working: Path | None = None
            processed_ok = False
            try:
                document = word.Documents.Open(
                    FileName=str(path),
                    ConfirmConversions=False,
                    ReadOnly=not save_source,
                    AddToRecentFiles=False,
                    Visible=False,
                )
                if not save_source:
                    # 改动只写进副本，源文件保持原样
                    target_dir = output_dir or temp_dir
                    target_dir.mkdir(parents=True, exist_ok=True)
                    working = target_dir / f"{index:03d}_{path.stem}.docx"
                    document.SaveAs2(str(working), wd_format_document_default)

                notes: list[str] = []
                if import_styles:
                    document.CopyStylesFromTemplate(str(template))
                    notes.append("已导入模板样式")
                    suppressed = _suppress_duplicate_heading_numbers_word(
                        document, log
                    )
                    if suppressed:
                        notes.append(f"取消 {suppressed} 个标题的自动编号")

                removed = 0
                if prune_styles:
                    if import_styles:
                        if template_names:
                            removed = _prune_word_styles(
                                document, None, template_names, log
                            )
                            notes.append(f"删除 {removed} 个样式")
                        else:
                            notes.append("模板样式读取失败，未清理样式")
                    elif path.suffix.lower() == ".docx" or working is not None:
                        document.Save()  # 先保存，便于统计样式引用情况
                        probe = working if working is not None else path
                        removed = _prune_word_styles(document, probe, None, log)
                        notes.append(f"删除 {removed} 个样式")
                    else:
                        notes.append("非 .docx 文件未选模板时无法统计样式，未清理")
                document.Save()

                if working is not None:
                    merge_files.append(working)
                elif export_docx:
                    if path.suffix.lower() == ".docx":
                        merge_files.append(path)
                    else:
                        converted = temp_dir / f"{index:03d}_{path.stem}.docx"
                        document.SaveAs2(str(converted), wd_format_document_default)
                        merge_files.append(converted)
                detail = "；".join(notes) or "无改动"
                if working is not None:
                    detail += f"；结果另存为 {working.name}（源文件未改动）"
                log(f"[{index}/{len(files)}] 已处理：{path.name}（{detail}）")
                processed_ok = True
            except Exception as exc:
                raise RuntimeError(f"处理样式失败（{path.name}）：{exc}") from exc
            finally:
                try:
                    if document is not None:
                        document.Close(SaveChanges=0)
                except Exception:
                    pass
                if prune_styles and processed_ok:
                    saved = working if working is not None else path
                    if saved.suffix.lower() == ".docx" and saved.is_file():
                        try:
                            _dedupe_docx_styles(saved, log)
                        except Exception as exc:
                            log(f"重复样式清理失败（{saved.name}）：{exc}")
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()

    return merge_files, base_docx


def _process_styles_python_docx(
    files: list[Path],
    template: Path | None,
    import_styles: bool,
    prune_styles: bool,
    backup_root: Path,
    log,
    save_source: bool = True,
    output_dir: Path | None = None,
) -> list[Path]:
    """无 Word 时的逐个文件样式处理（仅支持 .docx）。

    save_source 为 False 时不改动源文件，结果写到 output_dir。
    """
    import shutil

    from docx import Document

    if import_styles and template is None:
        raise ValueError("导入样式模板前必须选择模板文件。")

    non_docx = [p.name for p in files if p.suffix.lower() != ".docx"]
    if non_docx:
        raise RuntimeError(
            "未检测到可用的 Microsoft Word，只有 .docx 文件可以单独处理样式："
            + "、".join(non_docx)
        )

    results: list[Path] = []
    for index, path in enumerate(files, start=1):
        if save_source:
            backup = backup_root / f"{index:03d}_{path.name}"
            shutil.copy2(path, backup)
            target = path
        else:
            target_dir = output_dir or path.parent
            target_dir.mkdir(parents=True, exist_ok=True)
            target = target_dir / f"{index:03d}_{path.stem}.docx"

        document = Document(str(path))
        notes: list[str] = []
        names: set[str] | None = None
        if import_styles:
            names = _import_styles_python_docx(document, template, log)
            notes.append("已导入模板样式")
            suppressed = _suppress_duplicate_heading_numbers_docx(document, log)
            if suppressed:
                notes.append(f"取消 {suppressed} 个标题的自动编号")
        if prune_styles:
            if import_styles and not names:
                notes.append("模板样式读取失败，未清理样式")
            else:
                removed = _prune_python_docx_styles(
                    document, names if import_styles else None, log
                )
                notes.append(f"删除 {removed} 个样式")
        document.save(str(target))
        detail = "；".join(notes) or "无改动"
        if not save_source:
            detail += f"；结果另存为 {target.name}（源文件未改动）"
        log(f"[{index}/{len(files)}] 已处理：{path.name}（{detail}）")
        results.append(target)
    return results


def _make_output_dir(parent: Path, prefix: str = "样式处理结果") -> Path:
    """为“不保存源文件”的处理结果创建输出目录。"""
    import time

    stamp = time.strftime("%Y%m%d_%H%M%S")
    candidate = parent / f"{prefix}_{stamp}"
    counter = 1
    while candidate.exists():
        candidate = parent / f"{prefix}_{stamp}_{counter}"
        counter += 1
    candidate.mkdir(parents=True)
    return candidate


def apply_styles_to_files(
    files: list[Path],
    template: Path | None,
    import_styles: bool,
    prune_styles: bool,
    log,
    save_source: bool = True,
) -> Path:
    """单独执行样式模板导入 / 多余样式删除，不合并文档。

    save_source 为 True 时直接修改源文件（返回备份目录）；
    为 False 时源文件保持不变，结果输出到新建的“样式处理结果_时间”目录并返回该目录。
    """
    import shutil
    import tempfile

    if not files:
        raise ValueError("请先添加或选择要处理的 Word 文档。")
    if not import_styles and not prune_styles:
        raise ValueError("请至少勾选“导入样式模板”或“删除多余样式”中的一项。")
    if import_styles and template is None:
        raise ValueError("请先选择样式模板文件。")

    temp_dir = Path(tempfile.mkdtemp(prefix="wm_styles_"))
    if save_source:
        backup_root = _make_backup_dir(files[0].parent)
        output_dir = None
        log(f"开始单独处理 {len(files)} 个文档的样式（保存源文件）")
        log(f"源文件将备份到：{backup_root}")
    else:
        backup_root = files[0].parent  # 未使用，占位
        output_dir = _make_output_dir(files[0].parent)
        log(f"开始单独处理 {len(files)} 个文档的样式（不保存源文件）")
        log(f"处理结果将输出到：{output_dir}")
    try:
        try:
            _process_styles_with_word(
                files,
                template,
                import_styles,
                prune_styles,
                backup_root,
                temp_dir,
                log,
                save_source=save_source,
                output_dir=output_dir,
            )
        except WordUnavailableError:
            log("未检测到 Microsoft Word，改用 python-docx 处理（仅支持 .docx）。")
            _process_styles_python_docx(
                files,
                template,
                import_styles,
                prune_styles,
                backup_root,
                log,
                save_source=save_source,
                output_dir=output_dir,
            )
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)
    log("样式处理完成。")
    return output_dir if output_dir is not None else backup_root


def _load_template_base(base_docx: Path | None, template: Path, log):
    """加载模板基准文档（只保留样式与页面设置），作为 docxcompose 的主文档。"""
    from docx import Document
    from docx.oxml.ns import qn

    if base_docx is not None and base_docx.exists():
        document = Document(str(base_docx))
        body = document.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        return document

    document = Document()
    _import_styles_python_docx(document, template, log)
    return document


def _merge_with_unified_styles(files: list[Path], output: Path, options: MergeOptions, log) -> str:
    """先逐个文件统一到模板样式，再按章节顺序用 docxcompose 合并。"""
    import shutil
    import tempfile

    template = options.style_source
    save_source = options.save_source
    temp_dir = Path(tempfile.mkdtemp(prefix="wm_unify_"))
    backup_root = (
        _make_backup_dir(files[0].parent) if save_source else files[0].parent
    )

    log(f"第 1 步：按模板统一 {len(files)} 个文档的样式（模板：{template.name}）")
    if save_source:
        log(f"源文件将备份到：{backup_root}")
    else:
        log("已取消“保存源文件”：源文件保持不变，统一后的副本仅用于本次合并。")
    try:
        base_docx: Path | None = None
        try:
            merge_files, base_docx = _process_styles_with_word(
                files,
                template,
                True,
                options.prune_styles,
                backup_root,
                temp_dir,
                log,
                export_docx=True,
                save_source=save_source,
                output_dir=temp_dir,
            )
        except WordUnavailableError:
            log("未检测到 Microsoft Word，改用 python-docx 逐文件处理（仅支持 .docx）。")
            merge_files = _process_styles_python_docx(
                files,
                template,
                True,
                options.prune_styles,
                backup_root,
                log,
                save_source=save_source,
                output_dir=temp_dir,
            )

        log(f"第 2 步：按章节顺序用 docxcompose 合并 {len(merge_files)} 个文档（不再改动格式）")
        base_doc = _load_template_base(base_docx, template, log)
        _merge_with_docxcompose(merge_files, output, options, log, base_doc=base_doc)
        if options.make_toc:
            _update_fields_with_word(output, log)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

    return "样式统一 + docxcompose"


def _start_body_on_new_page(document) -> None:
    """删除目录与正文之间的空段落，并让第一章从新的一页开始。"""
    blank_indexes: list[int] = []
    first_index = None
    total = document.Paragraphs.Count
    for index in range(2, total + 1):  # 第 1 段是“目录”标题
        paragraph = document.Paragraphs(index)
        try:
            style_name = str(paragraph.Style.NameLocal or "")
        except Exception:
            style_name = ""
        if style_name.lower().startswith(("toc", "目录")):
            continue
        if paragraph.Range.Text.strip("\r\a\x07\x0c \t　"):
            first_index = index
            break
        blank_indexes.append(index)
    if first_index is None:
        return
    document.Paragraphs(first_index).Format.PageBreakBefore = True
    for index in reversed(blank_indexes):
        document.Paragraphs(index).Range.Delete()


def _word_last_paragraph_has_content(document) -> bool:
    """Word 文档最后一段是否还有内容（文字、图片、表格等）。"""
    try:
        count = int(document.Paragraphs.Count)
        if count < 1:
            return False
        paragraph = document.Paragraphs(count)
        if str(paragraph.Range.Text or "").strip("\r\n\a\x07\x0c \t　"):
            return True
        return bool(int(paragraph.Range.InlineShapes.Count))
    except Exception:
        return False


def _merge_with_word(files: list[Path], output: Path, options: MergeOptions, log):
    """使用 Microsoft Word 的 InsertFile 合并，最大程度保留原始格式。

    与默认引擎一致：第一个文档直接作为主文档打开，保留它的样式库、多级列表
    编号与页面设置，其余文档依次插入到文末。
    """
    import pythoncom
    import win32com.client

    wd_collapse_end = 0
    wd_page_break = 7
    wd_align_center = 1
    wd_format_document_default = 16

    if not files:
        raise ValueError("没有可合并的文档。")
    files = [Path(path).resolve() for path in files]

    pythoncom.CoInitialize()
    word = None
    document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            word.ScreenUpdating = False
        except Exception:
            pass
        try:
            word.AutomationSecurity = 3  # 禁用宏
        except Exception:
            pass

        document = word.Documents.Open(
            FileName=str(files[0]),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
            Visible=False,
        )
        log(f"以第一个文档为基准：{files[0].name}（保留其样式与页面设置）")

        toc = None
        if options.make_toc:
            title = (options.toc_title or "目录").strip() or "目录"
            start = document.Range(0, 0)
            start.InsertBefore(title + "\r")
            title_paragraph = document.Paragraphs(1)
            try:
                title_paragraph.Style = -1  # wdStyleNormal
            except Exception:
                pass
            title_range = title_paragraph.Range
            title_range.Bold = True
            title_range.Font.Size = 16
            try:
                title_range.ParagraphFormat.Alignment = wd_align_center
            except Exception:
                pass

            try:
                anchor = title_paragraph.Range.End
                toc = document.TablesOfContents.Add(
                    document.Range(anchor, anchor),
                    True,
                    1,
                    max(1, min(options.toc_levels, 5)),
                )
                # 目录与正文之间的分页交给 _start_body_on_new_page 处理：
                # 在这里插入分页符会把正文第一段（标题）拆成两个空标题段落，
                # 多出一个“第1章”的空标题，后面的章节编号也会整体后移。
            except Exception as exc:
                toc = None
                log(f"提示：自动生成目录失败（{exc}），文档已正常合并。")

        for index, path in enumerate(files, start=1):
            if index == 1:
                # 第一个文档本身就是主文档，无需再插入
                log(f"[{index}/{len(files)}] 已作为基准文档：{path.name}")
                continue
            target = document.Content
            target.Collapse(wd_collapse_end)
            if options.page_break:
                target.InsertBreak(wd_page_break)
                target.Collapse(wd_collapse_end)
            elif _word_last_paragraph_has_content(document):
                # 不分页时，末尾段落还有内容就补一个段落标记：
                # 否则插入的标题会被并进上一段，图片、样式一起被带乱
                target.InsertParagraphAfter()
                target.Collapse(wd_collapse_end)

            target.InsertFile(
                FileName=str(path),
                ConfirmConversions=False,
                Link=False,
                Attachment=False,
            )
            log(f"[{index}/{len(files)}] 已插入：{path.name}")

        # 仅按章节合并时不导入模板样式；只有显式勾选样式清理才处理样式
        style_names: set[str] | None = None
        if options.style_source and options.prune_styles:
            style_names = _import_styles_with_word(word, document, options.style_source, log)
        elif options.style_source:
            log(
                "已选择样式模板，但未勾选样式相关选项：本次合并不导入模板样式，"
                "各章节原有样式保持不变。"
            )

        # 合并后页码保持连续，避免各文件页码各自从 1 开始
        try:
            for i in range(1, document.Sections.Count + 1):
                document.Sections(i).PageNumbers.RestartNumberingAtSection = False
        except Exception:
            pass

        try:
            document.Fields.Update()
        except Exception:
            pass
        try:
            document.Repaginate()
        except Exception:
            pass

        toc_inserted = toc is not None
        if toc is not None:
            try:
                # 先让正文另起一页，再更新目录，页码才是调整后的
                _start_body_on_new_page(document)
                toc.Update()
                toc_text = toc.Range.Text or ""
                if "未找到" in toc_text or "No table of contents" in toc_text:
                    log("提示：未在正文中找到 Heading 1/2/3 标题样式，目录可能为空。")
            except Exception:
                pass

        # 目录更新完成后再清理样式，避免误删目录生成过程中用到的样式
        output.parent.mkdir(parents=True, exist_ok=True)
        if output.exists():
            output.unlink()
        document.SaveAs2(str(output), wd_format_document_default)

        if options.prune_styles:
            if options.style_source and not style_names:
                log("未能从模板中读取样式，已跳过样式清理，以免误删全部样式。")
            else:
                _prune_word_styles(document, output, style_names, log)

        log("目录已生成并更新页码。" if toc_inserted else "未生成目录。")
    finally:
        try:
            if document is not None:
                document.Close(SaveChanges=0)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()

    # Word 关闭后再清理重复样式：需要按正文引用的 styleId 区分同名样式
    if options.prune_styles and output.suffix.lower() == ".docx" and output.is_file():
        try:
            _dedupe_docx_styles(output, log)
        except Exception as exc:
            log(f"重复样式清理失败：{exc}")


def _set_page_break_before(element) -> bool:
    """让正文从新的一页开始：给段落加“段前分页”，而不是插入一个空的分页段落。

    空分页段落会在新页顶部留下一行空白，与原文档版式不一致。
    第一个元素不是段落（如表格）时，退回为在它前面插入分页段落。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if element is None or element.tag == qn("w:sectPr"):
        return False
    if element.tag != qn("w:p"):
        break_paragraph = _page_break_element()
        element.addprevious(break_paragraph)
        return True
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        element.insert(0, ppr)
    for existing in ppr.findall(qn("w:pageBreakBefore")):
        ppr.remove(existing)
    node = OxmlElement("w:pageBreakBefore")
    _replace_ordered_child(ppr, node, _PPR_ORDER)
    return True


def _add_toc_field(doc, title: str, levels: int = 3, at_start: bool = False) -> bool:
    """python-docx 回退方案：插入 TOC 域（打开文档时由 Word 更新）。

    返回 True 表示目录之后还有正文要追加，调用方需让第一段正文另起一页。
    at_start=True 时把目录放到正文最前面，并直接让原来的第一段正文另起一页：
    主文档不是空白文档时，add_paragraph 会把目录追加到已有正文之后。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    title_paragraph = doc.add_paragraph(title)
    try:
        title_paragraph.style = doc.styles["Title"]
    except Exception:
        for run in title_paragraph.runs:
            run.bold = True
    # 目录标题本身不进入目录：9 表示“正文级别”。模板的标题样式常带大纲级别，
    # 不处理的话目录里会多出一行指向目录自己。
    _set_ppr_outline(title_paragraph._p.get_or_add_pPr(), 9)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f'TOC \\o "1-{max(1, min(levels, 5))}" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "打开文档后按 F9 更新目录"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)

    pending_break = True
    if at_start:
        body = doc.element.body
        for offset, element in enumerate((title_paragraph._p, paragraph._p)):
            body.remove(element)
            body.insert(offset, element)
        children = list(body)
        if len(children) > 2:
            _set_page_break_before(children[2])
        pending_break = False

    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    return pending_break


def _strip_page_number_restart(doc) -> None:
    from docx.oxml.ns import qn

    for section in doc.sections:
        pg_num = section._sectPr.find(qn("w:pgNumType"))
        if pg_num is not None:
            pg_num.attrib.pop(qn("w:start"), None)


def _append_document(composer, master, document, page_break: bool) -> list:
    """把一份文档追加到主文档末尾；page_break 时让它的第一段另起一页。"""
    position = composer.append_index()
    added = _composer_insert(composer, master, position, document)
    if page_break and added:
        _set_page_break_before(added[0])
    return added


def _merge_with_docxcompose(
    files: list[Path],
    output: Path,
    options: MergeOptions,
    log,
    base_doc=None,
):
    """不依赖 Word 的 docxcompose 合并，要求源文件为 .docx。

    与「合并到文档」一致：第一个文档作为主文档，后续文档依次插入，
    主文档的样式库、多级列表编号、页面设置不会被改动。
    base_doc 为事先统一好样式的基准文档时，合并阶段不再改动任何样式。
    """
    from docx import Document
    from docxcompose.composer import Composer

    non_docx = [p.name for p in files if p.suffix.lower() != ".docx"]
    if non_docx:
        raise RuntimeError(
            "以下文件不是 .docx，无法用 docxcompose 合并："
            + "、".join(non_docx)
        )

    if base_doc is not None:
        master = base_doc
        remaining = list(files)
        break_before_each = False
    elif files:
        # 与“合并到文档”一致：用第一个文档作为主文档，它的样式库、编号定义、
        # 页面设置原样保留，后续文档只是插入进来，不会再被空白文档的默认样式覆盖。
        master = Document(str(files[0]))
        remaining = list(files[1:])
        break_before_each = True
        log(f"以第一个文档为基准：{files[0].name}（保留其样式与页面设置）")
    else:
        master = Document()
        remaining = []
        break_before_each = False

    toc_pending_break = False
    if options.make_toc:
        # 目录与正文之间必须分页；at_start 时直接让主文档原来的第一段另起一页，
        # 否则等到追加第一份文档时给它分页。
        toc_pending_break = _add_toc_field(
            master,
            options.toc_title or "目录",
            options.toc_levels,
            at_start=break_before_each,
        )

    # 纯章节合并时 preserve_styles=True：各章样式定义原样带进结果文档，
    # 同名但定义不同的样式只会另存一份副本，不会被主文档或前一章的样式覆盖。
    # base_doc 不为空表示样式已在合并前显式统一过，此时沿用原行为。
    composer = Composer(master, preserve_styles=base_doc is None)
    if base_doc is None:
        log("仅按章节拼接：各章原有格式与样式库保持不变，源文件只读。")
    # 与主文档样式相同的章节直接沿用主文档的样式与编号列表，标题编号由此连续；
    # 样式不同的章节按原样另带一份副本，并沿用该章节自己的编号起始值。
    for index, path in enumerate(remaining, start=1):
        need_break = options.page_break and (break_before_each or index > 1)
        if index == 1 and toc_pending_break:
            need_break = True
        _append_document(composer, master, Document(str(path)), need_break)
        log(f"[{index}/{len(remaining)}] 已合并：{path.name}")

    # base_doc 表示样式已在合并前统一过，这里不再导入/清理样式
    if base_doc is None:
        if options.prune_styles:
            style_names = None
            if options.style_source:
                if options.style_source.suffix.lower() == ".doc":
                    raise RuntimeError(
                        "无 Word 的回退方式不支持 .doc 样式模板，请先把模板另存为 .dotm 或 .dotx。"
                    )
                if not options.style_source.is_file():
                    raise RuntimeError(f"找不到样式模板文件：{options.style_source}")
                style_names = _import_styles_python_docx(
                    master, options.style_source, log
                )
            if options.style_source and not style_names:
                log("未能从模板中读取样式，已跳过样式清理，以免误删全部样式。")
            else:
                _prune_python_docx_styles(master, style_names, log)
        elif options.style_source:
            log(
                "已选择样式模板，但未勾选样式相关选项：本次合并不导入模板样式，"
                "各章节原有样式保持不变。"
            )

    _strip_page_number_restart(master)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        output.unlink()
    composer.save(str(output))
    if options.make_toc:
        log("已插入目录域，首次打开文档时 Word 会提示更新域，请选择“是”。")


def _convert_to_docx(files: list[Path], temp_dir: Path, log) -> list[Path]:
    """把 .doc/.docm/.rtf 用 Word 另存为临时 .docx，供 docxcompose 合并。"""
    if all(path.suffix.lower() == ".docx" for path in files):
        return files

    import pythoncom
    import win32com.client

    wd_format_document_default = 16
    converted: list[Path] = []
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            word.AutomationSecurity = 3
        except Exception:
            pass
        for index, path in enumerate(files, start=1):
            if path.suffix.lower() == ".docx":
                converted.append(path)
                continue
            document = word.Documents.Open(
                FileName=str(path),
                ConfirmConversions=False,
                ReadOnly=True,
                AddToRecentFiles=False,
                Visible=False,
            )
            try:
                target = temp_dir / f"{index:03d}_{path.stem}.docx"
                document.SaveAs2(str(target), wd_format_document_default)
                converted.append(target)
                log(f"已转换格式：{path.name} -> {target.name}")
            finally:
                try:
                    document.Close(SaveChanges=0)
                except Exception:
                    pass
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
    return converted


def _update_fields_with_word(output: Path, log) -> None:
    """可选：用 Word 更新目录等域的页码（不修改正文内容与编号）。"""
    try:
        import pythoncom
        import win32com.client
    except Exception:
        return
    try:
        pythoncom.CoInitialize()
    except Exception:
        pass
    word = None
    document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            word.AutomationSecurity = 3
        except Exception:
            pass
        document = word.Documents.Open(
            FileName=str(output),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
            Visible=False,
        )
        try:
            document.Fields.Update()
        except Exception:
            pass
        try:
            for index in range(1, document.TablesOfContents.Count + 1):
                document.TablesOfContents(index).Update()
        except Exception:
            pass
        try:
            document.Repaginate()
        except Exception:
            pass
        document.Save()
        log("已用 Word 更新目录页码。")
    except Exception as exc:
        log(f"提示：未能用 Word 更新目录页码（{exc}），打开文档后可选中目录按 F9 更新。")
    finally:
        try:
            if document is not None:
                document.Close(SaveChanges=0)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


# --------------------------------------------------------------------------
# 把选中的章节合并 / 替换到已有文档
# --------------------------------------------------------------------------

_RE_HEADING_STYLE = re.compile(r"^(?:heading|标题)([1-9])$", re.I)


@dataclass
class ChapterSpan:
    """目标文档中一章占用的正文元素区间 [start, end)。"""

    number: tuple[int, ...] | None
    heading: str
    start: int
    end: int


def _heading_level(paragraph) -> int | None:
    try:
        name = re.sub(r"\s+", "", paragraph.style.name or "")
    except Exception:
        return None
    match = _RE_HEADING_STYLE.match(name)
    return int(match.group(1)) if match else None


def _chapter_spans(document, any_heading: bool = False) -> list[ChapterSpan]:
    """按标题切分目标文档，返回每章的段落区间。

    any_heading 为 True 时接受任意级别的标题（用于没有一级标题的文档）。
    """
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    children = list(document.element.body)
    sect_index = len(children)
    for index, child in enumerate(children):
        if child.tag == qn("w:sectPr"):
            sect_index = index
            break

    spans: list[ChapterSpan] = []
    current: ChapterSpan | None = None
    for index, child in enumerate(children[:sect_index]):
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document._body)
        level = _heading_level(paragraph)
        if level is None or (not any_heading and level != 1):
            continue
        text = paragraph.text.strip()
        if not text:
            continue
        if current is not None:
            current.end = index
            spans.append(current)
        current = ChapterSpan(parse_chapter_number(text), text, index, sect_index)
    if current is not None:
        current.end = sect_index
        spans.append(current)
    return spans


def _document_spans(document) -> list[ChapterSpan]:
    spans = _chapter_spans(document) or _chapter_spans(document, any_heading=True)
    return _fill_span_numbers(document, spans)


def _fill_span_numbers(document, spans: list[ChapterSpan]) -> list[ChapterSpan]:
    """标题文字里没有章号的章节，用标题上的多级列表编号补出章号。"""
    if not spans or all(span.number is not None for span in spans):
        return spans
    from docx.oxml.ns import qn

    snapshot = _numbering_snapshot(document)
    if not snapshot:
        return spans
    children = list(document.element.body)
    filled: list[ChapterSpan] = []
    for span in spans:
        number = span.number
        if number is None and 0 <= span.start < len(children):
            child = children[span.start]
            if child.tag == qn("w:p"):
                info = snapshot.get(child)
                if info is not None:
                    number = (info[2],)
        filled.append(ChapterSpan(number, span.heading, span.start, span.end))
    return filled


def _page_break_element():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run.append(br)
    paragraph.append(run)
    return paragraph


def _chapter_label(number: tuple[int, ...] | None) -> str:
    if number:
        return "第" + ".".join(str(n) for n in number) + "章"
    return "未编号章节"


def _has_toc_field(document) -> bool:
    from docx.oxml.ns import qn

    for node in document.element.body.iter(qn("w:instrText")):
        text = (node.text or "").strip().upper()
        if text.startswith("TOC"):
            return True
    return False


def _strip_empty_edge_paragraphs(document) -> None:
    """去掉章节文档首尾的空段落，避免插入后多出空行。

    仅用于显式要求整理版式的场景；默认的章节合并要求原样保留章节内容，
    不会再调用此函数。
    """
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    body = document.element.body
    children = list(body)
    if not children:
        return

    def _is_safe_empty(child) -> bool:
        if child.tag != qn("w:p"):
            return False
        if list(child.iter(qn("w:drawing"))) or list(child.iter(qn("w:pict"))):
            return False
        if list(child.iter(qn("w:sectPr"))):
            return False
        return not Paragraph(child, document._body).text.strip()

    for child in children:
        if not _is_safe_empty(child):
            break
        body.remove(child)
    for child in reversed(children):
        if child.getparent() is None:
            continue
        if child.tag == qn("w:sectPr"):
            continue
        if not _is_safe_empty(child):
            break
        body.remove(child)


# --------------------------------------------------------------------------
# 合并 / 导入单章节：先选章节，再选总文档，同章节可替换
# --------------------------------------------------------------------------

_RE_SPLIT_SUFFIX = re.compile(r"[_\-]?\d{8}(?:[_\-]\d+)?$")


def _first_heading_text(document) -> str | None:
    """读取文档对象中的第一个标题（优先一级标题）。"""
    fallback: str | None = None
    for index, paragraph in enumerate(document.paragraphs):
        if index > 500:
            break
        text = paragraph.text.strip()
        if not text:
            continue
        level = _heading_level(paragraph)
        if level is None:
            continue
        if level == 1:
            return text
        if fallback is None:
            fallback = text
    return fallback


def _chapter_title_from_filename(path: Path) -> str:
    """从文件名中取章节标题（去掉章号、拆分时追加的日期后缀）。"""
    stem = _RE_SPLIT_SUFFIX.sub("", path.stem)
    return _chapter_title_from_heading(stem) or path.stem


def _normalized_title(text: str) -> str:
    """比较章节标题时统一格式：去掉编号、空白和下划线。"""
    title = _chapter_title_from_heading(text or "")
    return re.sub(r"[\s_]+", "", title).casefold()


def _match_chapter_span(spans: list, number, title: str):
    """在总文档的章节区间中找同章节：先按章号，再按标题文字。"""
    if number is not None:
        for span in spans:
            if span.number == number:
                return span
    key = _normalized_title(title)
    if key:
        for span in spans:
            if _normalized_title(span.heading) == key:
                return span
    return None


def _heading_numbering_map(document) -> dict[int, tuple[str, int]]:
    """总文档中各级标题实际使用的编号（numId, ilvl）。"""
    result: dict[int, tuple[str, int]] = {}
    for paragraph in _body_paragraphs(document):
        level = _heading_level(paragraph)
        if level is None or level in result:
            continue
        info = _paragraph_numbering(paragraph)
        if info is not None:
            result[level] = info
    return result


def _chapter_heading_numbering(document, span) -> dict[int, tuple[str, int]]:
    """某一章正文中各级标题实际使用的编号（numId, ilvl）。"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    children = list(document.element.body)
    result: dict[int, tuple[str, int]] = {}
    for position in range(span.start, min(span.end, len(children))):
        child = children[position]
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document._body)
        level = _heading_level(paragraph)
        if level is None or level in result:
            continue
        info = _paragraph_numbering(paragraph)
        if info is not None:
            result[level] = info
    return result


def _merge_heading_numbers(primary: dict, fallback: dict) -> dict:
    """以 primary 为主；缺的级别沿用同一列表的上级编号定义。"""
    base = primary.get(1) or fallback.get(1)
    if base is None:
        return dict(fallback)
    result = dict(primary)
    for level in range(1, 10):
        if level not in result:
            result[level] = (base[0], level - 1)
    return result


def _adopt_heading_numbering(document, elements: list, numbering_map: dict) -> None:
    """让导入章节的标题沿用总文档的编号列表。

    否则章节自带的编号会把列表拆成两段：被替换章节后面的章节
    会因为总文档的列表少了一个而整体前移（附录从第 3 章变成第 2 章）。
    """
    if not numbering_map:
        return
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    base = numbering_map.get(1)
    if base is None:
        return
    for element in elements:
        if element.tag != qn("w:p"):
            continue
        level = _heading_level(Paragraph(element, document._body))
        if level is None:
            continue
        num_id, ilvl = numbering_map.get(level, (base[0], level - 1))
        ppr = element.get_or_add_pPr()
        numpr = OxmlElement("w:numPr")
        ilvl_node = OxmlElement("w:ilvl")
        ilvl_node.set(qn("w:val"), str(ilvl))
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        numpr.append(ilvl_node)
        numpr.append(num_id_node)
        _replace_ordered_child(ppr, numpr, _PPR_ORDER)


def _composer_insert(composer, document, position: int, source) -> list:
    """在 position 处插入章节，返回插入正文中的元素。"""
    body = document.element.body
    before = len(body)
    composer.insert(position, source)
    children = list(body)
    added = len(children) - before
    if added <= 0:
        return []
    return children[position : position + added]


def _chapter_number_for_source(original: Path, converted: Path) -> tuple[int, ...] | None:
    """章节文档的章号：标题优先，标题里没有章号时用原文件名。"""
    number = parse_chapter_number(_first_heading(converted))
    if number is None:
        number = parse_chapter_number(original.stem)
    return number


def analyze_chapter_import(chapter: Path, target: Path, log) -> dict:
    """分析单章节导入：读章节标题/章号，并在总文档中查找同章节。"""
    import tempfile

    from docx import Document

    temp_dir = Path(tempfile.mkdtemp(prefix="wm_import_"))
    try:
        chapter_docx, target_docx = _convert_to_docx([chapter, target], temp_dir, log)
        chapter_document = Document(str(chapter_docx))
        heading = _first_heading_text(chapter_document) or ""
        number = parse_chapter_number(heading) or parse_chapter_number(chapter.stem)
        title = _chapter_title_from_heading(heading) if heading else ""
        if not title:
            title = _chapter_title_from_filename(chapter)
        spans = _document_spans(Document(str(target_docx)))
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

    match_index: int | None = None
    match = _match_chapter_span(spans, number, title)
    if match is not None:
        match_index = spans.index(match)
    return {
        "heading": heading or title,
        "number": number,
        "title": title,
        "spans": [{"heading": span.heading, "number": span.number} for span in spans],
        "match_index": match_index,
    }


def merge_chapters_into_document(
    target: Path,
    chapters: list[Path],
    output: Path,
    mode: str,
    options: MergeOptions,
    log,
) -> str:
    """把选中的章节合并（插入）或替换到目标文档中，结果保存为 output。

    mode == "replace"：按章号替换目标文档中的同名章节，找不到时按章号插入。
    mode == "merge"  ：按章号把章节插入目标文档的对应位置，原有内容保持不变。
    """
    import shutil
    import tempfile

    from docx import Document
    from docxcompose.composer import Composer

    temp_dir = Path(tempfile.mkdtemp(prefix="wm_into_"))
    try:
        target_docx = _convert_to_docx([target], temp_dir, log)[0]
        chapter_docx = _convert_to_docx(chapters, temp_dir, log)

        document = Document(str(target_docx))
        if not _document_spans(document):
            log("提示：目标文档中未找到章节标题，选中的章节将按章号追加到文末。")
        has_toc = _has_toc_field(document)
        # 只插入章节内容：章节自带的样式原样带进来，与目标文档同名但定义
        # 不同的样式另存副本，不改动目标文档已有的样式定义和章节内容。
        composer = Composer(document, preserve_styles=True)
        log("仅插入章节内容：目标文档样式库与各章格式保持不变，源文件只读。")

        heading_numbers = _heading_numbering_map(document)

        replaced = 0
        inserted = 0
        total = len(chapter_docx)
        for index, path in enumerate(chapter_docx, start=1):
            source = Document(str(path))
            number = _chapter_number_for_source(chapters[index - 1], path)

            if mode == "replace":
                span = _match_chapter_span(
                    _document_spans(document), number, _first_heading(path) or ""
                )
                if span is not None:
                    # 记下被替换章节原本使用的编号，替换后让新章节接上同一个列表
                    numbers = _merge_heading_numbers(
                        _chapter_heading_numbering(document, span), heading_numbers
                    )
                    body = document.element.body
                    for child in list(body)[span.start : span.end]:
                        body.remove(child)
                    added = _composer_insert(composer, document, span.start, source)
                    _adopt_heading_numbering(document, added, numbers)
                    replaced += 1
                    log(
                        f"[{index}/{total}] 已替换：{path.name} -> {span.heading}"
                    )
                    continue
                log(
                    f"目标文档中没有 {_chapter_label(number)}，改为按章号插入：{path.name}"
                )

            spans_now = _document_spans(document)
            position: int | None = None
            reference_span = spans_now[-1] if spans_now else None
            for span in spans_now:
                if (
                    span.number is not None
                    and number is not None
                    and span.number > number
                ):
                    position = span.start
                    break
                reference_span = span
            body = document.element.body
            if position is None:
                from docx.oxml.ns import qn

                sect = body.find(qn("w:sectPr"))
                position = list(body).index(sect) if sect is not None else len(body)
            if options.page_break and position > 0:
                body.insert(position, _page_break_element())
                position += 1
            numbers = (
                _merge_heading_numbers(
                    _chapter_heading_numbering(document, reference_span),
                    heading_numbers,
                )
                if reference_span is not None
                else heading_numbers
            )
            added = _composer_insert(composer, document, position, source)
            _adopt_heading_numbering(document, added, numbers)
            inserted += 1
            log(f"[{index}/{total}] 已合并：{path.name}")

        output.parent.mkdir(parents=True, exist_ok=True)
        if output.exists():
            output.unlink()
        document.save(str(output))
        log(f"已写入：{output}（替换 {replaced} 章，插入 {inserted} 章）")
        if has_toc:
            _update_fields_with_word(output, log)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

    return f"合并到文档（{mode == 'replace' and '替换' or '合并'}）"


# --------------------------------------------------------------------------
# 把总文档按章节拆分为单独文档
# --------------------------------------------------------------------------

_RE_FILENAME_BAD = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')
_RE_TITLE_PREFIX = re.compile(
    r"^\s*(?:"
    r"第\s*[零〇一二三四五六七八九十百两\d]+(?:[.．]\d+)*\s*[章篇部节]"
    r"|(?:chapter|chap\.?|ch\.?)\s*\d+"
    r"|\d+(?:[.．]\d+)*(?=[\s、,，.．:：_\-—])"
    r")\s*[\s、,，.．:：_\-—]*",
    re.I,
)


def _safe_filename_part(text: str, fallback: str, limit: int = 100) -> str:
    """去掉文件名中不允许的字符，并限制长度。"""
    cleaned = _RE_FILENAME_BAD.sub("_", text or "")
    cleaned = re.sub(r"\s+", " ", cleaned).strip().strip(".")
    cleaned = re.sub(r"\s*_\s*", "_", cleaned).strip("_")
    cleaned = re.sub(r"_+", "_", cleaned)
    return cleaned[:limit].rstrip(" ._") or fallback


def _chapter_title_from_heading(heading: str) -> str:
    """从标题中取章节标题，去掉开头的「第N章 / Chapter N / 1.2」等编号。"""
    text = (heading or "").strip()
    stripped = _RE_TITLE_PREFIX.sub("", text, count=1).strip()
    return stripped or text


def _chapter_number_text(number: tuple[int, ...] | None, index: int) -> str:
    """章号文本；识别不到章号时用章节顺序号。"""
    if number:
        return ".".join(str(part) for part in number)
    return str(index)


def _numbering_level_info(numbering_root) -> dict[str, dict[int, dict]]:
    """解析 numbering.xml：numId -> {ilvl: {start, restart}}。"""
    from docx.oxml.ns import qn

    def _int(node, default=None):
        if node is None:
            return default
        try:
            return int(node.get(qn("w:val")))
        except (TypeError, ValueError):
            return default

    def _ilvl_of(element, default=0):
        try:
            return int(element.get(qn("w:ilvl")))
        except (TypeError, ValueError):
            return default

    abstracts: dict[str, dict[int, dict]] = {}
    for abstract in numbering_root.findall(qn("w:abstractNum")):
        abstract_id = abstract.get(qn("w:abstractNumId"))
        levels: dict[int, dict] = {}
        for lvl in abstract.findall(qn("w:lvl")):
            ilvl = _ilvl_of(lvl)
            levels[ilvl] = {
                "start": _int(lvl.find(qn("w:start")), 1),
                # None 表示未声明（上一级递增时重置）；0 表示不重置
                "restart": _int(lvl.find(qn("w:lvlRestart"))),
            }
        abstracts[abstract_id] = levels

    result: dict[str, dict[int, dict]] = {}
    for num in numbering_root.findall(qn("w:num")):
        num_id = num.get(qn("w:numId"))
        ref = num.find(qn("w:abstractNumId"))
        if num_id is None or ref is None:
            continue
        levels = {
            ilvl: dict(info)
            for ilvl, info in abstracts.get(ref.get(qn("w:val")), {}).items()
        }
        for override in num.findall(qn("w:lvlOverride")):
            ilvl = _ilvl_of(override)
            start = _int(override.find(qn("w:startOverride")))
            if start is not None and ilvl in levels:
                levels[ilvl]["start"] = start
        result[num_id] = levels
    return result


def _numbering_from_ppr(ppr):
    """从 w:pPr 中读取 (numId, ilvl)；未声明时返回 None。"""
    from docx.oxml.ns import qn

    if ppr is None:
        return None
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        return None
    num_id_node = numpr.find(qn("w:numId"))
    num_id = num_id_node.get(qn("w:val")) if num_id_node is not None else None
    ilvl_node = numpr.find(qn("w:ilvl"))
    ilvl = None
    if ilvl_node is not None:
        try:
            ilvl = int(ilvl_node.get(qn("w:val")))
        except (TypeError, ValueError):
            ilvl = None
    return num_id, ilvl


def _paragraph_numbering(paragraph) -> tuple[str, int] | None:
    """段落实际使用的多级列表编号（numId, ilvl），编号来自段落或样式链。"""
    from docx.oxml.ns import qn

    num_id: str | None = None
    ilvl: int | None = None
    direct = _numbering_from_ppr(paragraph._p.pPr)
    if direct is not None:
        direct_id, direct_level = direct
        if direct_id is not None:
            if direct_id == "0":
                return None  # numId=0 表示取消编号
            num_id = direct_id
        ilvl = direct_level

    style = paragraph.style
    seen: set[str] = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        info = _numbering_from_ppr(style.element.find(qn("w:pPr")))
        if info is not None:
            style_id_value, style_level = info
            if num_id is None and style_id_value is not None:
                if style_id_value == "0":
                    return None
                num_id = style_id_value
            if ilvl is None:
                ilvl = style_level
        style = style.base_style

    if not num_id:
        return None
    return num_id, ilvl or 0


def _body_paragraphs(document):
    """按文档顺序遍历正文中的段落（含表格内的段落）。"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    for child in list(document.element.body):
        if child.tag == qn("w:p"):
            yield Paragraph(child, document._body)
        elif child.tag == qn("w:tbl"):
            for element in child.iter(qn("w:p")):
                yield Paragraph(element, document._body)


def _numbering_snapshot(document) -> dict:
    """模拟整篇文档的多级列表编号，返回 {段落元素: (numId, ilvl, 编号值)}。

    拆分会删掉前面的章节，Word 会把剩下的标题重新从 1 开始编号；
    这里先按 Word 的规则算出每个段落原本的编号值，拆分后再写回起始值。
    """
    from docx.oxml.ns import qn

    try:
        numbering_root = document.part.numbering_part.element
    except (Exception, NotImplementedError):
        return {}
    levels = _numbering_level_info(numbering_root)
    if not levels:
        return {}

    # 共用同一 abstractNum 的多个列表在 Word 中算同一个列表，计数要合在一起
    abstract_of: dict[str, str] = {}
    for num in numbering_root.findall(qn("w:num")):
        ref = num.find(qn("w:abstractNumId"))
        if ref is not None:
            abstract_of[num.get(qn("w:numId"))] = ref.get(qn("w:val"))

    counters: dict[str, dict[int, int]] = {}
    snapshot: dict = {}
    for paragraph in _body_paragraphs(document):
        info = _paragraph_numbering(paragraph)
        if info is None:
            continue
        num_id, ilvl = info
        definition = levels.get(num_id)
        if definition is None:
            continue
        state = counters.setdefault(abstract_of.get(num_id, num_id), {})
        if ilvl in state:
            state[ilvl] += 1
        else:
            state[ilvl] = int(definition.get(ilvl, {}).get("start", 1))
        for deeper in [level for level in state if level > ilvl]:
            restart = definition.get(deeper, {}).get("restart")
            # None：默认随上一级递增而重置；0：不重置；N：第 N 级递增时重置
            if restart is None or restart == ilvl + 1:
                del state[deeper]
        snapshot[paragraph._p] = (num_id, ilvl, state[ilvl])
    return snapshot


def _chapter_numbering_starts(snapshot: dict, children: list, span) -> dict[str, dict[int, int]]:
    """取本章各编号级别第一次出现时的编号值，用作拆分文档的起始编号。"""
    from docx.oxml.ns import qn

    starts: dict[str, dict[int, int]] = {}
    for position in range(span.start, min(span.end, len(children))):
        child = children[position]
        if child.tag != qn("w:p"):
            continue
        info = snapshot.get(child)
        if info is None:
            continue
        num_id, ilvl, value = info
        starts.setdefault(num_id, {}).setdefault(ilvl, value)
    return starts


def _numbering_levels_of(abstract) -> dict[int, object]:
    """abstractNum 中 ilvl -> w:lvl 元素。"""
    from docx.oxml.ns import qn

    levels: dict[int, object] = {}
    for lvl in abstract.findall(qn("w:lvl")):
        try:
            levels[int(lvl.get(qn("w:ilvl")))] = lvl
        except (TypeError, ValueError):
            continue
    return levels


def _set_level_start(lvl, value: str) -> None:
    """设置 w:lvl 的起始编号（w:start 必须是第一个子元素）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = lvl.find(qn("w:start"))
    if start is None:
        start = OxmlElement("w:start")
        lvl.insert(0, start)
    start.set(qn("w:val"), value)


def _apply_numbering_starts(document, starts: dict[str, dict[int, int]]) -> None:
    """把起始编号写回拆分文档的 numbering.xml。

    Word 中只写 w:startOverride 时，第 2 级及更深的级别在上一级递增后
    仍从 1 开始，因此这里改的是级别的起始编号（w:start）：
    共享同一 abstractNum 的多个列表会各自复制一份编号定义，互不影响。
    """
    if not starts:
        return
    from copy import deepcopy
    from random import randrange

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    try:
        numbering_root = document.part.numbering_part.element
    except (Exception, NotImplementedError):
        return

    nums = {
        num.get(qn("w:numId")): num for num in numbering_root.findall(qn("w:num"))
    }
    abstracts = {
        abstract.get(qn("w:abstractNumId")): abstract
        for abstract in numbering_root.findall(qn("w:abstractNum"))
    }
    usage: dict[str, int] = {}
    for num in nums.values():
        ref = num.find(qn("w:abstractNumId"))
        if ref is not None:
            usage[ref.get(qn("w:val"))] = usage.get(ref.get(qn("w:val")), 0) + 1

    numeric_ids = [int(key) for key in abstracts if str(key).isdigit()]
    next_id = max(numeric_ids, default=0) + 1

    for num_id, levels in starts.items():
        num = nums.get(num_id)
        if num is None:
            continue
        ref = num.find(qn("w:abstractNumId"))
        if ref is None:
            continue
        abstract = abstracts.get(ref.get(qn("w:val")))
        if abstract is None:
            continue

        if usage.get(ref.get(qn("w:val")), 0) > 1:
            # 同一编号定义被多个列表共用：复制一份，避免改动其他列表的编号
            abstract = deepcopy(abstract)
            abstract.set(qn("w:abstractNumId"), str(next_id))
            nsid = abstract.find(qn("w:nsid"))
            if nsid is not None:
                nsid.set(qn("w:val"), f"{randrange(0x10000000, 0xFFFFFFFF):08X}")
            last_abstract = numbering_root.findall(qn("w:abstractNum"))[-1]
            last_abstract.addnext(abstract)
            abstracts[str(next_id)] = abstract
            ref.set(qn("w:val"), str(next_id))
            next_id += 1

        definitions = _numbering_levels_of(abstract)
        for ilvl, value in levels.items():
            lvl = definitions.get(ilvl)
            if lvl is None:
                continue
            _set_level_start(lvl, str(value))
            # 该列表原有的级别覆盖也会盖住新的起始值，一并同步
            for override in num.findall(qn("w:lvlOverride")):
                if override.get(qn("w:ilvl")) != str(ilvl):
                    continue
                start_override = override.find(qn("w:startOverride"))
                if start_override is not None:
                    start_override.set(qn("w:val"), str(value))
                inner = override.find(qn("w:lvl"))
                if inner is not None:
                    _set_level_start(inner, str(value))


def _strip_trailing_page_breaks(document) -> None:
    """去掉章节末尾仅含分页符（或空白）的段落，避免拆分后多出一页空白。"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    body = document.element.body
    harmless = {
        qn("w:p"),
        qn("w:r"),
        qn("w:pPr"),
        qn("w:rPr"),
        qn("w:br"),
        qn("w:lastRenderedPageBreak"),
        qn("w:bookmarkStart"),
        qn("w:bookmarkEnd"),
        qn("w:proofErr"),
    }
    for child in reversed(list(body)):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag != qn("w:p"):
            break
        paragraph = Paragraph(child, document._body)
        if paragraph.text.strip():
            break
        ppr = child.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            break
        nodes = [
            node
            for node in child.iter()
            if node is not child
            and node.tag not in {qn("w:pPr"), qn("w:rPr"), qn("w:proofErr")}
        ]
        if any(
            node.tag == qn("w:br") and node.get(qn("w:type")) not in (None, "page")
            for node in nodes
        ):
            break  # 含换行、分栏等分隔符，保留
        if any(node.tag not in harmless for node in nodes):
            break  # 含图片、域、内容控件等，保留
        body.remove(child)


def split_document_by_chapters(
    source: Path,
    output_dir: Path,
    log,
    date_suffix: str | None = None,
) -> list[Path]:
    """把总文档按章节标题拆分为单独文档，格式保持不变。

    每个拆分文档都是原文档的完整副本，只删除其他章节的正文元素，
    因此样式、主题、编号、页眉页脚、图片、表格等格式全部原样保留。
    文件名为「章号_章节标题_日期.docx」，日期默认为当天（YYYYMMDD）。
    """
    from datetime import datetime

    from docx import Document
    from docx.oxml.ns import qn

    suffix = date_suffix or datetime.now().strftime("%Y%m%d")
    temp_dir = Path(tempfile.mkdtemp(prefix="wm_split_"))
    outputs: list[Path] = []
    try:
        docx_path = _convert_to_docx([source], temp_dir, log)[0]
        document = Document(str(docx_path))
        spans = _document_spans(document)
        if not spans:
            raise ValueError(
                "文档中未找到章节标题（优先使用一级标题，没有一级标题时使用其他级别标题），无法拆分。"
            )

        total = len(spans)
        output_dir.mkdir(parents=True, exist_ok=True)
        log(f"共识别 {total} 个章节，开始拆分：{source.name}")

        children = list(document.element.body)
        snapshot = _numbering_snapshot(document)

        used: set[str] = {source.stem.lower()}
        for index, span in enumerate(spans, start=1):
            number_text = _chapter_number_text(span.number, index)
            title = _chapter_title_from_heading(span.heading) or "章节"
            stem = _safe_filename_part(
                f"{number_text}_{title}_{suffix}",
                f"{number_text}_章节_{suffix}",
            )
            candidate = stem
            duplicate = 2
            while candidate.lower() in used:
                candidate = f"{stem}_{duplicate}"
                duplicate += 1
            used.add(candidate.lower())
            target = output_dir / f"{candidate}.docx"

            chapter = Document(str(docx_path))
            body = chapter.element.body
            for position, child in enumerate(list(body)):
                if child.tag == qn("w:sectPr"):
                    continue
                if span.start <= position < span.end:
                    continue
                body.remove(child)
            _strip_trailing_page_breaks(chapter)
            _apply_numbering_starts(
                chapter, _chapter_numbering_starts(snapshot, children, span)
            )
            chapter.save(str(target))
            outputs.append(target)
            log(f"[{index}/{total}] {span.heading} -> {target.name}")

        log(f"拆分完成，共生成 {len(outputs)} 个文档：{output_dir}")
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)
    return outputs


# --------------------------------------------------------------------------
# 删除标记区间：*ocs_start* … *ocs_end* 之间的内容
# --------------------------------------------------------------------------

DEFAULT_MARKER_START = "*ocs_start*"
DEFAULT_MARKER_END = "*ocs_end*"

# 删除位置两侧需要一并清掉的分隔字符（半角空格、制表符、不换行空格、全角空格）
_JOIN_WHITESPACE = " \t 　"

# 出现这些内容的空段落不删除：删掉会影响图片、域、页码、编号、分节、
# 书签或修订记录，可能让 Word 提示文档损坏
_EMPTY_PARA_KEEP_TAGS = (
    "w:drawing",
    "w:pict",
    "w:object",
    "w:sectPr",
    "w:footnoteReference",
    "w:endnoteReference",
    "w:commentReference",
    "w:numPr",
    "w:fldChar",
    "w:fldSimple",
    "w:txbxContent",
    "w:bookmarkStart",
    "w:bookmarkEnd",
    "w:commentRangeStart",
    "w:commentRangeEnd",
    "w:permStart",
    "w:permEnd",
    "w:ins",
    "w:del",
)


def _nested_paragraph(element, body) -> bool:
    """element 是否嵌套在另一个段落里（文本框中的段落），这类段落不参与查找。"""
    from docx.oxml.ns import qn

    parent = element.getparent()
    while parent is not None and parent is not body:
        if parent.tag == qn("w:p"):
            return True
        parent = parent.getparent()
    return False


def _all_paragraph_elements(document) -> list:
    """按正文顺序返回所有段落元素，包含表格里的段落，不含文本框内的段落。"""
    from docx.oxml.ns import qn

    body = document.element.body
    return [
        element
        for element in body.iter(qn("w:p"))
        if not _nested_paragraph(element, body)
    ]


def _paragraph_segments(paragraph_element) -> list[tuple]:
    """把段落拆成按顺序排列的文字片段：[(节点, 类型, 文字, 起始, 结束)]。

    类型为 text（w:t）或 separator（制表符、换行）。文本框里的内容不计入。
    """
    from docx.oxml.ns import qn

    segments: list[tuple] = []
    offset = 0
    for node in paragraph_element.iter():
        if node is paragraph_element:
            continue
        tag = node.tag
        if tag == qn("w:t"):
            text = node.text or ""
            kind = "text"
        elif tag == qn("w:tab"):
            text = "\t"
            kind = "separator"
        elif tag in (qn("w:br"), qn("w:cr")):
            text = "\n"
            kind = "separator"
        else:
            continue
        parent = node.getparent()
        nested = False
        while parent is not None and parent is not paragraph_element:
            if parent.tag == qn("w:p"):
                nested = True
                break
            parent = parent.getparent()
        if nested:
            continue
        segments.append((node, kind, text, offset, offset + len(text)))
        offset += len(text)
    return segments


def _paragraph_plain_text(paragraph_element) -> str:
    """段落中可见文字（含制表符与换行），与 Word 里看到的一致。"""
    return "".join(segment[2] for segment in _paragraph_segments(paragraph_element))


def _expand_join_whitespace(text: str, start: int, end: int) -> tuple[int, int]:
    """把删除区间向外扩到相邻的空格/制表符，避免删完留下多余空白。"""
    while start > 0 and text[start - 1] in _JOIN_WHITESPACE:
        start -= 1
    while end < len(text) and text[end] in _JOIN_WHITESPACE:
        end += 1
    return start, end


def _delete_text_in_paragraph(paragraph_element, start: int, end: int) -> None:
    """删除段落文字中的 [start, end)，只改文字，不动任何格式。"""
    if end <= start:
        return
    from docx.oxml.ns import qn

    for node, kind, text, seg_start, seg_end in _paragraph_segments(paragraph_element):
        if seg_end <= start or seg_start >= end:
            continue
        if kind == "separator":
            if node.getparent() is not None:
                node.getparent().remove(node)
            continue
        cut_from = max(start, seg_start) - seg_start
        cut_to = min(end, seg_end) - seg_start
        remaining = text[:cut_from] + text[cut_to:]
        node.text = remaining
        if remaining != remaining.strip():
            node.set(qn("xml:space"), "preserve")


def _paragraph_is_blank(element) -> bool:
    """段落是否已经没有可见内容（空段落或只剩空白）。"""
    from docx.oxml.ns import qn

    if element is None or element.tag != qn("w:p"):
        return False
    for tag in _EMPTY_PARA_KEEP_TAGS:
        if next(element.iter(qn(tag)), None) is not None:
            return False
    return not _paragraph_plain_text(element).strip(" \t\r\n 　")


def _remove_blank_paragraphs_between(parent, left, right) -> int:
    """删除 parent 中 left 与 right 之间的空段落（left/right 可为 None）。"""
    children = list(parent)
    try:
        start_index = children.index(left) + 1 if left is not None else 0
        end_index = children.index(right) if right is not None else len(children)
    except ValueError:
        return 0
    if end_index <= start_index:
        return 0
    removed = 0
    for child in children[start_index:end_index]:
        if _paragraph_is_blank(child):
            parent.remove(child)
            removed += 1
    return removed


def _neighbour_outside_blank(element, backwards: bool):
    """沿指定方向跳过空段落，返回第一个有内容的元素（可为 None）。"""
    node = element
    while node is not None and _paragraph_is_blank(node):
        node = node.getprevious() if backwards else node.getnext()
    return node


def _marker_variants(marker: str) -> list[str]:
    """标记本身与去掉外侧星号后的写法都认（Word 里常写成 *ocs_start*）。"""
    variants = [marker]
    stripped = marker.strip("*").strip()
    if stripped and stripped != marker:
        variants.append(stripped)
    return variants


def _find_marker(text: str, marker: str, start: int = 0):
    """在文字中查找标记，返回 (起始, 结束)；标记外包裹的星号一并算进区间。"""
    best = None
    for candidate in _marker_variants(marker):
        index = text.find(candidate, start)
        if index == -1:
            continue
        end = index + len(candidate)
        # 把紧贴标记的星号一起删掉，避免留下 ** 或 *
        while index > 0 and text[index - 1] == "*":
            index -= 1
        while end < len(text) and text[end] == "*":
            end += 1
        if best is None or index < best[0]:
            best = (index, end)
    return best


def _find_marker_pairs(document, start_marker: str, end_marker: str, log) -> list[tuple]:
    """找出所有标记对的 (起始段, 起始位置, 结束段, 结束位置)。"""
    pairs: list[tuple] = []
    pending = None
    for element in _all_paragraph_elements(document):
        text = _paragraph_plain_text(element)
        if not text:
            continue
        cursor = 0
        while cursor <= len(text):
            start_hit = _find_marker(text, start_marker, cursor)
            end_hit = _find_marker(text, end_marker, cursor)
            if start_hit is None and end_hit is None:
                break
            if start_hit is not None and (end_hit is None or start_hit[0] <= end_hit[0]):
                if pending is not None:
                    log(
                        f"提示：{start_marker} 后没有找到 {end_marker}，"
                        "上一处起始标记已忽略。"
                    )
                pending = (element, start_hit[0])
                cursor = max(start_hit[1], cursor + 1)
            else:
                if pending is None:
                    log(f"提示：发现多余的 {end_marker}，已忽略。")
                else:
                    pairs.append((pending[0], pending[1], element, end_hit[1]))
                    pending = None
                cursor = max(end_hit[1], cursor + 1)
    if pending is not None:
        log(f"提示：文档末尾的 {start_marker} 没有对应的 {end_marker}，未做删除。")
    return pairs


def _delete_marker_pair(start_p, start_off: int, end_p, end_off: int) -> int:
    """删除一对标记之间的内容，返回顺带清理掉的空段落数。"""
    if start_p is end_p:
        text = _paragraph_plain_text(start_p)
        begin, finish = _expand_join_whitespace(text, start_off, end_off)
        _delete_text_in_paragraph(start_p, begin, finish)
        if not _paragraph_is_blank(start_p):
            return 0
        parent = start_p.getparent()
        left_ctx = _neighbour_outside_blank(start_p.getprevious(), backwards=True)
        right_ctx = _neighbour_outside_blank(start_p.getnext(), backwards=False)
        parent.remove(start_p)
        return 1 + _remove_blank_paragraphs_between(parent, left_ctx, right_ctx)

    parent = start_p.getparent()
    if end_p.getparent() is not parent:
        raise ValueError(
            "标记跨越了表格或文本框的边界，无法安全删除，"
            "请把两个标记放在同一段落/同一单元格内。"
        )

    start_text = _paragraph_plain_text(start_p)
    end_text = _paragraph_plain_text(end_p)
    begin, _ = _expand_join_whitespace(start_text, start_off, len(start_text))
    _, finish = _expand_join_whitespace(end_text, 0, end_off)

    siblings = list(parent)
    first = siblings.index(start_p)
    last = siblings.index(end_p)
    # 标记之间的整块内容（段落、表格、图片等）全部删除
    for child in siblings[first + 1 : last]:
        parent.remove(child)

    _delete_text_in_paragraph(start_p, begin, len(start_text))
    _delete_text_in_paragraph(end_p, 0, finish)

    removed = 0
    if _paragraph_is_blank(start_p):
        left_ctx = _neighbour_outside_blank(start_p.getprevious(), backwards=True)
        parent.remove(start_p)
        removed += 1
    else:
        left_ctx = start_p
    if _paragraph_is_blank(end_p):
        right_ctx = _neighbour_outside_blank(end_p.getnext(), backwards=False)
        parent.remove(end_p)
        removed += 1
    else:
        right_ctx = end_p
    removed += _remove_blank_paragraphs_between(parent, left_ctx, right_ctx)
    return removed


def _ensure_containers_have_paragraph(document) -> None:
    """删除后保证正文和每个单元格至少还有一个段落，否则 Word 会报文档损坏。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    body = document.element.body
    for cell in body.iter(qn("w:tc")):
        if cell.find(qn("w:p")) is None:
            cell.append(OxmlElement("w:p"))
    if body.find(qn("w:p")) is None and body.find(qn("w:tbl")) is None:
        paragraph = OxmlElement("w:p")
        section = body.find(qn("w:sectPr"))
        if section is not None:
            section.addprevious(paragraph)
        else:
            body.append(paragraph)


def delete_marked_sections(
    source: Path,
    output: Path,
    start_marker: str = DEFAULT_MARKER_START,
    end_marker: str = DEFAULT_MARKER_END,
    log=print,
) -> dict:
    """删除文档中 start_marker … end_marker 之间的内容（含两个标记本身）。

    只删文字、不动其余段落的格式；删除位置两侧遗留的空段落和空格一并清理。
    返回 {"ranges": 删除的区间数, "paragraphs": 清理的空段落数}。
    """
    from docx import Document

    start_marker = (start_marker or "").strip()
    end_marker = (end_marker or "").strip()
    if not start_marker or not end_marker:
        raise ValueError("起始标记和结束标记不能为空。")
    if start_marker == end_marker:
        raise ValueError("起始标记和结束标记不能相同。")

    document = Document(str(source))
    pairs = _find_marker_pairs(document, start_marker, end_marker, log)
    cleaned = 0
    if pairs:
        # 从后往前删：同一段落里的后一处标记先删，前面记录的偏移量才不会失效
        for start_p, start_off, end_p, end_off in reversed(pairs):
            cleaned += _delete_marker_pair(start_p, start_off, end_p, end_off)
        _ensure_containers_have_paragraph(document)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return {"ranges": len(pairs), "paragraphs": cleaned}


def delete_marked_sections_in_files(
    files: list[Path],
    start_marker: str = DEFAULT_MARKER_START,
    end_marker: str = DEFAULT_MARKER_END,
    log=print,
    save_source: bool = False,
) -> tuple[Path, list[tuple[str, str]]]:
    """批量删除标记区间。

    save_source 为 True 时直接改源文件（先备份，返回备份目录）；
    为 False 时源文件不动，结果写入新建的“标记删除结果_时间”目录并返回该目录。
    返回（输出目录或备份目录，失败的 [(文件名, 原因)]）；单个文件失败不影响其他文件。
    """
    import shutil
    import tempfile

    if not files:
        raise ValueError("请先添加或选择要处理的 Word 文档。")

    failures: list[tuple[str, str]] = []
    temp_dir = Path(tempfile.mkdtemp(prefix="wm_marked_"))
    if save_source:
        backup_root = _make_backup_dir(files[0].parent, "_标记删除备份")
        output_dir = None
        log(f"开始删除标记区间（保存源文件），源文件将备份到：{backup_root}")
    else:
        backup_root = files[0].parent  # 占位，未使用
        output_dir = _make_output_dir(files[0].parent, "标记删除结果")
        log(f"开始删除标记区间（不保存源文件），结果输出到：{output_dir}")

    try:
        for index, path in enumerate(files, start=1):
            docx_path = _convert_to_docx([path], temp_dir, log)[0]
            if save_source:
                shutil.copy2(path, backup_root / f"{index:03d}_{path.name}")
                if path.suffix.lower() == ".docx":
                    target = path
                else:
                    target = path.with_suffix(".docx")
                    log(f"注意：{path.name} 不是 .docx，结果另存为 {target.name}")
            else:
                target = output_dir / f"{path.stem}.docx"

            try:
                result = delete_marked_sections(
                    docx_path, target, start_marker, end_marker, log
                )
            except Exception as exc:
                failures.append((path.name, str(exc)))
                log(f"[{index}/{len(files)}] {path.name}：删除失败（{exc}）")
                continue
            if result["ranges"]:
                log(
                    f"[{index}/{len(files)}] {path.name}："
                    f"删除 {result['ranges']} 处标记区间，"
                    f"清理 {result['paragraphs']} 个空段落"
                )
            else:
                log(
                    f"[{index}/{len(files)}] {path.name}："
                    f"未找到 {start_marker} … {end_marker}，已原样复制"
                )
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

    return (output_dir if output_dir is not None else backup_root), failures


# --------------------------------------------------------------------------
# 分析文档中实际用到的样式，并按模板逐个样式替换格式
# --------------------------------------------------------------------------

def _rpr_format(rpr) -> dict:
    """从 w:rPr 中提取字体信息。"""
    if rpr is None:
        return {}
    W = f"{{{_W_NS}}}"
    result: dict = {}
    fonts = rpr.find(f"{W}rFonts")
    if fonts is not None:
        east = fonts.get(f"{W}eastAsia")
        ascii_name = fonts.get(f"{W}ascii")
        hansi = fonts.get(f"{W}hAnsi")
        result["font"] = east or ascii_name or hansi
        if east and ascii_name and east != ascii_name:
            result["font_detail"] = f"{east} / {ascii_name}"
    size = rpr.find(f"{W}sz")
    if size is not None:
        try:
            result["size"] = int(size.get(f"{W}val")) / 2
        except (TypeError, ValueError):
            pass
    bold = rpr.find(f"{W}b")
    if bold is not None:
        val = bold.get(f"{W}val")
        result["bold"] = val not in ("0", "false", "off")
    return result


def _style_format_table(root) -> tuple[dict, dict]:
    """解析 styles.xml 根元素，返回（按 styleId 索引的样式表，默认字体信息）。"""
    W = f"{{{_W_NS}}}"
    defaults: dict = {}
    doc_defaults = root.find(f"{W}docDefaults")
    if doc_defaults is not None:
        rpr = doc_defaults.find(f"{W}rPrDefault/{W}rPr")
        defaults = _rpr_format(rpr)

    table: dict[str, dict] = {}
    for element in root.findall(f"{W}style"):
        style_id = (element.get(f"{W}styleId") or "").strip()
        if not style_id:
            continue
        name_element = element.find(f"{W}name")
        name = (
            (name_element.get(f"{W}val") if name_element is not None else "") or style_id
        ).strip()
        based_element = element.find(f"{W}basedOn")
        based_id = (
            (based_element.get(f"{W}val") or "").strip()
            if based_element is not None
            else ""
        )
        outline = None
        ppr = element.find(f"{W}pPr")
        if ppr is not None:
            outline_element = ppr.find(f"{W}outlineLvl")
            if outline_element is not None:
                try:
                    outline = int(outline_element.get(f"{W}val"))
                except (TypeError, ValueError):
                    outline = None
        table[style_id] = {
            "name": name,
            "type": element.get(f"{W}type") or "paragraph",
            "based": based_id,
            "outline": outline,
            "fmt": _rpr_format(element.find(f"{W}rPr")),
            "default": element.get(f"{W}default") in ("1", "true", "on"),
        }

    resolved: dict[str, dict] = {}

    def resolve(style_id: str, seen: tuple[str, ...] = ()) -> dict:
        if style_id in resolved:
            return resolved[style_id]
        if style_id in seen or style_id not in table:
            return {}
        info = table[style_id]
        merged: dict = {}
        based = info["based"]
        if based:
            merged.update(resolve(based, seen + (style_id,)))
        merged.update(info["fmt"])
        resolved[style_id] = merged
        return merged

    for style_id in list(table):
        resolve(style_id)
    return table, defaults


# pPr / rPr 子元素的架构顺序：插入或替换格式属性时保持合法顺序，避免 Word 提示修复
_PPR_ORDER = {
    name: index
    for index, name in enumerate(
        (
            "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
            "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
            "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
            "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
            "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
            "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl",
            "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
        )
    )
}

_RPR_ORDER = {
    name: index
    for index, name in enumerate(
        (
            "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
            "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
            "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
            "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
            "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
            "eastAsianLayout", "specVanish", "oMath",
        )
    )
}

# 套用模板样式时会整体替换的字符格式属性（移除目标原有值，写入模板值）
_RPR_FORMAT_TAGS = (
    "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
    "outline", "shadow", "emboss", "imprint", "color", "spacing", "w",
    "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr",
    "shd", "fitText", "vertAlign", "rtl", "cs", "em",
)

# 段落格式中不随模板覆盖的属性：编号、大纲级别、分页与结构属性保持不变
_PPR_KEEP_TAGS = {
    "pStyle", "numPr", "outlineLvl", "rPr", "sectPr", "pPrChange",
    "framePr", "pageBreakBefore", "divId", "cnfStyle",
}

# “清除直接段落格式”时移除的段落属性（不影响编号与分页）
_PPR_DIRECT_TAGS = (
    "ind", "spacing", "jc", "textAlignment", "contextualSpacing",
    "tabs", "mirrorIndents",
)


def _local_name(tag: str) -> str:
    return tag.split("}")[-1] if "}" in tag else tag


def _replace_ordered_child(parent, child, order: dict) -> None:
    """用 child 替换 parent 中同名子元素，并按架构顺序插入。"""
    name = _local_name(child.tag)
    rank = order.get(name, len(order))
    for existing in list(parent):
        if existing.tag == child.tag:
            parent.remove(existing)
    index = len(parent)
    for position, existing in enumerate(parent):
        if order.get(_local_name(existing.tag), len(order)) > rank:
            index = position
            break
    parent.insert(index, child)


def _style_chain(style_id: str, table: dict) -> list[str]:
    """沿 basedOn 找到继承链（从最上级到当前样式）。"""
    chain: list[str] = []
    seen: set[str] = set()
    current = style_id
    while current and current in table and current not in seen:
        seen.add(current)
        chain.append(current)
        current = table[current]["based"]
    chain.reverse()
    return chain


def _merged_template_rpr(root, chain: list[str]):
    """把样式及其 basedOn 上级、文档默认值的 rPr 逐层合并，得到完整字符格式。"""
    from copy import deepcopy

    merged = _defaults_rpr_element(root)
    for style_id in chain:
        rpr = _style_rpr_element(root, style_id)
        if rpr is None:
            continue
        if merged is None:
            merged = deepcopy(rpr)
            continue
        for child in rpr:
            name = _local_name(child.tag)
            if name in ("rStyle", "rPrChange"):
                continue
            _replace_ordered_child(merged, deepcopy(child), _RPR_ORDER)
    return merged


def _merged_template_ppr(root, chain: list[str]):
    """把样式及其 basedOn 上级、文档默认值的 pPr 逐层合并，得到完整段落格式。"""
    from copy import deepcopy

    W = f"{{{_W_NS}}}"
    base = root.find(f"{W}docDefaults/{W}pPrDefault/{W}pPr")
    merged = deepcopy(base) if base is not None else None
    for style_id in chain:
        element = _find_style_element(root, style_id)
        if element is None:
            continue
        ppr = element.find(f"{W}pPr")
        if ppr is None:
            continue
        if merged is None:
            merged = deepcopy(ppr)
            continue
        for child in ppr:
            name = _local_name(child.tag)
            if name in _PPR_KEEP_TAGS:
                continue
            _replace_ordered_child(merged, deepcopy(child), _PPR_ORDER)
    if merged is None:
        return None
    for child in list(merged):
        if _local_name(child.tag) in _PPR_KEEP_TAGS:
            merged.remove(child)
    return merged if len(merged) else None


def _resolve_theme_fonts_in_rpr(rpr, themes: dict) -> None:
    """把主题字体引用（如 majorEastAsia）还原为模板中的实际字体名。"""
    from docx.oxml.ns import qn

    if rpr is None or not themes:
        return
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        return
    for slot, theme in (
        ("ascii", "asciiTheme"),
        ("hAnsi", "hAnsiTheme"),
        ("eastAsia", "eastAsiaTheme"),
        ("cs", "cstheme"),
    ):
        value = fonts.get(qn(f"w:{theme}"))
        if value in themes:
            fonts.set(qn(f"w:{slot}"), themes[value])
            fonts.attrib.pop(qn(f"w:{theme}"), None)


def _style_chain_numbering(root, chain: list[str]):
    """沿继承链找到样式实际使用的编号定义，返回（numId, ilvl）或 None。"""
    from docx.oxml.ns import qn

    for style_id in reversed(chain):  # 先看当前样式，再看上级
        element = _find_style_element(root, style_id)
        if element is None:
            continue
        ppr = element.find(qn("w:pPr"))
        if ppr is None:
            continue
        numpr = ppr.find(qn("w:numPr"))
        if numpr is None:
            continue
        num_id = numpr.find(qn("w:numId"))
        value = num_id.get(qn("w:val")) if num_id is not None else None
        if not value or value == "0":
            return None
        ilvl = numpr.find(qn("w:ilvl"))
        try:
            level = int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
        except (TypeError, ValueError):
            level = 0
        return {"num_id": value, "ilvl": level}
    return None


def _style_chain_outline(table: dict, style_id: str) -> int | None:
    """沿继承链解析样式的大纲级别（0 表示标题 1）。"""
    for parent_id in reversed(_style_chain(style_id, table)):
        outline = table[parent_id]["outline"]
        if outline is not None:
            return outline
    return None


def _template_style_catalog(template: Path) -> list[dict]:
    """列出模板中的全部段落样式及其完整字符/段落格式（供“任意样式”选择使用）。"""
    root = _read_styles_xml(template)
    table, defaults = _style_format_table(root)
    themes = _theme_fonts(template)

    entries: list[dict] = []
    for style_id, info in table.items():
        if info["type"] != "paragraph":
            continue
        chain = _style_chain(style_id, table)
        fmt = dict(defaults)
        for parent_id in chain:
            fmt.update(
                {
                    k: v
                    for k, v in table[parent_id]["fmt"].items()
                    if k in ("font", "font_detail", "size", "bold")
                }
            )
        rpr = _merged_template_rpr(root, chain)
        _resolve_theme_fonts_in_rpr(rpr, themes)
        if rpr is not None:
            fmt.update(_rpr_format(rpr))
        entries.append(
            {
                "style_id": style_id,
                "name": info["name"],
                "outline": _style_chain_outline(table, style_id),
                "default": info["default"],
                "fmt": fmt,
                "rpr": rpr,
                "ppr": _merged_template_ppr(root, chain),
                "numbering": _style_chain_numbering(root, chain),
                "based": info["based"],
            }
        )

    def sort_key(entry: dict):
        generated = entry["name"].startswith("样式 ") or " + " in entry["name"]
        return (
            generated,
            0 if entry["default"] else 1,
            entry["outline"] if entry["outline"] is not None else 10,
            entry["name"].casefold(),
        )

    entries.sort(key=sort_key)
    return entries


def _template_style_label(entry: dict) -> str:
    """模板样式在下拉框中的显示文本。"""
    fmt = entry["fmt"]
    font = fmt.get("font") or "继承字体"
    size = f"{fmt['size']:g}pt" if fmt.get("size") else "继承字号"
    bold = " 加粗" if fmt.get("bold") else ""
    return f"{entry['name']}（{font} {size}{bold}）"


def _normalized_style_name(name: str) -> str:
    return re.sub(r"\s+", "", name or "").casefold()


def _heading_level_from_name(name: str) -> int | None:
    match = _RE_HEADING_STYLE.match(re.sub(r"\s+", "", name or ""))
    return int(match.group(1)) if match else None


def _suggest_template_style(
    style_name: str, level: int | None, is_default: bool, catalog: list[dict]
):
    """为文档中的一个样式推荐模板中对应的样式。

    优先同名样式；标题再按大纲级别匹配，并优先选择带多级列表编号的变体
    （例如模板中的“标题1”）；正文优先“正文 / 段落正文 / Body Text”。
    """
    if not catalog:
        return None
    key = _normalized_style_name(style_name)
    exact = [e for e in catalog if _normalized_style_name(e["name"]) == key]
    if exact:
        numbered = [e for e in exact if e.get("numbering")]
        if numbered:
            return numbered[0]
        if level is not None:
            # 同名的是不带编号的 heading N，模板里还有带多级列表的“标题 N”
            variants = [
                e
                for e in catalog
                if e["outline"] == level - 1
                and e.get("numbering")
                and re.fullmatch(
                    rf"(标题|heading)\s*{level}", re.sub(r"\s+", "", e["name"]), re.I
                )
            ]
            if variants:
                return variants[0]
        return exact[0]

    if level is not None:
        candidates = [e for e in catalog if e["outline"] == level - 1]
        named = [
            e
            for e in candidates
            if re.fullmatch(
                rf"(标题|heading)\s*{level}", re.sub(r"\s+", "", e["name"]), re.I
            )
        ]
        for pool in (named, candidates):
            if pool:
                numbered = [e for e in pool if e.get("numbering")]
                return (numbered or pool)[0]

    if is_default or key in ("normal", "正文", "bodytext", "段落正文"):
        for target in ("正文", "段落正文", "bodytext"):
            match = next(
                (e for e in catalog if _normalized_style_name(e["name"]) == target),
                None,
            )
            if match is not None:
                return match
        default = next(
            (e for e in catalog if e["default"] and e["outline"] is None), None
        )
        if default is not None:
            return default

    if key:
        # 表格文字 / 图注等常见样式名，优先按用途匹配模板中的专用样式
        purposes = (
            (("表字体", "表格文字", "表格文本", "表内文字"), ("表字体", "表格")),
            (("表注", "表题", "表标题"), ("表注", "表题", "表标题")),
            (("图注", "图题", "图标题"), ("图注", "图题", "图标题")),
            (("注意事项", "提示", "注意"), ("注意事项", "提示")),
        )
        for triggers, targets in purposes:
            if not any(word in key for word in triggers):
                continue
            for target in targets:
                match = next(
                    (
                        e
                        for e in catalog
                        if _normalized_style_name(e["name"]) == target
                    ),
                    None,
                )
                if match is not None:
                    return match
            loose = next(
                (
                    e
                    for e in catalog
                    if any(
                        word in _normalized_style_name(e["name"]) for word in targets
                    )
                ),
                None,
            )
            if loose is not None:
                return loose

        fuzzy = [
            e
            for e in catalog
            if len(_normalized_style_name(e["name"])) >= 2
            and (
                key in _normalized_style_name(e["name"])
                or _normalized_style_name(e["name"]) in key
            )
        ]
        if fuzzy:
            return min(
                fuzzy,
                key=lambda e: abs(len(_normalized_style_name(e["name"])) - len(key)),
            )
    return None


def _defaults_rpr_element(root):
    """模板的文档默认 rPr（作为样式继承链的基底）。"""
    from copy import deepcopy

    W = f"{{{_W_NS}}}"
    node = root.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr")
    if node is None:
        return None
    rpr = deepcopy(node)
    return rpr if len(rpr) else None


def _style_rpr_element(root, style_id: str):
    """按 styleId 找到样式元素里的 w:rPr（用于复制格式）。"""
    from lxml import etree

    W = f"{{{_W_NS}}}"
    for element in root.findall(f"{W}style"):
        if (element.get(f"{W}styleId") or "").strip() == style_id:
            rpr = element.find(f"{W}rPr")
            return etree.fromstring(etree.tostring(rpr)) if rpr is not None else None
    return None


def _raw_style_id(paragraph) -> str:
    """读取段落 w:pStyle 的原始值（样式定义缺失时 python-docx 会抛异常）。"""
    try:
        from docx.oxml.ns import qn

        ppr = paragraph._p.pPr
        if ppr is not None:
            node = ppr.find(qn("w:pStyle"))
            if node is not None:
                return (node.get(qn("w:val")) or "").strip()
    except Exception:
        pass
    return ""


def _paragraph_style_name(paragraph) -> str:
    try:
        return paragraph.style.name or ""
    except Exception:
        return _raw_style_id(paragraph)


def _paragraph_style_id(paragraph) -> str:
    try:
        return paragraph.style.style_id or ""
    except Exception:
        # 样式定义缺失时 python-docx 会抛异常，退回 pStyle 的原始值
        return _raw_style_id(paragraph)


def _paragraph_range(document, chapter_index: int | None):
    """返回要处理的正文元素区间（整篇或某一章）。"""
    body = document.element.body
    children = list(body)
    if chapter_index is None:
        return 0, len(children)
    spans = _document_spans(document)
    if 0 <= chapter_index < len(spans):
        return spans[chapter_index].start, spans[chapter_index].end
    return 0, len(children)


def _range_paragraph_entries(document, chapter_index: int | None) -> list[dict]:
    """返回段落及其来源信息：是否在表格内、属于第几个表格（从 0 开始）。"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    body = document.element.body
    start, end = _paragraph_range(document, chapter_index)
    entries: list[dict] = []
    table_index = 0
    for element in list(body)[start:end]:
        if element.tag == qn("w:p"):
            entries.append(
                {"paragraph": Paragraph(element, document._body), "table": None}
            )
        elif element.tag == qn("w:tbl"):
            for inner in element.iter(qn("w:p")):
                entries.append(
                    {
                        "paragraph": Paragraph(inner, document._body),
                        "table": table_index,
                    }
                )
            table_index += 1
    return entries


def _range_paragraphs(document, chapter_index: int | None) -> list:
    """返回要处理的段落列表（含表格内的段落），按文档顺序排列。"""
    return [
        item["paragraph"]
        for item in _range_paragraph_entries(document, chapter_index)
    ]


def _table_style_element(element):
    """表格的 w:tblStyle 元素（表格样式的 styleId）。"""
    from docx.oxml.ns import qn

    tbl_pr = element.find(qn("w:tblPr"))
    if tbl_pr is None:
        return None
    return tbl_pr.find(qn("w:tblStyle"))


def _describe_tables(document, chapter_index: int | None) -> list[dict]:
    """识别文档中的表格：样式名、行数、列数、单元格内用到的段落样式。"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    body = document.element.body
    start, end = _paragraph_range(document, chapter_index)
    style_table, _ = _style_format_table(document.styles.element)
    tables: list[dict] = []
    for element in list(body)[start:end]:
        if element.tag != qn("w:tbl"):
            continue
        style_node = _table_style_element(element)
        style_id = (style_node.get(qn("w:val")) or "").strip() if style_node is not None else ""
        style_name = style_table.get(style_id, {}).get("name") or ""
        rows = len(element.findall(qn("w:tr")))
        columns = 0
        for row in element.findall(qn("w:tr")):
            columns = max(columns, len(row.findall(qn("w:tc"))))
        paragraph_styles: dict[str, int] = {}
        for inner in element.iter(qn("w:p")):
            if not Paragraph(inner, document._body).text.strip():
                continue
            name = _paragraph_style_name(Paragraph(inner, document._body))
            if name:
                paragraph_styles[name] = paragraph_styles.get(name, 0) + 1
        tables.append(
            {
                "index": len(tables) + 1,
                "style_id": style_id,
                "style": style_name or (style_id or "（无表格样式）"),
                "rows": rows,
                "columns": columns,
                "text_styles": paragraph_styles,
            }
        )
    return tables


def _paragraph_direct_format(paragraph):
    """取段落文字上直接设置的字体与字号（没有设置时返回 None）。"""
    from docx.oxml.ns import qn

    font = size = None
    for run in paragraph.runs:
        rpr = run._element.rPr
        if rpr is not None and not font:
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                font = (
                    fonts.get(qn("w:eastAsia"))
                    or fonts.get(qn("w:ascii"))
                    or fonts.get(qn("w:hAnsi"))
                )
        if not font and run.font.name:
            font = run.font.name
        if size is None and run.font.size is not None:
            size = run.font.size.pt
        if font and size is not None:
            break
    return font, size


def _effective_style_format(style_table: dict, style_defaults: dict, style_id: str) -> dict:
    """样式继承链合并后的字体与字号。"""
    fmt = dict(style_defaults)
    if style_id:
        for parent_id in _style_chain(style_id, style_table):
            fmt.update(
                {
                    k: v
                    for k, v in style_table[parent_id]["fmt"].items()
                    if k in ("font", "size")
                }
            )
    return fmt


# 标题文字中已经带有的手动编号（第 1 章、1.2 特性、一、（一）等）
_RE_MANUAL_NUMBER = re.compile(
    r"^\s*(?:"
    r"第\s*[0-9零〇一二三四五六七八九十百两]+\s*[章节篇部]"
    r"|(?:chapter|chap\.?|ch\.?)\s*[0-9]+"
    r"|[0-9]+(?:\.[0-9]+)+[\s　、.．]"
    r"|[0-9]+[.．、]\s"
    r"|[（(]\s*(?:[0-9]+|[一二三四五六七八九十]+)\s*[)）]"
    r"|(?:[0-9]+|[一二三四五六七八九十]+)\s*[、.．](?![0-9])"
    r"|[一二三四五六七八九十]+、"
    r")",
    re.I,
)


def analyze_document_styles(
    path: Path, template: Path, chapter_index: int | None, log
) -> dict:
    """分析文档中实际用到的每个段落样式、表格，并推荐模板中对应的样式。

    返回 {"styles": [样式行...], "tables": [表格信息...]}。
    """
    import shutil
    import tempfile

    from docx import Document

    temp_dir = Path(tempfile.mkdtemp(prefix="wm_analyze_"))
    try:
        docx_path = _convert_to_docx([path], temp_dir, log)[0]
        document = Document(str(docx_path))

        style_root = document.styles.element
        style_table, style_defaults = _style_format_table(style_root)
        catalog = _template_style_catalog(template)

        stats: dict[str, dict] = {}
        for item in _range_paragraph_entries(document, chapter_index):
            paragraph = item["paragraph"]
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = _paragraph_style_name(paragraph)
            lowered = style_name.lower().replace(" ", "")
            if lowered.startswith(("toc", "目录")):
                continue
            style_id = _paragraph_style_id(paragraph)
            key = style_id or f"name:{style_name}"
            entry = stats.setdefault(
                key,
                {
                    "style_id": style_id,
                    "style": style_name,
                    "count": 0,
                    "table_count": 0,
                    "tables": set(),
                    "fonts": {},
                    "sizes": {},
                    "texts": [],
                },
            )
            entry["count"] += 1
            entry["texts"].append(text)
            if item["table"] is not None:
                entry["table_count"] += 1
                entry["tables"].add(item["table"])
            font, size = _paragraph_direct_format(paragraph)
            if font:
                entry["fonts"][font] = entry["fonts"].get(font, 0) + 1
            if size is not None:
                entry["sizes"][size] = entry["sizes"].get(size, 0) + 1

        rows: list[dict] = []
        for entry in stats.values():
            style_id = entry["style_id"]
            style_name = entry["style"]
            if not style_id:
                # 个别情况下拿不到 styleId，按名称回查
                style_id = next(
                    (
                        sid
                        for sid, info in style_table.items()
                        if info["name"] == style_name
                    ),
                    "",
                )
            info = style_table.get(style_id, {})
            outline = _style_chain_outline(style_table, style_id) if style_id else None
            level = outline + 1 if outline is not None else _heading_level_from_name(style_name)
            fmt = _effective_style_format(style_table, style_defaults, style_id)
            font = (
                max(entry["fonts"], key=entry["fonts"].get)
                if entry["fonts"]
                else fmt.get("font")
            )
            size = (
                max(entry["sizes"], key=entry["sizes"].get)
                if entry["sizes"]
                else fmt.get("size")
            )
            numbering = (
                _style_chain_numbering(style_root, _style_chain(style_id, style_table))
                if style_id
                else None
            )
            if numbering is not None and not _numbering_id_exists(
                document, numbering["num_id"]
            ):
                # 样式引用了编号定义但文档里找不到（常见于样式导入/合并后）
                numbering = None
            key = _normalized_style_name(style_name)
            bodyish = bool(info.get("default")) or key in (
                "normal", "正文", "段落正文", "bodytext"
            ) or "正文" in key
            tpl = _suggest_template_style(
                style_name, level, bool(info.get("default")), catalog
            )
            table_only = (
                entry["table_count"] > 0
                and entry["table_count"] == entry["count"]
            )
            rows.append(
                {
                    "style_id": style_id,
                    "style": style_name,
                    "count": entry["count"],
                    "table_count": entry["table_count"],
                    "table_only": table_only,
                    "tables": sorted(entry["tables"]),
                    "position": (
                        "表格"
                        if table_only
                        else ("正文" if entry["table_count"] == 0 else "正文+表格")
                    ),
                    "level": level or 0,
                    "label": (
                        f"标题 {level}"
                        if level
                        else ("正文" if bodyish else "自定义")
                    ),
                    "font": font or "（继承）",
                    "size": f"{size:g} pt" if size else "（继承）",
                    "numbered": bool(numbering),
                    "manual_numbered": any(
                        _RE_MANUAL_NUMBER.match(item) for item in entry["texts"]
                    )
                    if level
                    else False,
                    "template_style": tpl["name"] if tpl else None,
                    "template_style_id": tpl["style_id"] if tpl else None,
                    "template_numbered": bool(tpl and tpl.get("numbering")),
                }
            )

        rows.sort(
            key=lambda row: (
                1 if row["table_only"] else 0,
                row["level"] if row["level"] else 99,
                -row["count"],
                row["style"].casefold(),
            )
        )
        return {"styles": rows, "tables": _describe_tables(document, chapter_index)}
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def _apply_template_rpr(target_rpr, source_rpr) -> None:
    """把模板样式的字符格式写入目标 rPr，并清掉目标原有的同名格式。"""
    from copy import deepcopy

    from docx.oxml.ns import qn

    if source_rpr is None:
        return

    # 字体按属性合并：模板没声明的字体（例如东亚字体）沿用目标文档原有设置。
    # Word 中主题字体（asciiTheme 等）优先于显式字体（ascii），所以设置显式字体时
    # 必须去掉对应的主题字体，否则模板字体不会生效。
    source_fonts = source_rpr.find(qn("w:rFonts"))
    if source_fonts is not None:
        target_fonts = target_rpr.find(qn("w:rFonts"))
        if target_fonts is None:
            target_fonts = deepcopy(source_fonts)
            _replace_ordered_child(target_rpr, target_fonts, _RPR_ORDER)

        slots = (
            ("ascii", "asciiTheme"),
            ("hAnsi", "hAnsiTheme"),
            ("eastAsia", "eastAsiaTheme"),
            ("cs", "cstheme"),
        )
        explicit = {slot: source_fonts.get(qn(f"w:{slot}")) for slot, _ in slots}
        themes = {slot: source_fonts.get(qn(f"w:{theme}")) for slot, theme in slots}
        # 模板只给了西文字体时，中文字体也沿用同一字体
        fallback = explicit["eastAsia"] or explicit["ascii"] or explicit["hAnsi"]

        for slot, theme in slots:
            if explicit[slot]:
                target_fonts.set(qn(f"w:{slot}"), explicit[slot])
                target_fonts.attrib.pop(qn(f"w:{theme}"), None)
            elif themes[slot]:
                target_fonts.set(qn(f"w:{theme}"), themes[slot])
                target_fonts.attrib.pop(qn(f"w:{slot}"), None)
        if fallback and not explicit["eastAsia"]:
            target_fonts.set(qn("w:eastAsia"), fallback)
            target_fonts.attrib.pop(qn("w:eastAsiaTheme"), None)

        known = {qn(f"w:{key}") for pair in slots for key in pair}
        for key, value in source_fonts.attrib.items():
            if key not in known:
                target_fonts.set(key, value)

    # 其余字符格式属性整体以模板为准，避免目标样式残留加粗、字号、颜色等旧格式
    for tag in _RPR_FORMAT_TAGS:
        for node in target_rpr.findall(qn(f"w:{tag}")):
            target_rpr.remove(node)
    for child in source_rpr:
        if _local_name(child.tag) in _RPR_FORMAT_TAGS:
            _replace_ordered_child(target_rpr, deepcopy(child), _RPR_ORDER)


def _apply_template_ppr(target_ppr, source_ppr) -> None:
    """把模板样式的段落格式写入目标 pPr，编号与大纲级别保持不变。"""
    from copy import deepcopy

    if source_ppr is None:
        return
    for child in list(target_ppr):
        if _local_name(child.tag) not in _PPR_KEEP_TAGS:
            target_ppr.remove(child)
    for child in source_ppr:
        if _local_name(child.tag) in _PPR_KEEP_TAGS:
            continue
        _replace_ordered_child(target_ppr, deepcopy(child), _PPR_ORDER)


def _clear_direct_char_format(paragraph) -> None:
    """清除段落与文字上直接设置的字符格式，让样式定义完全生效。"""
    from docx.oxml.ns import qn

    tags = ("rFonts",) + _RPR_FORMAT_TAGS

    def strip(rpr) -> None:
        for tag in tags:
            for node in rpr.findall(qn(f"w:{tag}")):
                rpr.remove(node)

    ppr = paragraph._p.pPr
    if ppr is not None:
        rpr = ppr.find(qn("w:rPr"))
        if rpr is not None:
            strip(rpr)
    for run in paragraph.runs:
        rpr = run._element.rPr
        if rpr is not None:
            strip(rpr)


def _clear_direct_paragraph_format(paragraph) -> None:
    """清除段落上直接设置的缩进、行距、对齐等，让样式定义完全生效。"""
    from docx.oxml.ns import qn

    ppr = paragraph._p.pPr
    if ppr is None:
        return
    for tag in _PPR_DIRECT_TAGS:
        for node in ppr.findall(qn(f"w:{tag}")):
            ppr.remove(node)


def _ensure_style_ppr(style_element):
    """取样式元素中的 pPr，没有则按架构顺序创建。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = style_element.find(qn("w:pPr"))
    if ppr is not None:
        return ppr
    ppr = OxmlElement("w:pPr")
    rpr = style_element.find(qn("w:rPr"))
    if rpr is not None:
        rpr.addprevious(ppr)
    else:
        style_element.append(ppr)
    return ppr


def _ensure_style_rpr(style_element):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rpr = style_element.find(qn("w:rPr"))
    if rpr is not None:
        return rpr
    rpr = OxmlElement("w:rPr")
    style_element.append(rpr)
    return rpr


def _set_ppr_outline(ppr, outline: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    node = OxmlElement("w:outlineLvl")
    node.set(qn("w:val"), str(outline))
    _replace_ordered_child(ppr, node, _PPR_ORDER)


def _set_style_numbering(ppr, num_id: int, ilvl: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    numpr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), str(ilvl))
    numpr.append(level)
    value = OxmlElement("w:numId")
    value.set(qn("w:val"), str(num_id))
    numpr.append(value)
    _replace_ordered_child(ppr, numpr, _PPR_ORDER)


def _numbering_id_exists(document, num_id) -> bool:
    """文档的 numbering.xml 中是否存在指定 numId 的定义。"""
    if not num_id:
        return False
    try:
        root = document.part.numbering_part.element
    except Exception:
        return False
    from docx.oxml.ns import qn

    return any(
        (node.get(qn("w:numId")) or "").strip() == str(num_id)
        for node in root.findall(qn("w:num"))
    )


def _numbering_root(document):
    """取文档的 numbering.xml 根元素；文档没有编号部件时新建一个。"""
    from docx.opc.constants import CONTENT_TYPE as CT
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.parts.numbering import NumberingPart

    try:
        return document.part.numbering_part.element
    except NotImplementedError:
        # 源文档没有 numbering.xml，python-docx 无法自动创建，这里手动补上
        root = parse_xml(f'<w:numbering {nsdecls("w")}/>')
        part = NumberingPart(
            PackURI("/word/numbering.xml"),
            CT.WML_NUMBERING,
            root,
            document.part.package,
        )
        document.part.relate_to(part, RT.NUMBERING)
        return root


def _import_numbering_definition(
    document, template: Path, template_num_id: str, cache: dict
) -> int | None:
    """把模板中的多级列表定义复制到目标文档，返回新的 numId。

    同一模板 numId 在一次套用中只导入一次，保证“标题 1/2/3”共用同一套
    多级列表，编号能正确呈现为 1、1.1、1.1.1。
    """
    import uuid
    import zipfile
    from copy import deepcopy

    from docx.oxml.ns import qn
    from lxml import etree

    if template_num_id in cache:
        return cache[template_num_id]

    try:
        with zipfile.ZipFile(str(template)) as archive:
            source_root = etree.fromstring(archive.read("word/numbering.xml"))
    except Exception:
        return None

    source_num = next(
        (
            node
            for node in source_root.findall(qn("w:num"))
            if (node.get(qn("w:numId")) or "").strip() == str(template_num_id)
        ),
        None,
    )
    if source_num is None:
        return None
    abstract_ref = source_num.find(qn("w:abstractNumId"))
    if abstract_ref is None:
        return None
    abstract_id = (abstract_ref.get(qn("w:val")) or "").strip()
    source_abstract = next(
        (
            node
            for node in source_root.findall(qn("w:abstractNum"))
            if (node.get(qn("w:abstractNumId")) or "").strip() == abstract_id
        ),
        None,
    )
    if source_abstract is None:
        return None

    target_root = _numbering_root(document)
    existing_abstract = [
        int(node.get(qn("w:abstractNumId")) or 0)
        for node in target_root.findall(qn("w:abstractNum"))
    ]
    existing_num = [
        int(node.get(qn("w:numId")) or 0)
        for node in target_root.findall(qn("w:num"))
    ]
    new_abstract_id = max(existing_abstract + [-1]) + 1
    new_num_id = max(existing_num + [0]) + 1

    new_abstract = deepcopy(source_abstract)
    new_abstract.set(qn("w:abstractNumId"), str(new_abstract_id))
    nsid = new_abstract.find(qn("w:nsid"))
    if nsid is not None:
        nsid.set(qn("w:val"), uuid.uuid4().hex[:8].upper())
    # 去掉指向模板样式的链接，避免在目标文档中绑定到同名的无关样式
    for tag in ("w:numStyleLink", "w:styleLink"):
        for node in new_abstract.findall(qn(tag)):
            new_abstract.remove(node)
    for level in new_abstract.findall(qn("w:lvl")):
        for node in level.findall(qn("w:pStyle")):
            level.remove(node)
    if not new_abstract.findall(qn("w:lvl")):
        return None

    first_num = target_root.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(new_abstract)
    else:
        cleanup = target_root.find(qn("w:numIdMacAtCleanup"))
        if cleanup is not None:
            cleanup.addprevious(new_abstract)
        else:
            target_root.append(new_abstract)

    new_num = deepcopy(source_num)
    new_num.set(qn("w:numId"), str(new_num_id))
    new_abstract_ref = new_num.find(qn("w:abstractNumId"))
    if new_abstract_ref is not None:
        new_abstract_ref.set(qn("w:val"), str(new_abstract_id))
    cleanup = target_root.find(qn("w:numIdMacAtCleanup"))
    if cleanup is not None:
        cleanup.addprevious(new_num)
    else:
        target_root.append(new_num)

    cache[template_num_id] = new_num_id
    return new_num_id


# 套用模板样式后，用来抵消“从上级样式继承来的格式”的中性值。
# 目标文档的样式往往 basedOn 另一套样式（例如标题 1 basedOn 正文），
# 只改自身属性时，模板没有声明的属性仍会从上级继承，导致“实际格式与模板样式不符”。
_PPR_NEUTRAL = {
    "ind": {
        "firstLine": "0", "firstLineChars": "0",
        "hanging": "0", "hangingChars": "0",
        "left": "0", "leftChars": "0",
        "right": "0", "rightChars": "0",
    },
    "spacing": {
        "before": "0", "beforeLines": "0",
        "after": "0", "afterLines": "0",
        "line": "240", "lineRule": "auto",
    },
    "jc": {"val": "left"},
    "textAlignment": {"val": "auto"},
    "keepNext": {"val": "0"},
    "keepLines": {"val": "0"},
    "pageBreakBefore": {"val": "0"},
    "contextualSpacing": {"val": "0"},
    "mirrorIndents": {"val": "0"},
    "suppressAutoHyphens": {"val": "0"},
    "bidi": {"val": "0"},
    "widowControl": {"val": "1"},
    "snapToGrid": {"val": "1"},
    "wordWrap": {"val": "1"},
    "kinsoku": {"val": "1"},
    "topLinePunct": {"val": "1"},
    "overflowPunct": {"val": "1"},
    "autoSpaceDE": {"val": "1"},
    "autoSpaceDN": {"val": "1"},
    "adjustRightInd": {"val": "1"},
    "tabs": {},
    "pBdr": {},
    "shd": {},
}

_RPR_NEUTRAL = {
    "b": {"val": "0"}, "bCs": {"val": "0"},
    "i": {"val": "0"}, "iCs": {"val": "0"},
    "caps": {"val": "0"}, "smallCaps": {"val": "0"},
    "strike": {"val": "0"}, "dstrike": {"val": "0"},
    "outline": {"val": "0"}, "shadow": {"val": "0"},
    "emboss": {"val": "0"}, "imprint": {"val": "0"},
    "color": {"val": "auto"},
    "u": {"val": "none"},
    "kern": {"val": "0"},
    "spacing": {"val": "0"},
    "w": {"val": "0"},
    "position": {"val": "0"},
    "highlight": {"val": "none"},
    "vertAlign": {"val": "baseline"},
    "rtl": {"val": "0"},
    "cs": {"val": "0"},
    "em": {"val": "none"},
    "shd": {},
    "bdr": {},
}


def _inherited_format_tags(root, style_id: str, table: dict, kind: str) -> set[str]:
    """样式从文档默认值和上级样式继承到的格式属性名集合。"""
    W = f"{{{_W_NS}}}"
    tags: set[str] = set()
    defaults = root.find(f"{W}docDefaults/{W}{kind}Default/{W}{kind}")
    if defaults is not None:
        tags |= {_local_name(child.tag) for child in defaults}
    chain = _style_chain(style_id, table)
    for parent_id in chain[:-1]:
        element = _find_style_element(root, parent_id)
        if element is None:
            continue
        node = element.find(f"{W}{kind}")
        if node is not None:
            tags |= {_local_name(child.tag) for child in node}
    return tags


def _guard_inherited_format(
    container, kind: str, inherited: set[str], keep: set[str]
) -> int:
    """为模板未声明、但会从上级样式继承到的格式补上中性值。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if container is None:
        return 0
    neutrals = _PPR_NEUTRAL if kind == "pPr" else _RPR_NEUTRAL
    order = _PPR_ORDER if kind == "pPr" else _RPR_ORDER
    own = {_local_name(child.tag) for child in container}
    added = 0
    for tag in sorted(inherited - own - keep):
        if tag not in neutrals:
            continue
        element = OxmlElement(f"w:{tag}")
        for key, value in neutrals[tag].items():
            element.set(qn(f"w:{key}"), value)
        _replace_ordered_child(container, element, order)
        added += 1
    return added


def _merge_format_chain(root, style_id: str, table: dict, kind: str):
    """合并“文档默认值 + 样式继承链”，得到样式的完整有效格式元素。"""
    from copy import deepcopy

    W = f"{{{_W_NS}}}"
    base = root.find(f"{W}docDefaults/{W}{kind}Default/{W}{kind}")
    merged = deepcopy(base) if base is not None else None
    for sid in _style_chain(style_id, table):
        element = _find_style_element(root, sid)
        if element is None:
            continue
        node = element.find(f"{W}{kind}")
        if node is None:
            continue
        if merged is None:
            merged = deepcopy(node)
            continue
        for child in node:
            name = _local_name(child.tag)
            if kind == "pPr" and name in ("pStyle", "numPr", "outlineLvl", "rPr", "sectPr"):
                continue
            if kind == "rPr" and name in ("rStyle", "rPrChange"):
                continue
            _replace_ordered_child(
                merged,
                deepcopy(child),
                _PPR_ORDER if kind == "pPr" else _RPR_ORDER,
            )
    return merged


def _format_spec(rpr, ppr) -> dict:
    """把 rPr / pPr 归纳成便于比较的格式描述。"""
    from docx.oxml.ns import qn

    spec: dict = {}

    def _bool(node) -> bool:
        return (node.get(qn("w:val")) or "1") not in ("0", "false", "off")

    if rpr is not None:
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is not None:
            east = fonts.get(qn("w:eastAsia"))
            west = fonts.get(qn("w:ascii")) or fonts.get(qn("w:hAnsi"))
            if east:
                spec["中文字体"] = east
            if west:
                spec["西文字体"] = west
        size = rpr.find(qn("w:sz"))
        if size is not None:
            try:
                spec["字号"] = int(size.get(qn("w:val"))) / 2
            except (TypeError, ValueError):
                pass
        for tag, label in (("b", "加粗"), ("i", "倾斜")):
            node = rpr.find(qn(f"w:{tag}"))
            spec[label] = _bool(node) if node is not None else False
        color = rpr.find(qn("w:color"))
        spec["字体颜色"] = (
            (color.get(qn("w:val")) or "auto") if color is not None else "auto"
        )
        underline = rpr.find(qn("w:u"))
        spec["下划线"] = (
            (underline.get(qn("w:val")) or "single") if underline is not None else "none"
        )

    if ppr is not None:
        jc = ppr.find(qn("w:jc"))
        spec["对齐"] = (jc.get(qn("w:val")) or "left") if jc is not None else "left"
        indent = ppr.find(qn("w:ind"))
        for key, label in (
            ("firstLineChars", "首行缩进(字符)"),
            ("firstLine", "首行缩进"),
            ("leftChars", "左缩进(字符)"),
            ("left", "左缩进"),
            ("right", "右缩进"),
        ):
            value = indent.get(qn(f"w:{key}")) if indent is not None else None
            if value is not None:
                try:
                    spec[label] = int(value)
                except (TypeError, ValueError):
                    pass
        spacing = ppr.find(qn("w:spacing"))
        if spacing is not None:
            for key, label in (
                ("before", "段前"),
                ("after", "段后"),
                ("line", "行距"),
            ):
                value = spacing.get(qn(f"w:{key}"))
                if value is not None:
                    try:
                        spec[label] = int(value)
                    except (TypeError, ValueError):
                        pass
            rule = spacing.get(qn("w:lineRule"))
            if rule:
                spec["行距方式"] = rule
    return spec


def _verify_style_format(
    root, style_table: dict, style_id: str, tpl: dict, themes: dict, char_only: bool = False
) -> list[str]:
    """比较目标样式实际生效的格式与模板样式，返回不一致项说明。

    char_only 为 True 时只比较字符格式（表格专用样式不改段落格式）。
    """
    from copy import deepcopy

    actual_rpr = _merge_format_chain(root, style_id, style_table, "rPr")
    actual_ppr = None if char_only else _merge_format_chain(root, style_id, style_table, "pPr")
    _resolve_theme_fonts_in_rpr(actual_rpr, themes)
    actual = _format_spec(actual_rpr, actual_ppr)
    expected = _format_spec(
        deepcopy(tpl.get("rpr")), None if char_only else deepcopy(tpl.get("ppr"))
    )

    differences: list[str] = []
    # 没有写出的属性按 Word 默认值比较，避免把“未设置”误判为不一致
    defaults = {
        "加粗": False, "倾斜": False, "对齐": "left",
        "字体颜色": "auto", "下划线": "none",
        "首行缩进(字符)": 0, "首行缩进": 0,
        "左缩进(字符)": 0, "左缩进": 0, "右缩进": 0,
        "段前": 0, "段后": 0, "行距": 240, "行距方式": "auto",
    }
    for key in sorted(set(actual) | set(expected)):
        left = actual.get(key)
        right = expected.get(key)
        if left is None:
            left = defaults.get(key, None)
        if right is None:
            right = defaults.get(key, None)
        if left == right:
            continue
        if key == "字号" and (left is None or right is None):
            continue
        differences.append(
            f"{key} 实际 {left if left is not None else '默认'}"
            f" / 模板 {right if right is not None else '默认'}"
        )
    return differences


def apply_template_formats(
    path: Path,
    template: Path,
    selection: dict,
    options: dict,
    chapter_index: int | None,
    output: Path,
    log,
) -> dict:
    """按“文档样式 → 模板样式”的映射套用模板格式，返回处理统计。

    selection 为 {文档 styleId: 模板 styleId}；处理整个文档时修改的是文档
    样式定义，处理单章时只给该章文字直接设置格式。
    """
    import shutil
    import tempfile

    from docx import Document
    from docx.oxml.ns import qn

    temp_dir = Path(tempfile.mkdtemp(prefix="wm_apply_"))
    try:
        docx_path = _convert_to_docx([path], temp_dir, log)[0]
        document = Document(str(docx_path))

        catalog = {
            entry["style_id"]: entry for entry in _template_style_catalog(template)
        }
        style_root = document.styles.element
        style_table, _ = _style_format_table(style_root)

        by_style: dict[str, list] = {}
        table_styles: dict[str, int] = {}
        table_paragraph_ids: set[int] = set()
        table_indexes: set[int] = set()
        for item in _range_paragraph_entries(document, chapter_index):
            paragraph = item["paragraph"]
            if item["table"] is not None:
                table_indexes.add(item["table"])
            if not paragraph.text.strip():
                continue
            style_name = _paragraph_style_name(paragraph)
            if style_name.lower().replace(" ", "").startswith(("toc", "目录")):
                continue
            style_id = _paragraph_style_id(paragraph)
            if style_id in selection:
                by_style.setdefault(style_id, []).append(paragraph)
                if item["table"] is not None:
                    table_styles[style_id] = table_styles.get(style_id, 0) + 1
                    table_paragraph_ids.add(id(paragraph))

        clear_chars = bool(options.get("clear_direct_chars", True))
        clear_paragraph = bool(options.get("clear_direct_paragraph", True))
        add_numbering = bool(options.get("add_numbering", True))
        keep_table = bool(options.get("keep_table_settings", True))
        scoped = chapter_index is not None

        themes = _theme_fonts(template)
        inherited_cache: dict[str, set[str]] = {}

        def inherited(style_id: str, kind: str) -> set[str]:
            key = f"{kind}:{style_id}"
            if key not in inherited_cache:
                inherited_cache[key] = _inherited_format_tags(
                    style_root, style_id, style_table, kind
                )
            return inherited_cache[key]

        numbering_cache: dict[str, int] = {}
        processed = 0
        applied_styles = 0
        numbered_styles = 0
        applied: list[tuple[str, dict, str, bool]] = []

        for style_id, template_id in selection.items():
            paragraphs = by_style.get(style_id) or []
            tpl = catalog.get(template_id)
            if not paragraphs or tpl is None:
                continue
            if tpl.get("rpr") is None and tpl.get("ppr") is None:
                continue
            style_name = _paragraph_style_name(paragraphs[0]) or style_id
            style_element = _find_style_element(style_root, style_id)
            # 只用于表格内文字的样式：按选项只改字符格式，保留单元格/段落设置
            table_only = len(paragraphs) == table_styles.get(style_id, 0)
            char_only = keep_table and table_only

            if scoped or style_element is None:
                # 样式定义是全篇共享的：按单章处理时只改这一章的文字格式
                needs_outline = (
                    not char_only
                    and tpl.get("outline") is not None
                    and _style_chain_outline(style_table, style_id) is None
                )
                inherited_ppr = inherited(style_id, "pPr")
                inherited_rpr = inherited(style_id, "rPr")
                for paragraph in paragraphs:
                    if clear_chars:
                        _clear_direct_char_format(paragraph)
                    for run in paragraph.runs:
                        rpr = run._element.get_or_add_rPr()
                        _apply_template_rpr(rpr, tpl.get("rpr"))
                        # 样式本身不改动，需补中性值抵消原样式继承来的格式
                        _guard_inherited_format(rpr, "rPr", inherited_rpr, set())
                    # 表格内的文字只套用字符格式，保留单元格与段落设置
                    if char_only or (
                        keep_table and id(paragraph) in table_paragraph_ids
                    ):
                        continue
                    if clear_paragraph:
                        _clear_direct_paragraph_format(paragraph)
                    ppr = paragraph._p.get_or_add_pPr()
                    _apply_template_ppr(ppr, tpl.get("ppr"))
                    if needs_outline:
                        _set_ppr_outline(ppr, tpl["outline"])
                    _guard_inherited_format(
                        ppr, "pPr", inherited_ppr, _PPR_KEEP_TAGS
                    )
                if scoped and add_numbering and tpl.get("numbering") and not char_only:
                    log(f"提示：按单章处理时不改动多级列表（「{style_name}」）。")
            else:
                _apply_template_rpr(_ensure_style_rpr(style_element), tpl.get("rpr"))
                if char_only:
                    ppr = style_element.find(qn("w:pPr"))
                    log(
                        f"「{style_name}」只用于表格文字，仅套用字符格式，"
                        "表格样式、边框、底纹与单元格设置保持不变。"
                    )
                else:
                    ppr = _ensure_style_ppr(style_element)
                if ppr is not None and not char_only:
                    _apply_template_ppr(ppr, tpl.get("ppr"))

                # 目标样式继承链没有大纲级别时，按模板补上，保证标题能被目录识别
                if (
                    not char_only
                    and tpl.get("outline") is not None
                    and _style_chain_outline(style_table, style_id) is None
                ):
                    _set_ppr_outline(ppr, tpl["outline"])

                source_numbering = _style_chain_numbering(
                    style_root, _style_chain(style_id, style_table)
                )
                if source_numbering is not None and not _numbering_id_exists(
                    document, source_numbering["num_id"]
                ):
                    log(
                        f"提示：「{style_name}」原有的多级列表定义缺失，"
                        "将按模板重新补全。"
                    )
                    source_numbering = None
                if (
                    add_numbering
                    and not char_only
                    and tpl.get("numbering")
                    and tpl.get("outline") is not None
                    and source_numbering is None
                ):
                    texts = [p.text.strip() for p in paragraphs if p.text.strip()]
                    if any(_RE_MANUAL_NUMBER.match(text) for text in texts):
                        log(
                            f"提示：「{style_name}」的标题文字已带手动编号，"
                            "未自动添加多级列表，避免编号重复。"
                        )
                    else:
                        num_id = _import_numbering_definition(
                            document,
                            template,
                            tpl["numbering"]["num_id"],
                            numbering_cache,
                        )
                        if num_id is not None:
                            _set_style_numbering(
                                ppr, num_id, tpl["numbering"]["ilvl"]
                            )
                            numbered_styles += 1
                            log(f"已为「{style_name}」补上模板的多级列表编号。")
                        else:
                            log(
                                f"提示：未能读取模板的编号定义，"
                                f"未为「{style_name}」补编号。"
                            )

                if clear_chars or clear_paragraph:
                    for paragraph in paragraphs:
                        if clear_chars:
                            _clear_direct_char_format(paragraph)
                        if clear_paragraph and not (
                            keep_table and id(paragraph) in table_paragraph_ids
                        ):
                            _clear_direct_paragraph_format(paragraph)

            processed += len(paragraphs)
            applied_styles += 1
            if not scoped and style_element is not None:
                applied.append((style_id, tpl, style_name, char_only))
            if table_styles.get(style_id) and not char_only:
                log(
                    f"提示：「{style_name}」同时用于正文与表格文字，"
                    "表格内文字会随样式一起更新，表格边框、底纹等设置不变。"
                )
            font = tpl["fmt"].get("font") or "继承字体"
            size = (
                f"{tpl['fmt']['size']:g}pt"
                if tpl["fmt"].get("size")
                else "继承字号"
            )
            log(
                f"「{style_name}」→ 模板「{tpl['name']}」："
                f"{len(paragraphs)} 段（{font} {size}）"
            )

        # 所有样式都改完后，再统一检查继承泄漏并补中性值：
        # 例如正文样式加上首行缩进后，basedOn 正文的标题会继承缩进，
        # 与模板标题样式不符，这里显式写回模板的取值（或默认值）。
        mismatched: list[str] = []
        if applied:
            inherited_cache.clear()
            for style_id, tpl, style_name, char_only in applied:
                style_element = _find_style_element(style_root, style_id)
                if style_element is None:
                    continue
                rpr = style_element.find(qn("w:rPr"))
                _guard_inherited_format(
                    rpr,
                    "rPr",
                    _inherited_format_tags(style_root, style_id, style_table, "rPr"),
                    set(),
                )
                # 表格专用样式不修改段落格式，也就无需补段落中性值
                if char_only:
                    continue
                ppr = style_element.find(qn("w:pPr"))
                _guard_inherited_format(
                    ppr,
                    "pPr",
                    _inherited_format_tags(style_root, style_id, style_table, "pPr"),
                    _PPR_KEEP_TAGS,
                )

            # 校验：逐项比较实际生效格式与模板样式，发现不符时写进日志
            for style_id, tpl, style_name, char_only in applied:
                differences = _verify_style_format(
                    style_root, style_table, style_id, tpl, themes, char_only
                )
                if differences:
                    mismatched.append(style_name)
                    log(
                        f"校验：「{style_name}」与模板「{tpl['name']}」仍有差异 —— "
                        + "；".join(differences)
                    )
            if not mismatched:
                log("校验：所选样式的实际格式与模板样式一致。")

        if keep_table:
            log("表格已识别：表格样式、边框、底纹、合并单元格等设置保持不变，只更新文字格式。")
        else:
            log("表格仅更新文字格式，边框、底纹等表格格式保持不变。")

        output.parent.mkdir(parents=True, exist_ok=True)
        if output.exists():
            output.unlink()
        document.save(str(output))
        log(f"已保存：{output}")
        return {
            "paragraphs": processed,
            "styles": applied_styles,
            "numbered": numbered_styles,
            "mismatched": mismatched,
            "tables": len(table_indexes),
        }
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


# --------------------------------------------------------------------------
# 生成样式参考汇总模板
# --------------------------------------------------------------------------

_CN_FONT_SIZES = {
    42: "初号", 36: "小初", 26: "一号", 24: "小一", 22: "二号", 18: "小二",
    16: "三号", 15: "小三", 14: "四号", 12: "小四", 10.5: "五号", 9: "小五",
    7.5: "六号", 6.5: "小六", 5.5: "七号", 5: "八号",
}

_ALIGN_NAMES = {
    "left": "左对齐",
    "center": "居中",
    "right": "右对齐",
    "both": "两端对齐",
    "justify": "两端对齐",
    "distribute": "分散对齐",
}


def _local_attr(element) -> dict:
    """把带命名空间的属性转成 {短名: 值}。"""
    if element is None:
        return {}
    return {(key.split("}")[-1]) if "}" in key else key: value
            for key, value in element.attrib.items()}


def _find_style_element(root, style_id: str):
    W = f"{{{_W_NS}}}"
    for element in root.findall(f"{W}style"):
        if (element.get(f"{W}styleId") or "").strip() == style_id:
            return element
    return None


def _theme_fonts(reference: Path) -> dict:
    """读取模板主题字体（用于把 majorEastAsia 之类的主题字体还原成实际字体）。"""
    import zipfile

    from lxml import etree

    result: dict = {}
    try:
        with zipfile.ZipFile(str(reference)) as archive:
            name = next(
                (n for n in archive.namelist() if n.startswith("word/theme/") and n.endswith(".xml")),
                None,
            )
            if name is None:
                return result
            root = etree.fromstring(archive.read(name))
    except Exception:
        return result

    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    scheme = root.find(f".//{A}fontScheme")
    if scheme is None:
        return result
    for kind in ("major", "minor"):
        node = scheme.find(f"{A}{kind}Font")
        if node is None:
            continue
        latin = node.find(f"{A}latin")
        east = node.find(f"{A}ea")
        if latin is not None and latin.get("typeface"):
            result[f"{kind}HAnsi"] = latin.get("typeface")
            result[f"{kind}Ascii"] = latin.get("typeface")
        east_name = east.get("typeface") if east is not None else None
        if not east_name:
            # 主题里的中文字体通常写在 <a:font script="Hans" typeface="宋体"/>
            for font in node.findall(f"{A}font"):
                if (font.get("script") or "") in ("Hans", "Hant"):
                    east_name = font.get("typeface")
                    if east_name:
                        break
        if east_name:
            result[f"{kind}EastAsia"] = east_name
    return result


def _resolve_font_name(value: str | None, themes: dict) -> str | None:
    if not value:
        return None
    if value in themes:
        return f"{themes[value]}（主题）"
    return value


def _style_description(root, style_id: str, themes: dict) -> dict:
    """沿 basedOn 链合并出一个样式完整的字体与段落格式描述。"""
    W = f"{{{_W_NS}}}"
    table, _defaults = _style_format_table(root)

    chain: list[str] = []
    seen: set[str] = set()
    current = style_id
    while current and current in table and current not in seen:
        seen.add(current)
        chain.append(current)
        current = table[current]["based"]
    chain.reverse()

    desc: dict = {}
    # 文档默认字体（样式没写字体时用它兜底）
    defaults_rpr = root.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr")
    if defaults_rpr is not None:
        fonts = defaults_rpr.find(f"{W}rFonts")
        if fonts is not None:
            east = _resolve_font_name(
                fonts.get(f"{W}eastAsia") or fonts.get(f"{W}eastAsiaTheme"), themes
            )
            west = _resolve_font_name(
                fonts.get(f"{W}ascii") or fonts.get(f"{W}asciiTheme"), themes
            )
            if east:
                desc["east"] = east
            if west:
                desc["west"] = west
        size = defaults_rpr.find(f"{W}sz")
        if size is not None:
            try:
                desc["size"] = int(size.get(f"{W}val")) / 2
            except (TypeError, ValueError):
                pass

    for sid in chain:
        element = _find_style_element(root, sid)
        if element is None:
            continue
        rpr = element.find(f"{W}rPr")
        if rpr is not None:
            fonts = rpr.find(f"{W}rFonts")
            if fonts is not None:
                east = _resolve_font_name(
                    fonts.get(f"{W}eastAsia") or fonts.get(f"{W}eastAsiaTheme"), themes
                )
                west = _resolve_font_name(
                    fonts.get(f"{W}ascii") or fonts.get(f"{W}asciiTheme"), themes
                )
                if east:
                    desc["east"] = east
                if west:
                    desc["west"] = west
            size = rpr.find(f"{W}sz")
            if size is not None:
                try:
                    desc["size"] = int(size.get(f"{W}val")) / 2
                except (TypeError, ValueError):
                    pass
            bold = rpr.find(f"{W}b")
            if bold is not None:
                desc["bold"] = (bold.get(f"{W}val") or "1") not in ("0", "false", "off")
        ppr = element.find(f"{W}pPr")
        if ppr is not None:
            jc = ppr.find(f"{W}jc")
            if jc is not None:
                desc["align"] = jc.get(f"{W}val")
            indent = ppr.find(f"{W}ind")
            if indent is not None:
                desc["indent"] = _local_attr(indent)
            spacing = ppr.find(f"{W}spacing")
            if spacing is not None:
                desc["spacing"] = _local_attr(spacing)
            if ppr.find(f"{W}numPr") is not None:
                desc["numbered"] = True
            outline = ppr.find(f"{W}outlineLvl")
            if outline is not None:
                try:
                    desc["outline"] = int(outline.get(f"{W}val"))
                except (TypeError, ValueError):
                    pass
    return desc


def _spec_text(desc: dict) -> str:
    """把样式描述转成中文格式说明，例如“中文字体-宋体，字号-小四（12 磅）…”。"""
    parts: list[str] = []
    if desc.get("east"):
        parts.append(f"中文字体-{desc['east']}")
    if desc.get("west") and desc.get("west") != desc.get("east"):
        parts.append(f"西文字体-{desc['west']}")
    size = desc.get("size")
    if size:
        name = _CN_FONT_SIZES.get(size)
        parts.append(f"字号-{name}（{size:g} 磅）" if name else f"字号-{size:g} 磅")
    parts.append("加粗" if desc.get("bold") else "加粗-无")

    align = desc.get("align")
    if align:
        parts.append(_ALIGN_NAMES.get(align, align))

    indent = desc.get("indent") or {}
    first_chars = indent.get("firstLineChars")
    first = indent.get("firstLine")
    hanging_chars = indent.get("hangingChars")
    hanging = indent.get("hanging")
    left = indent.get("left") or indent.get("start")
    if first_chars:
        parts.append(f"首行缩进 {int(first_chars) / 100:g} 字符")
    elif first:
        parts.append(f"首行缩进 {int(first) / 20:g} 磅")
    elif hanging_chars:
        parts.append(f"悬挂缩进 {int(hanging_chars) / 100:g} 字符")
    elif hanging:
        parts.append(f"悬挂缩进 {int(hanging) / 20:g} 磅")
    if left and int(left) > 0:
        parts.append(f"左缩进 {int(left) / 20:g} 磅")

    spacing = desc.get("spacing") or {}
    line = spacing.get("line")
    rule = spacing.get("lineRule") or "auto"
    if line:
        value = int(line)
        if rule == "auto":
            if value == 240:
                parts.append("单倍行距")
            else:
                parts.append(f"{round(value / 240, 2):g} 倍行距")
        else:
            parts.append(f"固定行距 {value / 20:g} 磅")
    for key, label in (("before", "段前"), ("after", "段后")):
        if spacing.get(key):
            parts.append(f"{label} {int(spacing[key]) / 20:g} 磅")
        elif spacing.get(f"{key}Lines"):
            parts.append(f"{label} {int(spacing[f'{key}Lines']) / 100:g} 行")

    if desc.get("numbered"):
        parts.append("带多级列表编号")
    return "，".join(parts) if parts else "继承默认格式"


def _style_examples(document) -> dict:
    """收集参考文档中每种样式的第一段文字，作为汇总模板的示例内容。"""
    examples: dict[str, str] = {}
    for paragraph in _range_paragraphs(document, None):
        text = paragraph.text.strip()
        if not text or len(text) < 2:
            continue
        name = _paragraph_style_name(paragraph)
        if name and name not in examples:
            examples[name] = text[:60] + ("…" if len(text) > 60 else "")
    return examples


def _is_meaningful_style(name: str) -> bool:
    """过滤 Word 自动生成的样式（“样式 标题 1 + …”“列出段落3”等），避免汇总模板过于杂乱。"""
    if not name:
        return False
    if name.startswith("样式 ") or " + " in name:
        return False
    if re.match(r"^(列出段落|无间隔|修订|index\s*\d|heading [6-9]|Normal Indent)", name, re.I):
        return False
    return len(name) <= 24


def _summary_style_entries(root, examples: dict) -> list[dict]:
    """挑选要写进汇总模板的样式。

    参考文档有正文时只收录其中实际用到的样式（标题按级别排在前面）；
    参考文档只有样式库（如 .dotm）时，收录标题样式与自定义样式。
    """
    W = f"{{{_W_NS}}}"
    styles: list[dict] = []
    for element in root.findall(f"{W}style"):
        if (element.get(f"{W}type") or "paragraph") != "paragraph":
            continue
        style_id = (element.get(f"{W}styleId") or "").strip()
        name_element = element.find(f"{W}name")
        name = (
            (name_element.get(f"{W}val") if name_element is not None else "") or style_id
        ).strip()
        if not name or name.lower().replace(" ", "").startswith(("toc", "目录")):
            continue
        ppr = element.find(f"{W}pPr")
        outline = None
        if ppr is not None:
            node = ppr.find(f"{W}outlineLvl")
            if node is not None:
                try:
                    outline = int(node.get(f"{W}val"))
                except (TypeError, ValueError):
                    outline = None
        if outline is None:
            # 没有 outlineLvl 时，按“标题N / heading N”名称判定标题级别
            match = re.fullmatch(r"(?:标题|heading)\s*([1-9])", name, re.I)
            if match:
                outline = int(match.group(1)) - 1
        styles.append(
            {
                "style_id": style_id,
                "name": name,
                "outline": outline,
                "custom": element.get(f"{W}customStyle") == "1",
                "default": element.get(f"{W}default") in ("1", "true", "on"),
            }
        )

    by_name = {}
    for entry in styles:
        by_name.setdefault(entry["name"], entry)

    chosen: list[dict] = []
    used_ids: set[str] = set()

    def add(entry: dict | None):
        if entry is None or entry["style_id"] in used_ids:
            return
        used_ids.add(entry["style_id"])
        chosen.append(entry)

    used = [by_name[name] for name in examples if name in by_name]
    if used:
        headings = sorted(
            [entry for entry in used if entry["outline"] is not None],
            key=lambda entry: (entry["outline"], used.index(entry)),
        )
        others = [entry for entry in used if entry["outline"] is None]
        for entry in headings + others:
            add(entry)
        add(next((entry for entry in styles if entry["default"]), None))
    else:
        for level in range(9):
            level_styles = [
                entry
                for entry in styles
                if entry["outline"] == level and _is_meaningful_style(entry["name"])
            ]
            if not level_styles:
                continue
            cjk = [
                entry
                for entry in level_styles
                if not entry["name"].lower().startswith("heading")
            ]
            add((cjk or level_styles)[0])
        for entry in styles:
            if entry["custom"] and _is_meaningful_style(entry["name"]):
                add(entry)
        add(next((entry for entry in styles if entry["default"]), None))
    return chosen[:40]


def _page_setup_text(document) -> str:
    """把参考文档的纸张与页边距转成说明文字。"""
    try:
        section = document.sections[0]
        width_cm = section.page_width.cm if section.page_width else None
        height_cm = section.page_height.cm if section.page_height else None
        names = {(21.0, 29.7): "A4", (21.59, 27.94): "Letter"}
        paper = None
        if width_cm and height_cm:
            for (w, h), label in names.items():
                if abs(width_cm - w) < 0.3 and abs(height_cm - h) < 0.3:
                    paper = label
                    break
        bits = []
        if paper:
            bits.append(f"页面-{paper}（{width_cm:.1f} × {height_cm:.1f} 厘米）")
        elif width_cm and height_cm:
            bits.append(f"页面-{width_cm:.1f} × {height_cm:.1f} 厘米")
        margins = [
            f"上 {section.top_margin.cm:g}",
            f"下 {section.bottom_margin.cm:g}",
            f"左 {section.left_margin.cm:g}",
            f"右 {section.right_margin.cm:g}",
        ] if all(
            getattr(section, key) is not None
            for key in ("top_margin", "bottom_margin", "left_margin", "right_margin")
        ) else []
        if margins:
            bits.append("页边距-厘米：" + " / ".join(margins))
        return "；".join(bits)
    except Exception:
        return ""


def _set_paragraph_style_id(paragraph, style_id: str) -> None:
    """直接按 styleId 应用样式，避免样式名（本地化名称）匹配失败。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.find(qn("w:pStyle"))
    if existing is not None:
        ppr.remove(existing)
    element = OxmlElement("w:pStyle")
    element.set(qn("w:val"), style_id)
    ppr.insert(0, element)


def _write_style_summary(
    document, reference: Path, entries: list[dict], examples: dict, root, log
) -> int:
    """把样式说明与示例写入已清空正文的文档，返回写入的样式数量。"""
    import time

    from docx.shared import Pt, RGBColor

    themes = _theme_fonts(reference)

    title = document.add_paragraph()
    title.alignment = 1  # 居中
    run = title.add_run("样式参考汇总模板")
    run.bold = True
    run.font.size = Pt(22)

    meta = document.add_paragraph()
    meta.alignment = 1
    meta_run = meta.add_run(
        f"参考文档：{reference.name}　　生成时间：{time.strftime('%Y-%m-%d %H:%M')}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    page_info = _page_setup_text(document)
    intro_lines = [
        "使用说明：每种样式上方是一行灰色小字，说明该样式的字体、字号、缩进、行距等设置；"
        "下方则是按该样式排版的示例文字。",
    ]
    if page_info:
        intro_lines.append(f"版面：{page_info}。")
    intro_lines.append("页眉、页脚请参考参考文档中的实际设置。")
    for line in intro_lines:
        paragraph = document.add_paragraph(line)
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x50, 0x50, 0x50)

    groups = (
        ("一、标题样式", lambda entry: entry["outline"] is not None),
        ("二、正文与段落样式", lambda entry: entry["outline"] is None and not _is_table_style(entry)),
        ("三、表格与图注样式", _is_table_style),
    )

    written = 0
    tables_added = False
    for group_title, predicate in groups:
        members = [entry for entry in entries if predicate(entry)]
        if not members:
            continue

        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(4)
        heading_run = heading.add_run(group_title)
        heading_run.bold = True
        heading_run.font.size = Pt(14)
        heading_run.font.color.rgb = RGBColor(0x1F, 0x3F, 0x66)

        for entry in members:
            desc = _style_description(root, entry["style_id"], themes)
            label = document.add_paragraph()
            label.paragraph_format.space_before = Pt(8)
            label.paragraph_format.space_after = Pt(0)
            display_name = (
                f"{entry['name']}（正文默认）" if entry["default"] else entry["name"]
            )
            label_run = label.add_run(f"【{display_name}】{_spec_text(desc)}")
            label_run.font.size = Pt(9)
            label_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

            sample_text = examples.get(entry["name"]) or _sample_text(entry)
            sample = document.add_paragraph()
            _set_paragraph_style_id(sample, entry["style_id"])
            sample.add_run(sample_text)
            written += 1

        if group_title.startswith("三") and not tables_added:
            _add_sample_table(document, root, entries, log)
            tables_added = True

    # 表格示例：即使模板里没有“表字体”样式也给一个
    if not tables_added:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        run = heading.add_run("三、表格与图注样式")
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1F, 0x3F, 0x66)
        _add_sample_table(document, root, entries, log)
    return written


def _is_table_style(entry: dict) -> bool:
    name = entry["name"]
    return ("表" in name) or name.lower() in ("caption", "caption text") or "图注" in name


def _sample_text(entry: dict) -> str:
    if entry["outline"] is not None:
        return f"这是{entry['name']}样式示例，用于查看标题的字体、字号与编号效果。"
    return (
        f"这是“{entry['name']}”样式的示例文字，用于查看该样式的字体、字号、缩进与行距效果。"
        "中文与 English 及 123 数字混排时字体的表现也一并展示。"
    )


def _add_sample_table(document, root, entries: list[dict], log) -> None:
    """在汇总模板中插入一个示例表格（使用模板的表格样式与“表字体”段落样式）。"""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt, RGBColor

    W = f"{{{_W_NS}}}"
    table_style_name = None
    for element in root.findall(f"{W}style"):
        if (element.get(f"{W}type") or "") != "table":
            continue
        name_element = element.find(f"{W}name")
        name = (name_element.get(f"{W}val") if name_element is not None else "") or ""
        if name == "Table Grid":
            table_style_name = name
            break
        if table_style_name is None:
            table_style_name = name

    cell_style_id = next(
        (entry["style_id"] for entry in entries if "表字体" in entry["name"]), None
    )

    label = document.add_paragraph()
    label.paragraph_format.space_before = Pt(8)
    label_run = label.add_run("【表格示例】表内文字使用上面的表格文字样式，表格自身沿用模板的表格样式")
    label_run.font.size = Pt(9)
    label_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    table = document.add_table(rows=3, cols=3)
    if table_style_name:
        try:
            table.style = table_style_name
        except Exception:
            pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("寄存器", "地址", "说明")
    rows = (
        ("ADC_CR", "0x4001_2460", "控制寄存器"),
        ("ADC_TGSR", "0x4001_2450", "采样组配置"),
    )
    for column, text in enumerate(headers):
        cell = table.rows[0].cells[column]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        if cell_style_id:
            _set_paragraph_style_id(paragraph, cell_style_id)
        paragraph.add_run(text)
    for row_index, values in enumerate(rows, start=1):
        for column, text in enumerate(values):
            cell = table.rows[row_index].cells[column]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            if cell_style_id:
                _set_paragraph_style_id(paragraph, cell_style_id)
            paragraph.add_run(text)
    log("已在汇总模板中加入表格示例。")


def _load_reference_base(reference: Path, temp_dir: Path, log):
    """打开参考文档作为汇总模板的底稿，返回（文档对象，styles.xml 根元素）。"""
    from docx import Document

    suffix = reference.suffix.lower()
    if suffix in (".docx", ".docm"):
        document = Document(str(reference))
        return document, document.styles.element

    # .dotm / .dotx / .doc / .rtf：用 Word 转成临时 docx，保留样式与页面设置
    try:
        converted = _convert_to_docx([reference], temp_dir, log)[0]
        document = Document(str(converted))
        return document, document.styles.element
    except Exception as exc:
        log(f"未能直接打开模板（{exc}），改用导入样式的方式生成。")

    document = Document()
    _import_styles_python_docx(document, reference, log)
    return document, _read_styles_xml(reference)


def _clear_body_keep_section(document) -> None:
    """清空正文，但保留页面设置以及参考文档的页眉页脚。"""
    from copy import deepcopy

    from docx.oxml.ns import qn

    body = document.element.body
    final_sect = body.find(qn("w:sectPr"))

    # 参考文档可能分多节，页眉页脚引用常写在靠前的分节符里，这里先取出来
    first_sect = None
    for element in body.iter(qn("w:sectPr")):
        first_sect = element
        break

    preserved = []
    if first_sect is not None and first_sect is not final_sect:
        for child in first_sect:
            if child.tag in (
                qn("w:headerReference"),
                qn("w:footerReference"),
                qn("w:titlePg"),
            ):
                preserved.append(deepcopy(child))

    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    if final_sect is not None and preserved:
        for item in preserved:
            if item.tag == qn("w:titlePg"):
                continue
            kind = item.get(qn("w:type"))
            for existing in final_sect.findall(item.tag):
                if existing.get(qn("w:type")) == kind:
                    final_sect.remove(existing)
            final_sect.append(item)
        if any(item.tag == qn("w:titlePg") for item in preserved):
            if final_sect.find(qn("w:titlePg")) is None:
                final_sect.append(
                    next(item for item in preserved if item.tag == qn("w:titlePg"))
                )


def build_style_summary(reference: Path, output: Path, log) -> Path:
    """以参考文档为样板，生成带样式说明与示例的汇总模板。"""
    import shutil
    import tempfile

    from docx import Document  # noqa: F401 - 确保依赖可用

    temp_dir = Path(tempfile.mkdtemp(prefix="wm_summary_"))
    try:
        log(f"正在读取参考文档：{reference.name}")
        document, root = _load_reference_base(reference, temp_dir, log)

        examples = _style_examples(document)
        entries = _summary_style_entries(root, examples)
        if not entries:
            raise RuntimeError("参考文档中没有可用的段落样式。")

        _clear_body_keep_section(document)

        written = _write_style_summary(document, reference, entries, examples, root, log)

        output.parent.mkdir(parents=True, exist_ok=True)
        if output.exists():
            output.unlink()
        document.save(str(output))
        log(f"汇总模板已生成：{output}（共 {written} 种样式）")
        return output
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def merge_documents(files: list[Path], output: Path, options: MergeOptions, log) -> str:
    """合并入口，返回实际使用的引擎名称。

    默认用 docxcompose 按顺序拼接文档内容，并以第一个文档为主文档：
    主文档的样式库、多级列表编号和页面设置原样保留，后续文档依次插入，
    同名但定义不同的样式随所在文档各带一份副本，不改变源文件。
    非 .docx 文件先用 Word 转换为临时 .docx。
    """
    if options.unify_styles and options.style_source:
        return _merge_with_unified_styles(files, output, options, log)

    import shutil
    import tempfile

    temp_dir: Path | None = None
    try:
        if any(path.suffix.lower() != ".docx" for path in files):
            temp_dir = Path(tempfile.mkdtemp(prefix="wm_merge_"))
            merge_files = _convert_to_docx(files, temp_dir, log)
        else:
            merge_files = files
        _merge_with_docxcompose(merge_files, output, options, log)
        if options.make_toc:
            _update_fields_with_word(output, log)
        return "docxcompose"
    except Exception as exc:
        log(f"使用 docxcompose 合并失败：{exc}")
        try:
            if output.exists():
                output.unlink()
        except Exception:
            pass
        log("改用 Microsoft Word 合并（部分格式可能略有差异）…")
        _merge_with_word(files, output, options, log)
        return "Microsoft Word"
    finally:
        if temp_dir is not None:
            shutil.rmtree(temp_dir, ignore_errors=True)


# --------------------------------------------------------------------------
# 图形界面
# --------------------------------------------------------------------------

class MergeApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.items: list[DocumentItem] = []
        self.messages: queue.Queue = queue.Queue()
        self.busy = False
        self.output_auto = True
        self._setting_output = False
        self._open_dir: Path | None = None

        root.title("Word 文档按章节合并")
        root.geometry("1020x940")
        root.minsize(980, 780)

        self.auto_sort = tk.BooleanVar(value=True)
        self.page_break = tk.BooleanVar(value=True)
        self.make_toc = tk.BooleanVar(value=True)
        self.toc_title = tk.StringVar(value="目录")
        self.style_source = tk.StringVar(value="")
        self.prune_styles = tk.BooleanVar(value=False)
        self.unify_styles = tk.BooleanVar(value=False)
        self.standalone_import = tk.BooleanVar(value=True)
        self.standalone_prune = tk.BooleanVar(value=True)
        self.marker_start = tk.StringVar(value=DEFAULT_MARKER_START)
        self.marker_end = tk.StringVar(value=DEFAULT_MARKER_END)
        self.save_source = tk.BooleanVar(value=False)
        self.output_var = tk.StringVar(value="")
        self.status_var = tk.StringVar(value="就绪")

        self._build_ui()
        self.output_var.trace_add("write", self._on_output_edited)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self._poll_messages()

    # ---------------- UI ----------------
    def _build_ui(self):
        style = ttk.Style(self.root)
        try:
            style.theme_use("vista")
        except Exception:
            pass

        outer = ttk.Frame(self.root, padding=12)
        outer.pack(fill="both", expand=True)

        header = ttk.Frame(outer)
        header.pack(fill="x")
        ttk.Label(
            header, text="按章节合并 Word 文档", font=("Microsoft YaHei UI", 14, "bold")
        ).pack(anchor="w")
        hint = (
            "添加或拖入多份文档，每份文档作为一章；仅按章号排序并拼接，"
            "不修改章节文档，也不改动原有格式和样式库。"
        )
        if not HAS_DND:
            hint += "（提示：安装 tkinterdnd2 后可拖拽文件）"
        ttk.Label(header, text=hint, foreground="#555").pack(anchor="w", pady=(4, 0))

        toolbar = ttk.Frame(outer)
        toolbar.pack(fill="x", pady=(10, 6))
        ttk.Button(toolbar, text="添加文件…", command=self.add_files).pack(side="left")
        ttk.Button(toolbar, text="添加文件夹…", command=self.add_folder).pack(
            side="left", padx=(6, 0)
        )
        ttk.Button(
            toolbar, text="合并/导入单章节…", command=self.import_chapter_into_document
        ).pack(side="left", padx=(6, 0))
        ttk.Button(
            toolbar, text="合并到文档…", command=self.merge_into_document
        ).pack(side="left", padx=(6, 0))
        self.split_button = ttk.Button(
            toolbar, text="拆分总文档…", command=self.split_document
        )
        self.split_button.pack(side="left", padx=(6, 0))
        self.format_button = ttk.Button(
            toolbar, text="套用模板样式…", command=self.apply_template_to_file
        )
        self.format_button.pack(side="left", padx=(6, 0))
        self.summary_button = ttk.Button(
            toolbar, text="生成汇总模板…", command=self.build_summary_template
        )
        self.summary_button.pack(side="left", padx=(6, 0))
        ttk.Separator(toolbar, orient="vertical").pack(
            side="left", fill="y", padx=10
        )
        ttk.Button(toolbar, text="上移", command=lambda: self.move(-1)).pack(side="left")
        ttk.Button(toolbar, text="下移", command=lambda: self.move(1)).pack(
            side="left", padx=(6, 0)
        )
        ttk.Button(toolbar, text="移除", command=self.remove_selected).pack(
            side="left", padx=(6, 0)
        )
        ttk.Button(toolbar, text="清空", command=self.clear_all).pack(
            side="left", padx=(6, 0)
        )

        list_frame = ttk.Frame(outer)
        list_frame.pack(fill="both", expand=True)
        self.listbox = tk.Listbox(
            list_frame, selectmode="extended", activestyle="none", height=10
        )
        scrollbar = ttk.Scrollbar(
            list_frame, orient="vertical", command=self.listbox.yview
        )
        self.listbox.configure(yscrollcommand=scrollbar.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        if HAS_DND:
            for widget in (self.root, self.listbox):
                try:
                    widget.drop_target_register(DND_FILES)
                    widget.dnd_bind("<<Drop>>", self._on_drop)
                except Exception:
                    pass

        options = ttk.LabelFrame(outer, text="合并选项", padding=8)
        options.pack(fill="x", pady=(10, 6))
        ttk.Checkbutton(
            options,
            text="仅按章号自动排序",
            variable=self.auto_sort,
            command=self._apply_sort,
        ).grid(row=0, column=0, sticky="w", padx=(0, 18))
        ttk.Checkbutton(
            options, text="每章另起一页", variable=self.page_break
        ).grid(row=0, column=1, sticky="w", padx=(0, 18))
        ttk.Checkbutton(
            options, text="自动生成目录", variable=self.make_toc
        ).grid(row=0, column=2, sticky="w", padx=(0, 18))
        ttk.Label(options, text="目录标题：").grid(row=0, column=3, sticky="e")
        ttk.Entry(options, textvariable=self.toc_title, width=12).grid(
            row=0, column=4, sticky="w"
        )

        style_row = ttk.Frame(options)
        style_row.grid(row=1, column=0, columnspan=5, sticky="we", pady=(8, 0))
        ttk.Label(style_row, text="样式模板（.dotm/.dotx）：").pack(side="left")
        ttk.Entry(style_row, textvariable=self.style_source).pack(
            side="left", fill="x", expand=True
        )
        ttk.Button(style_row, text="选择模板…", command=self.choose_style_source).pack(
            side="left", padx=(6, 0)
        )
        ttk.Button(
            style_row, text="清除", command=lambda: self.style_source.set("")
        ).pack(side="left", padx=(6, 0))
        ttk.Label(
            options,
            text="默认仅按章节拼接：以第一个文档为基准（样式、多级列表、页面设置不变），"
            "不修改任何章节文档，不导入/删除样式。",
            foreground="#555",
        ).grid(row=2, column=0, columnspan=5, sticky="w", pady=(6, 0))
        ttk.Checkbutton(
            options,
            text="删除模板之外的样式（需手动勾选；未选模板时只删除正文未使用的样式；同名重复的未使用样式一并清理）",
            variable=self.prune_styles,
        ).grid(row=3, column=0, columnspan=5, sticky="w", pady=(2, 0))
        self.unify_check = ttk.Checkbutton(
            options,
            text="合并前先逐个文件统一为模板样式（需手动勾选；默认只处理临时副本，不影响源文件）",
            variable=self.unify_styles,
        )
        self.unify_check.grid(row=4, column=0, columnspan=5, sticky="w", pady=(4, 0))
        self.save_source_check = ttk.Checkbutton(
            options,
            text="写回源文件（默认关闭：合并与样式处理只操作临时副本/新文件夹，勾选后才会修改源文件并自动备份）",
            variable=self.save_source,
        )
        self.save_source_check.grid(
            row=5, column=0, columnspan=5, sticky="w", pady=(4, 0)
        )
        self.style_source.trace_add("write", lambda *_: self._sync_style_options())
        self._sync_style_options()

        solo = ttk.LabelFrame(outer, text="单独处理样式（不合并）", padding=8)
        solo.pack(fill="x", pady=(0, 6))
        solo_row = ttk.Frame(solo)
        solo_row.pack(fill="x")
        ttk.Checkbutton(
            solo_row, text="导入样式模板", variable=self.standalone_import
        ).pack(side="left")
        ttk.Checkbutton(
            solo_row,
            text="删除多余样式（选了模板时删除模板之外的样式；同名重复的未使用样式一并清理）",
            variable=self.standalone_prune,
        ).pack(side="left", padx=(14, 0))
        self.style_button = ttk.Button(
            solo_row, text="处理选中文件…", command=self.start_style_only
        )
        self.style_button.pack(side="right")
        ttk.Label(
            solo,
            text="未选中文件时处理列表中的全部文件；默认不修改源文件（输出到新文件夹），勾选「写回源文件」后才会备份并修改原文件。",
            foreground="#555",
        ).pack(anchor="w", pady=(4, 0))

        marked = ttk.LabelFrame(outer, text="删除标记区间（不合并）", padding=8)
        marked.pack(fill="x", pady=(0, 6))
        marked_row = ttk.Frame(marked)
        marked_row.pack(fill="x")
        ttk.Label(marked_row, text="起始标记：").pack(side="left")
        ttk.Entry(marked_row, textvariable=self.marker_start, width=16).pack(
            side="left"
        )
        ttk.Label(marked_row, text="结束标记：").pack(side="left", padx=(12, 0))
        ttk.Entry(marked_row, textvariable=self.marker_end, width=16).pack(side="left")
        self.mark_button = ttk.Button(
            marked_row,
            text="删除选中文件中的标记区间…",
            command=self.start_delete_marked,
        )
        self.mark_button.pack(side="right")
        ttk.Label(
            marked,
            text="删除两个标记之间的全部内容及标记本身，并清理删除位置两侧遗留的空段落与空格；"
            "其余内容的文字、样式、编号均保持不变。",
            foreground="#555",
        ).pack(anchor="w", pady=(4, 0))

        output_frame = ttk.Frame(outer)
        output_frame.pack(fill="x", pady=(2, 6))
        ttk.Label(output_frame, text="输出文件：").pack(side="left")
        self.output_entry = ttk.Entry(output_frame, textvariable=self.output_var)
        self.output_entry.pack(side="left", fill="x", expand=True)
        ttk.Button(output_frame, text="浏览…", command=self.choose_output).pack(
            side="left", padx=(6, 0)
        )

        action = ttk.Frame(outer)
        action.pack(fill="x", pady=(2, 6))
        self.merge_button = ttk.Button(
            action, text="开始合并", command=self.start_merge
        )
        self.merge_button.pack(side="left")
        self.progress = ttk.Progressbar(action, mode="indeterminate", length=180)
        self.progress.pack(side="left", padx=10)
        ttk.Label(action, textvariable=self.status_var, foreground="#555").pack(
            side="left"
        )
        self.open_button = ttk.Button(
            action, text="打开输出位置", command=self.open_output_folder, state="disabled"
        )
        self.open_button.pack(side="right")

        log_frame = ttk.LabelFrame(outer, text="日志", padding=4)
        log_frame.pack(fill="both", expand=False)
        self.log_text = tk.Text(log_frame, height=8, wrap="word", state="disabled")
        log_scroll = ttk.Scrollbar(
            log_frame, orient="vertical", command=self.log_text.yview
        )
        self.log_text.configure(yscrollcommand=log_scroll.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        log_scroll.pack(side="right", fill="y")

        self._refresh_output_default()

    # ---------------- 文件列表 ----------------
    def _on_drop(self, event):
        try:
            paths = self.root.tk.splitlist(event.data)
        except Exception:
            paths = str(event.data).split()
        for raw in paths:
            path = Path(raw)
            if path.is_dir():
                self._add_directory(path)
            else:
                self._add_path(path)

    def add_files(self):
        paths = filedialog.askopenfilenames(
            title="选择要合并的 Word 文档",
            filetypes=[
                ("Word 文档", "*.docx *.doc *.docm *.rtf"),
                ("所有文件", "*.*"),
            ],
        )
        for raw in paths:
            self._add_path(Path(raw))

    def add_folder(self):
        folder = filedialog.askdirectory(title="选择包含 Word 文档的文件夹")
        if folder:
            self._add_directory(Path(folder))

    # ---------------- 合并 / 导入单章节 ----------------
    def import_chapter_into_document(self):
        """先选单章节、再选总文档；总文档已有同章节时可选择替换。"""
        if self.busy:
            return

        chapter = self._ask_word_file("第 1 步：选择要导入的单个章节文档")
        if chapter is None:
            return
        target = self._ask_word_file("第 2 步：选择要合并到的总文档")
        if target is None:
            return
        try:
            same_file = chapter.resolve() == target.resolve()
        except Exception:
            same_file = chapter == target
        if same_file:
            messagebox.showwarning(
                "文件重复", "章节文档不能与总文档是同一个文件，请重新选择。"
            )
            return

        self.status_var.set("正在分析总文档的章节…")
        self.root.update_idletasks()
        try:
            info = analyze_chapter_import(chapter, target, self._log)
        except Exception as exc:
            messagebox.showerror("分析失败", str(exc))
            return
        self.status_var.set("就绪")

        action = self._ask_import_chapter_action(chapter, target, info)
        if action is None:
            return

        save_name = filedialog.asksaveasfilename(
            title="保存合并结果",
            defaultextension=".docx",
            initialdir=str(target.parent),
            initialfile=f"{target.stem}_更新.docx",
            filetypes=[("Word 文档", "*.docx")],
        )
        if not save_name:
            return
        output = Path(save_name).with_suffix(".docx")
        for source, label in ((target, "总文档"), (chapter, "章节文档")):
            try:
                if output.resolve() == source.resolve():
                    messagebox.showwarning(
                        "路径冲突",
                        f"输出文件不能与{label}相同，请另选一个文件名。",
                    )
                    return
            except Exception:
                pass
        if output.exists():
            if not messagebox.askyesno(
                "文件已存在", f"输出文件已存在，要覆盖吗？\n{output}"
            ):
                return

        options = MergeOptions(
            make_toc=False,
            page_break=self.page_break.get(),
        )

        self.busy = True
        self.merge_button.configure(state="disabled")
        self.style_button.configure(state="disabled")
        self.format_button.configure(state="disabled")
        self.summary_button.configure(state="disabled")
        self.split_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.progress.start(12)
        self.status_var.set("正在导入单章节…")
        self._append_log("=" * 60)
        self._append_log(
            f"单章节{'替换' if action == 'replace' else '导入'}：{chapter.name}"
            f" -> {target.name}（输出 {output}）"
        )

        worker = threading.Thread(
            target=self._into_worker,
            args=(target, [chapter], output, action, options),
            daemon=True,
        )
        worker.start()

    def _ask_word_file(self, title: str) -> Path | None:
        """弹窗选择一个 Word 文档，校验后返回绝对路径。"""
        filename = filedialog.askopenfilename(
            title=title,
            filetypes=[
                ("Word 文档", "*.docx *.doc *.docm *.rtf"),
                ("所有文件", "*.*"),
            ],
        )
        if not filename:
            return None
        path = Path(filename)
        try:
            path = path.resolve()
        except Exception:
            pass
        if path.name.startswith("~$") or not path.is_file():
            messagebox.showwarning("文件无效", f"无法读取该文件：\n{path}")
            return None
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            messagebox.showwarning(
                "文件类型不支持", "仅支持 .docx / .doc / .docm / .rtf 文档。"
            )
            return None
        return path

    def _ask_import_chapter_action(self, chapter: Path, target: Path, info: dict):
        """展示总文档的章节，确认是替换已有同章节还是插入新章节。"""
        dialog = tk.Toplevel(self.root)
        dialog.title("合并 / 导入单章节")
        dialog.transient(self.root)
        dialog.resizable(False, False)

        frame = ttk.Frame(dialog, padding=14)
        frame.pack(fill="both", expand=True)
        ttk.Label(
            frame,
            text=f"章节文档：{chapter.name}",
            font=("Microsoft YaHei UI", 10, "bold"),
        ).pack(anchor="w")
        ttk.Label(
            frame,
            text=f"识别章号：{_chapter_label(info['number'])}"
            f"    标题：{info['title'] or '（未识别）'}",
            foreground="#555",
        ).pack(anchor="w", pady=(2, 0))
        ttk.Label(
            frame,
            text=f"总文档：{target.name}（共 {len(info['spans'])} 章）",
            font=("Microsoft YaHei UI", 10, "bold"),
        ).pack(anchor="w", pady=(10, 0))

        spans = info["spans"]
        match_index = info["match_index"]
        if match_index is not None:
            match_heading = spans[match_index]["heading"]
            ttk.Label(
                frame,
                text=f"总文档中已有同章节：【{match_heading}】，可替换该章节。",
                foreground="#1f6f3f",
            ).pack(anchor="w", pady=(6, 4))
        elif spans:
            ttk.Label(
                frame,
                text=f"总文档中没有找到【{_chapter_label(info['number'])}】，将作为新章节插入。",
                foreground="#8a4b00",
            ).pack(anchor="w", pady=(6, 4))
        else:
            ttk.Label(
                frame,
                text="总文档中没有找到章节标题，本章将追加到文末。",
                foreground="#8a4b00",
            ).pack(anchor="w", pady=(6, 4))

        if spans:
            listbox = tk.Listbox(
                frame, height=min(8, max(4, len(spans))), width=70, activestyle="none"
            )
            listbox.pack(fill="x", pady=(0, 10))
            for index, span in enumerate(spans, start=1):
                listbox.insert("end", f"{index}. {span['heading']}")
            if match_index is not None:
                listbox.selection_set(match_index)
                listbox.see(match_index)
            listbox.configure(state="disabled")

        ttk.Label(frame, text="请选择处理方式：").pack(anchor="w")
        action = tk.StringVar(value="replace" if match_index is not None else "insert")
        if match_index is not None:
            ttk.Radiobutton(
                frame,
                text=f"替换：用本章替换【{match_heading}】"
                "（只替换这一章，总文档其余内容与格式保持不变）",
                variable=action,
                value="replace",
            ).pack(anchor="w", pady=(6, 0))
        ttk.Radiobutton(
            frame,
            text="插入：按章号插入到总文档的对应位置（原有章节保持不变）",
            variable=action,
            value="insert",
        ).pack(anchor="w", pady=(4, 0))

        buttons = ttk.Frame(frame)
        buttons.pack(fill="x", pady=(14, 0))
        result: dict[str, str | None] = {"value": None}

        def confirm():
            result["value"] = action.get()
            dialog.destroy()

        ttk.Button(buttons, text="取消", command=dialog.destroy).pack(side="right")
        ttk.Button(buttons, text="确定", command=confirm).pack(
            side="right", padx=(0, 6)
        )
        dialog.bind("<Return>", lambda _event: confirm())
        dialog.bind("<Escape>", lambda _event: dialog.destroy())
        dialog.update_idletasks()
        try:
            x = (
                self.root.winfo_rootx()
                + (self.root.winfo_width() - dialog.winfo_width()) // 2
            )
            y = (
                self.root.winfo_rooty()
                + (self.root.winfo_height() - dialog.winfo_height()) // 3
            )
            dialog.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        except Exception:
            pass
        dialog.grab_set()
        self.root.wait_window(dialog)
        return result["value"]

    # ---------------- 合并到已有文档 ----------------
    def merge_into_document(self):
        """把选中的章节合并（插入）或替换到另一个已有文档中。"""
        if self.busy:
            return
        selection = list(self.listbox.curselection())
        items = [self.items[i] for i in selection] if selection else list(self.items)
        if not items:
            messagebox.showwarning("提示", "请先添加要合并的章节文档。")
            return

        filename = filedialog.askopenfilename(
            title="选择要合并到的文档",
            filetypes=[
                ("Word 文档", "*.docx *.doc *.docm *.rtf"),
                ("所有文件", "*.*"),
            ],
        )
        if not filename:
            return
        target = Path(filename)
        try:
            target = target.resolve()
        except Exception:
            pass
        if not target.is_file() or target.name.startswith("~$"):
            messagebox.showwarning("目标文档无效", f"无法读取该文件：\n{target}")
            return

        chapters: list[Path] = []
        for item in items:
            try:
                if item.path.resolve() == target:
                    continue
            except Exception:
                pass
            chapters.append(item.path)
        if not chapters:
            messagebox.showwarning(
                "提示", "目标文档不能同时作为要合并的章节，请重新选择。"
            )
            return

        mode = self._ask_into_action(target, len(chapters))
        if mode is None:
            return

        default_name = f"{target.stem}_更新.docx"
        save_name = filedialog.asksaveasfilename(
            title="保存合并结果",
            defaultextension=".docx",
            initialdir=str(target.parent),
            initialfile=default_name,
            filetypes=[("Word 文档", "*.docx")],
        )
        if not save_name:
            return
        output = Path(save_name).with_suffix(".docx")
        try:
            if output.resolve() == target.resolve():
                messagebox.showwarning(
                    "路径冲突", "输出文件不能与目标文档相同，请另选一个文件名。"
                )
                return
        except Exception:
            pass
        for chapter in chapters:
            try:
                if output.resolve() == chapter.resolve():
                    messagebox.showwarning(
                        "路径冲突", f"输出文件不能与章节文档相同：\n{chapter.name}"
                    )
                    return
            except Exception:
                pass
        if output.exists():
            if not messagebox.askyesno(
                "文件已存在", f"输出文件已存在，要覆盖吗？\n{output}"
            ):
                return

        options = MergeOptions(
            make_toc=False,
            page_break=self.page_break.get(),
        )

        self.busy = True
        self.merge_button.configure(state="disabled")
        self.style_button.configure(state="disabled")
        self.format_button.configure(state="disabled")
        self.summary_button.configure(state="disabled")
        self.split_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.progress.start(12)
        self.status_var.set("正在合并到文档…")
        self._append_log("=" * 60)
        self._append_log(
            f"把 {len(chapters)} 个章节{'替换到' if mode == 'replace' else '合并到'}"
            f" {target.name} -> {output}"
        )

        worker = threading.Thread(
            target=self._into_worker,
            args=(target, chapters, output, mode, options),
            daemon=True,
        )
        worker.start()

    def _ask_into_action(self, target: Path, count: int) -> str | None:
        """询问：把章节合并（插入）到目标文档，还是替换目标文档中的同名章节。"""
        dialog = tk.Toplevel(self.root)
        dialog.title("合并到文档")
        dialog.transient(self.root)
        dialog.resizable(False, False)

        frame = ttk.Frame(dialog, padding=14)
        frame.pack(fill="both", expand=True)
        ttk.Label(
            frame,
            text=f"目标文档：{target.name}",
            font=("Microsoft YaHei UI", 10, "bold"),
        ).pack(anchor="w")
        ttk.Label(
            frame,
            text=f"将处理 {count} 个章节文档，请选择方式：",
            foreground="#555",
        ).pack(anchor="w", pady=(2, 10))

        action = tk.StringVar(value="merge")
        ttk.Radiobutton(
            frame,
            text="合并：按章号把章节插入目标文档的对应位置，原有内容保持不变",
            variable=action,
            value="merge",
        ).pack(anchor="w")
        ttk.Radiobutton(
            frame,
            text="替换：按章号替换目标文档中的同名章节（找不到时按章号插入）",
            variable=action,
            value="replace",
        ).pack(anchor="w", pady=(4, 0))
        ttk.Label(
            frame,
            text="结果会另存为新文件，目标文档本身不会被修改。",
            foreground="#555",
        ).pack(anchor="w", pady=(10, 0))

        buttons = ttk.Frame(frame)
        buttons.pack(fill="x", pady=(14, 0))
        result: dict[str, str | None] = {"value": None}

        def confirm():
            result["value"] = action.get()
            dialog.destroy()

        ttk.Button(buttons, text="取消", command=dialog.destroy).pack(side="right")
        ttk.Button(buttons, text="确定", command=confirm).pack(
            side="right", padx=(0, 6)
        )
        dialog.bind("<Return>", lambda _event: confirm())
        dialog.bind("<Escape>", lambda _event: dialog.destroy())
        dialog.update_idletasks()
        try:
            x = (
                self.root.winfo_rootx()
                + (self.root.winfo_width() - dialog.winfo_width()) // 2
            )
            y = (
                self.root.winfo_rooty()
                + (self.root.winfo_height() - dialog.winfo_height()) // 3
            )
            dialog.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        except Exception:
            pass
        dialog.grab_set()
        self.root.wait_window(dialog)
        return result["value"]

    def _into_worker(
        self,
        target: Path,
        chapters: list[Path],
        output: Path,
        mode: str,
        options: MergeOptions,
    ):
        try:
            engine = merge_chapters_into_document(
                target, chapters, output, mode, options, self._log
            )
            self.messages.put(("into_done", (output, engine)))
        except Exception as exc:
            self.messages.put(("error", (exc, traceback.format_exc())))

    def _on_into_done(self, output: Path, engine: str):
        self._finish()
        self.status_var.set(f"已完成（{engine}）")
        self._append_log(f"已生成：{output}")
        self._open_dir = output.parent
        self.open_button.configure(state="normal")
        if messagebox.askyesno(
            "处理完成", f"文档已生成：\n{output}\n\n是否立即打开？"
        ):
            try:
                os.startfile(str(output))
            except Exception as exc:
                messagebox.showerror("无法打开文档", str(exc))

    # ---------------- 按章节拆分总文档 ----------------
    def split_document(self):
        """把一本总文档按章节拆分为多个单独文档，格式保持不变。"""
        if self.busy:
            return

        filename = filedialog.askopenfilename(
            title="选择要按章节拆分的总文档",
            filetypes=[
                ("Word 文档", "*.docx *.doc *.docm *.rtf"),
                ("所有文件", "*.*"),
            ],
        )
        if not filename:
            return
        path = Path(filename)
        try:
            path = path.resolve()
        except Exception:
            pass
        if not path.is_file() or path.name.startswith("~$"):
            messagebox.showwarning("文件无效", f"无法读取该文件：\n{path}")
            return
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            messagebox.showwarning(
                "文件类型不支持", "仅支持 .docx / .doc / .docm / .rtf 文档。"
            )
            return

        self.status_var.set("正在读取章节…")
        self.root.update_idletasks()
        temp_dir = Path(tempfile.mkdtemp(prefix="wm_split_preview_"))
        try:
            docx_path = _convert_to_docx([path], temp_dir, self._log)[0]
            from docx import Document

            spans = _document_spans(Document(str(docx_path)))
        except Exception as exc:
            messagebox.showerror("读取失败", str(exc))
            return
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

        if not spans:
            messagebox.showwarning(
                "未找到章节",
                "文档中未找到章节标题（优先使用一级标题，没有一级标题时使用其他级别标题），"
                "无法拆分。",
            )
            return

        names = "\n".join(f"  {span.heading}" for span in spans[:8])
        if len(spans) > 8:
            names += f"\n  …等 {len(spans)} 章"
        if not messagebox.askyesno(
            "按章节拆分",
            f"识别到 {len(spans)} 个章节：\n\n{names}\n\n"
            f"将拆分为 {len(spans)} 个单独文档，是否继续？",
        ):
            return

        folder = filedialog.askdirectory(
            title="选择拆分后文档的保存文件夹", initialdir=str(path.parent)
        )
        if not folder:
            return
        output_dir = Path(folder)

        self.busy = True
        self.merge_button.configure(state="disabled")
        self.split_button.configure(state="disabled")
        self.style_button.configure(state="disabled")
        self.format_button.configure(state="disabled")
        self.summary_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.progress.start(12)
        self.status_var.set("正在拆分文档…")
        self._append_log("=" * 60)
        self._append_log(f"按章节拆分：{path.name} -> {output_dir}")

        worker = threading.Thread(
            target=self._split_worker, args=(path, output_dir), daemon=True
        )
        worker.start()

    def _split_worker(self, path: Path, output_dir: Path):
        try:
            outputs = split_document_by_chapters(path, output_dir, self._log)
            self.messages.put(("split_done", (output_dir, outputs)))
        except Exception as exc:
            self.messages.put(("error", (exc, traceback.format_exc())))

    def _on_split_done(self, output_dir: Path, outputs: list[Path]):
        self._finish()
        self.status_var.set(f"拆分完成（{len(outputs)} 个文档）")
        self._append_log(
            f"拆分完成，共生成 {len(outputs)} 个文档：{output_dir}"
        )
        self._open_dir = output_dir
        self.open_button.configure(state="normal")
        if messagebox.askyesno(
            "拆分完成",
            f"已拆分为 {len(outputs)} 个文档：\n{output_dir}\n\n是否打开输出文件夹？",
        ):
            try:
                os.startfile(str(output_dir))
            except Exception as exc:
                messagebox.showerror("无法打开文件夹", str(exc))

    # ---------------- 分析并套用模板样式 ----------------
    def apply_template_to_file(self):
        """分析文档中实际用到的样式及段落数，并按模板逐个替换格式。"""
        if self.busy:
            return
        template = self._template_path()
        if template is None:
            return

        filename = filedialog.askopenfilename(
            title="选择要分析并套用样式的文档（单章或已合并的文档）",
            filetypes=[
                ("Word 文档", "*.docx *.doc *.docm *.rtf"),
                ("所有文件", "*.*"),
            ],
        )
        if not filename:
            return
        path = Path(filename)
        try:
            path = path.resolve()
        except Exception:
            pass
        if not path.is_file() or path.name.startswith("~$"):
            messagebox.showwarning("文件无效", f"无法读取该文件：\n{path}")
            return

        self.status_var.set("正在分析文档样式…")
        self.root.update_idletasks()
        temp_dir = Path(tempfile.mkdtemp(prefix="wm_preview_"))
        try:
            docx_path = _convert_to_docx([path], temp_dir, self._log)
            docx_path = docx_path[0]
            from docx import Document

            document = Document(str(docx_path))
            spans = _document_spans(document)
            chapter_index: int | None = None
            if len(spans) > 1:
                choice = self._ask_chapter_scope(path, spans)
                if choice is _CANCELLED:
                    return
                chapter_index = choice  # None 表示整个文档
            analysis = analyze_document_styles(path, template, chapter_index, self._log)
            rows = analysis["styles"]
            tables = analysis["tables"]
            catalog = _template_style_catalog(template)
        except Exception as exc:
            messagebox.showerror("分析失败", str(exc))
            return
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

        if not rows:
            messagebox.showinfo("提示", "文档中没有可分析的正文或标题段落。")
            return
        if not catalog:
            messagebox.showwarning(
                "模板中没有可用样式", "所选模板中没有段落样式，无法套用格式。"
            )
            return

        selection = self._ask_style_selection(path, rows, catalog, tables)
        if selection is None:
            return
        style_selection, options = selection

        default_name = f"{path.stem}_套用模板.docx"
        save_name = filedialog.asksaveasfilename(
            title="保存套用样式后的文档",
            defaultextension=".docx",
            initialdir=str(path.parent),
            initialfile=default_name,
            filetypes=[("Word 文档", "*.docx")],
        )
        if not save_name:
            return
        output = Path(save_name).with_suffix(".docx")
        try:
            if output.resolve() == path.resolve() or any(
                output.resolve() == item.path.resolve() for item in self.items
            ):
                messagebox.showwarning(
                    "路径冲突", "输出文件不能与源文档相同，请另选一个文件名。"
                )
                return
        except Exception:
            pass
        if output.exists():
            if not messagebox.askyesno(
                "文件已存在", f"输出文件已存在，要覆盖吗？\n{output}"
            ):
                return

        self.busy = True
        self.merge_button.configure(state="disabled")
        self.style_button.configure(state="disabled")
        self.format_button.configure(state="disabled")
        self.summary_button.configure(state="disabled")
        self.split_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.progress.start(12)
        self.status_var.set("正在套用模板样式…")
        self._append_log("=" * 60)
        self._append_log(
            f"按模板 {template.name} 套用格式：{path.name} -> {output}"
        )

        worker = threading.Thread(
            target=self._apply_format_worker,
            args=(path, template, style_selection, options, chapter_index, output),
            daemon=True,
        )
        worker.start()

    def _ask_chapter_scope(self, path: Path, spans: list) -> object:
        """多章节文档：选择处理整个文档还是其中某一章。"""
        dialog = tk.Toplevel(self.root)
        dialog.title("选择章节范围")
        dialog.transient(self.root)
        dialog.resizable(False, False)

        frame = ttk.Frame(dialog, padding=14)
        frame.pack(fill="both", expand=True)
        ttk.Label(
            frame,
            text=f"「{path.name}」包含 {len(spans)} 个章节，请选择处理范围：",
            font=("Microsoft YaHei UI", 10, "bold"),
        ).pack(anchor="w")
        ttk.Label(frame, text="（不选具体章节则处理整个文档）").pack(
            anchor="w", pady=(2, 8)
        )

        listbox = tk.Listbox(frame, height=min(10, max(4, len(spans))), width=60)
        listbox.pack(fill="both", expand=True)
        listbox.insert("end", "整个文档")
        for index, span in enumerate(spans, start=1):
            listbox.insert("end", f"{index}. {span.heading}")

        buttons = ttk.Frame(frame)
        buttons.pack(fill="x", pady=(12, 0))
        result: dict[str, object] = {"value": _CANCELLED}

        def confirm():
            selection = listbox.curselection()
            if not selection or selection[0] == 0:
                result["value"] = None
            else:
                result["value"] = selection[0] - 1
            dialog.destroy()

        ttk.Button(buttons, text="取消", command=dialog.destroy).pack(side="right")
        ttk.Button(buttons, text="确定", command=confirm).pack(
            side="right", padx=(0, 6)
        )
        listbox.selection_set(0)
        dialog.grab_set()
        self.root.wait_window(dialog)
        return result["value"]

    def _ask_style_selection(
        self,
        path: Path,
        rows: list[dict],
        catalog: list[dict],
        tables: list[dict] | None = None,
    ):
        """展示文档中实际用到的样式及段落数，返回（{文档样式: 模板样式}，选项）。"""
        tables = tables or []
        dialog = tk.Toplevel(self.root)
        dialog.title("分析文档样式并按模板套用")
        dialog.transient(self.root)

        frame = ttk.Frame(dialog, padding=12)
        frame.pack(fill="both", expand=True)
        ttk.Label(
            frame,
            text=f"{path.name}：下面列出文档中实际用到的每个样式及段落数。",
            font=("Microsoft YaHei UI", 10, "bold"),
        ).pack(anchor="w", pady=(0, 4))
        ttk.Label(
            frame,
            text=(
                "勾选要套用的样式，并在右侧选择模板中替换它的样式；"
                "已按名称/标题级别自动匹配建议项。"
            ),
            foreground="#555",
        ).pack(anchor="w", pady=(0, 4))
        if tables:
            table_text = "；".join(
                f"表格 {item['index']}：{item['style']}（{item['rows']} 行 × {item['columns']} 列）"
                for item in tables[:4]
            )
            if len(tables) > 4:
                table_text += f" 等共 {len(tables)} 个"
            else:
                table_text = f"共 {len(tables)} 个表格 —— " + table_text
            ttk.Label(
                frame,
                text=f"识别到{table_text}。表格样式与设置保持不变，只套用所选文字样式。",
                foreground="#1f6f3f",
                wraplength=980,
                justify="left",
            ).pack(anchor="w", pady=(0, 8))
        else:
            ttk.Label(frame, text="未识别到表格。", foreground="#555").pack(
                anchor="w", pady=(0, 8)
            )

        container = ttk.Frame(frame)
        container.pack(fill="both", expand=True)
        canvas = tk.Canvas(
            container,
            borderwidth=0,
            highlightthickness=0,
            width=1040,
            height=min(540, 60 + len(rows) * 26),
        )
        scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        table = ttk.Frame(canvas)
        window = canvas.create_window((0, 0), window=table, anchor="nw")
        table.bind(
            "<Configure>",
            lambda event: canvas.configure(scrollregion=canvas.bbox("all")),
        )
        canvas.bind(
            "<Configure>",
            lambda event: canvas.itemconfigure(window, width=event.width),
        )
        canvas.bind_all(
            "<MouseWheel>",
            lambda event: canvas.yview_scroll(int(-event.delta / 120), "units"),
        )
        dialog.bind("<Destroy>", lambda event: canvas.unbind_all("<MouseWheel>"))
        headers = (
            "套用",
            "当前样式",
            "段落数",
            "位置",
            "级别",
            "当前字体",
            "当前字号",
            "多级列表",
            "套用模板样式（可改选）",
        )
        for column, text in enumerate(headers):
            ttk.Label(
                table, text=text, font=("Microsoft YaHei UI", 9, "bold")
            ).grid(row=0, column=column, sticky="w", padx=6, pady=(0, 4))

        labels = [_template_style_label(entry) for entry in catalog]
        if len(set(labels)) != len(labels):
            labels = [
                f"{label} [{entry['style_id']}]"
                for label, entry in zip(labels, catalog)
            ]
        style_ids = [entry["style_id"] for entry in catalog]
        values = ["（保留原样式）"] + labels

        variables: dict[str, tk.BooleanVar] = {}
        combos: dict[str, ttk.Combobox] = {}
        keys: list[str] = []
        for row_index, row in enumerate(rows, start=1):
            key = row.get("style_id") or f"name:{row['style']}"
            keys.append(key)
            suggestion = row.get("template_style_id")
            var = tk.BooleanVar(value=suggestion is not None)
            variables[key] = var
            ttk.Checkbutton(table, variable=var).grid(
                row=row_index, column=0, sticky="w", padx=6
            )

            if row.get("numbered"):
                numbering = "有"
            elif row.get("template_numbered"):
                numbering = (
                    "手动编号"
                    if row.get("manual_numbered")
                    else "无（模板可补）"
                )
            else:
                numbering = "无"

            for column, value in enumerate(
                (
                    row["style"],
                    str(row["count"]),
                    row.get("position") or "正文",
                    row["label"],
                    str(row["font"]),
                    row["size"],
                    numbering,
                ),
                start=1,
            ):
                ttk.Label(table, text=value).grid(
                    row=row_index, column=column, sticky="w", padx=6, pady=1
                )

            combo = ttk.Combobox(
                table,
                values=values,
                state="readonly",
                width=38,
            )
            if suggestion in style_ids:
                combo.current(style_ids.index(suggestion) + 1)
            else:
                combo.current(0)
            combo.grid(row=row_index, column=8, sticky="w", padx=6, pady=1)
            combos[key] = combo

        options_frame = ttk.Frame(frame)
        options_frame.pack(fill="x", pady=(10, 0))
        clear_chars = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            options_frame,
            text="清除文字的直接字符格式（字体、字号、加粗、颜色等），让模板样式完全生效",
            variable=clear_chars,
        ).pack(anchor="w")
        clear_paragraph = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            options_frame,
            text="同时清除段落的直接段落格式（缩进、行距、对齐），让模板的段落格式生效",
            variable=clear_paragraph,
        ).pack(anchor="w")
        add_numbering = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            options_frame,
            text="标题缺少多级列表时，按模板补上多级列表编号（文字已带手动编号的标题会自动跳过）",
            variable=add_numbering,
        ).pack(anchor="w")
        keep_table = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            options_frame,
            text="表格只更新文字格式，保持表格样式、边框、底纹、单元格对齐等设置不变",
            variable=keep_table,
        ).pack(anchor="w")

        ttk.Label(
            frame,
            text=(
                "说明：处理整个文档时修改的是每个样式本身的字符与段落定义；"
                "处理单章时只给该章文字设置格式。"
                "已有的大纲级别、多级列表编号和表格设置保持不变。"
            ),
            foreground="#555",
            wraplength=900,
            justify="left",
        ).pack(anchor="w", pady=(6, 0))

        buttons = ttk.Frame(frame)
        buttons.pack(fill="x", pady=(12, 0))
        result: dict[str, object] = {"value": None}

        def confirm():
            chosen: dict[str, str] = {}
            for key in keys:
                if not variables[key].get():
                    continue
                index = combos[key].current()
                if index <= 0:  # 0 表示保留原样式
                    continue
                chosen[key] = style_ids[index - 1]
            if not chosen:
                messagebox.showwarning(
                    "提示",
                    "请至少勾选一个样式，并在右侧选择要替换成的模板样式。",
                    parent=dialog,
                )
                return
            result["value"] = (
                chosen,
                {
                    "clear_direct_chars": clear_chars.get(),
                    "clear_direct_paragraph": clear_paragraph.get(),
                    "add_numbering": add_numbering.get(),
                    "keep_table_settings": keep_table.get(),
                },
            )
            dialog.destroy()

        ttk.Button(buttons, text="取消", command=dialog.destroy).pack(side="right")
        ttk.Button(buttons, text="开始套用", command=confirm).pack(
            side="right", padx=(0, 6)
        )
        dialog.grab_set()
        self.root.wait_window(dialog)
        return result["value"]

    def _apply_format_worker(
        self,
        path: Path,
        template: Path,
        style_selection: dict,
        options: dict,
        chapter_index: int | None,
        output: Path,
    ):
        try:
            result = apply_template_formats(
                path,
                template,
                style_selection,
                options,
                chapter_index,
                output,
                self._log,
            )
            self.messages.put(("format_done", (output, result)))
        except Exception as exc:
            self.messages.put(("error", (exc, traceback.format_exc())))

    def _on_format_done(self, output: Path, result: dict):
        self._finish()
        paragraphs = result.get("paragraphs", 0)
        styles = result.get("styles", 0)
        numbered = result.get("numbered", 0)
        mismatched = result.get("mismatched") or []
        tables = result.get("tables", 0)
        summary = f"{styles} 种样式、{paragraphs} 段"
        if numbered:
            summary += f"，补全 {numbered} 个多级列表"
        if tables:
            summary += f"，{tables} 个表格（样式未改动）"
        self.status_var.set(f"已套用样式（{paragraphs} 段）")
        self._append_log(f"套用完成，共处理 {summary}：{output}")
        self._open_dir = output.parent
        self.open_button.configure(state="normal")
        note = ""
        if mismatched:
            note = (
                "\n\n校验提示：以下样式与所选模板样式仍有差异（详见日志）："
                + "、".join(mismatched)
            )
        if messagebox.askyesno(
            "套用完成",
            f"已处理 {summary}。{note}\n\n已生成：\n{output}\n\n是否立即打开？",
        ):
            try:
                os.startfile(str(output))
            except Exception as exc:
                messagebox.showerror("无法打开文档", str(exc))

    # ---------------- 生成样式参考汇总模板 ----------------
    def build_summary_template(self):
        """以参考文档（样式模板）为样板，生成带样式说明与示例的汇总模板。"""
        if self.busy:
            return

        reference: Path | None = None
        text = self.style_source.get().strip().strip('"')
        if text:
            candidate = Path(text)
            if candidate.is_file():
                reference = candidate
            else:
                messagebox.showwarning(
                    "样式模板无效", f"找不到样式模板文件：\n{candidate}"
                )
                return
        if reference is None:
            filename = filedialog.askopenfilename(
                title="选择参考文档（用于生成样式汇总模板）",
                filetypes=[
                    ("Word 文档 / 模板", "*.docx *.dotx *.docm *.dotm *.doc"),
                    ("所有文件", "*.*"),
                ],
            )
            if not filename:
                return
            reference = Path(filename)
            if not reference.is_file():
                messagebox.showwarning(
                    "参考文档无效", f"无法读取该文件：\n{reference}"
                )
                return

        save_name = filedialog.asksaveasfilename(
            title="保存样式汇总模板",
            defaultextension=".docx",
            initialdir=str(reference.parent),
            initialfile=f"{reference.stem}_样式汇总模板.docx",
            filetypes=[("Word 文档", "*.docx")],
        )
        if not save_name:
            return
        output = Path(save_name).with_suffix(".docx")
        try:
            if output.resolve() == reference.resolve():
                messagebox.showwarning(
                    "路径冲突", "输出文件不能与参考文档相同，请另选一个文件名。"
                )
                return
        except Exception:
            pass
        if output.exists():
            if not messagebox.askyesno(
                "文件已存在", f"输出文件已存在，要覆盖吗？\n{output}"
            ):
                return

        self.busy = True
        self.merge_button.configure(state="disabled")
        self.style_button.configure(state="disabled")
        self.format_button.configure(state="disabled")
        self.summary_button.configure(state="disabled")
        self.split_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.progress.start(12)
        self.status_var.set("正在生成汇总模板…")
        self._append_log("=" * 60)
        self._append_log(f"按参考文档 {reference.name} 生成样式汇总模板 -> {output}")

        worker = threading.Thread(
            target=self._summary_worker,
            args=(reference, output),
            daemon=True,
        )
        worker.start()

    def _summary_worker(self, reference: Path, output: Path):
        try:
            result = build_style_summary(reference, output, self._log)
            self.messages.put(("summary_done", result))
        except Exception as exc:
            self.messages.put(("error", (exc, traceback.format_exc())))

    def _on_summary_done(self, output: Path):
        self._finish()
        self.status_var.set("汇总模板已生成")
        self._append_log(f"汇总模板已生成：{output}")
        self._open_dir = output.parent
        self.open_button.configure(state="normal")
        if messagebox.askyesno(
            "生成完成", f"汇总模板已生成：\n{output}\n\n是否立即打开？"
        ):
            try:
                os.startfile(str(output))
            except Exception as exc:
                messagebox.showerror("无法打开文档", str(exc))

    def _add_directory(self, folder: Path):
        found = sorted(
            (
                p
                for p in folder.iterdir()
                if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
            ),
            key=lambda p: p.name.casefold(),
        )
        if not found:
            messagebox.showinfo("提示", f"文件夹中没有 Word 文档：\n{folder}")
            return
        for path in found:
            self._add_path(path)

    def _add_path(self, path: Path):
        try:
            path = path.resolve()
        except Exception:
            pass
        if path.name.startswith("~$"):
            return
        if not path.is_file():
            return
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            self._log(f"已跳过不支持的文件：{path.name}")
            return
        if any(item.path == path for item in self.items):
            return

        output_text = self.output_var.get().strip()
        if output_text:
            try:
                if Path(output_text).resolve() == path:
                    return
            except Exception:
                pass

        self.status_var.set(f"正在读取：{path.name}")
        self.root.update_idletasks()
        item = detect_document(path)
        self.items.append(item)
        if self.auto_sort.get():
            self._sort_items()
        self.refresh_list()
        if not self.output_var.get().strip() or self.output_auto:
            self._refresh_output_default()
        self.status_var.set(f"已添加 {len(self.items)} 个文件")

    def refresh_list(self):
        self.listbox.delete(0, "end")
        for item in self.items:
            self.listbox.insert("end", item.label)

    def _sort_items(self):
        """只按各文档的章号排序；章号相同的文档保持原有顺序。"""
        self.items.sort(key=lambda item: item.sort_key)

    def _apply_sort(self):
        if self.auto_sort.get():
            self._sort_items()
            self.refresh_list()

    def move(self, direction: int):
        selection = list(self.listbox.curselection())
        if not selection:
            return
        if self.auto_sort.get():
            self.auto_sort.set(False)
        if direction < 0 and selection[0] == 0:
            return
        if direction > 0 and selection[-1] == len(self.items) - 1:
            return

        indexes = selection if direction < 0 else reversed(selection)
        new_selection = []
        for index in indexes:
            target = index + direction
            self.items[index], self.items[target] = self.items[target], self.items[index]
            new_selection.append(target)
        self.refresh_list()
        for index in new_selection:
            self.listbox.selection_set(index)

    def remove_selected(self):
        selection = list(self.listbox.curselection())
        for index in reversed(selection):
            del self.items[index]
        self.refresh_list()
        if self.output_auto or not self.items:
            self._refresh_output_default()

    def clear_all(self):
        self.items.clear()
        self.refresh_list()
        self._refresh_output_default()

    # ---------------- 输出路径 ----------------
    def _refresh_output_default(self):
        if self.items:
            default = self.items[0].path.parent / "合并结果.docx"
        else:
            desktop = Path.home() / "Desktop"
            default = (desktop if desktop.exists() else Path.home()) / "合并结果.docx"
        self._setting_output = True
        self.output_var.set(str(default))
        self._setting_output = False
        self.output_auto = True

    def _on_output_edited(self, *_):
        if not self._setting_output:
            self.output_auto = False
            self._open_dir = None

    def choose_output(self):
        initial = self.output_var.get().strip()
        initial_dir = str(Path(initial).parent) if initial else str(Path.home())
        filename = filedialog.asksaveasfilename(
            title="保存合并后的文档",
            defaultextension=".docx",
            initialdir=initial_dir,
            initialfile=Path(initial).name if initial else "合并结果.docx",
            filetypes=[("Word 文档", "*.docx")],
        )
        if filename:
            self.output_var.set(filename)

    def _sync_style_options(self):
        """未选择模板时，逐文件统一格式无从谈起，禁用该选项。"""
        has_template = bool(self.style_source.get().strip())
        self.unify_check.configure(state="normal" if has_template else "disabled")

    def choose_style_source(self):
        filename = filedialog.askopenfilename(
            title="选择样式模板文件",
            filetypes=[
                ("Word 模板", "*.dotm *.dotx"),
                ("Word 文档", "*.docx *.docm"),
                ("旧版 Word 文档", "*.doc"),
                ("所有文件", "*.*"),
            ],
        )
        if filename:
            self.style_source.set(filename)

    def open_output_folder(self):
        if self._open_dir is not None and self._open_dir.exists():
            folder = self._open_dir
        else:
            text = self.output_var.get().strip()
            if not text:
                return
            path = Path(text)
            folder = path.parent if path.parent.exists() else Path.home()
        try:
            os.startfile(str(folder))
        except Exception as exc:
            messagebox.showerror("无法打开文件夹", str(exc))

    # ---------------- 单独处理样式 ----------------
    def _template_path(self) -> Path | None:
        """读取并校验界面中的样式模板路径，无效时提示并返回 None。"""
        text = self.style_source.get().strip().strip('"')
        if not text:
            messagebox.showwarning("未选择模板", "请先点击「选择模板…」选择样式模板文件。")
            return None
        template = Path(text)
        if not template.is_file():
            messagebox.showwarning("样式模板无效", f"找不到样式模板文件：\n{template}")
            return None
        if template.suffix.lower() not in STYLE_SOURCE_EXTENSIONS:
            messagebox.showwarning(
                "样式模板无效",
                "样式模板应为 Word 模板或文档：.docx / .dotx / .docm / .dotm / .doc",
            )
            return None
        return template

    def start_style_only(self):
        if self.busy:
            return
        selection = list(self.listbox.curselection())
        items = [self.items[i] for i in selection] if selection else list(self.items)
        if not items:
            messagebox.showwarning("提示", "请先添加要处理的 Word 文档。")
            return

        import_styles = self.standalone_import.get()
        prune_styles = self.standalone_prune.get()
        if not import_styles and not prune_styles:
            messagebox.showwarning(
                "提示", "请至少勾选“导入样式模板”或“删除多余样式”中的一项。"
            )
            return

        template: Path | None = None
        if import_styles:
            template = self._template_path()
            if template is None:
                return

        names = "\n".join(f"  {item.path.name}" for item in items[:8])
        if len(items) > 8:
            names += f"\n  …等 {len(items)} 个文件"
        actions = "、".join(
            filter(
                None,
                [
                    "导入模板样式" if import_styles else "",
                    "删除多余样式" if prune_styles else "",
                ],
            )
        )
        save_source = self.save_source.get()
        if save_source:
            note = "源文件会被直接修改（处理前自动备份）。"
        else:
            note = (
                f"源文件保持原样，处理结果会输出到「{items[0].path.parent.name}"
                "」目录下新建的“样式处理结果_时间”文件夹。"
            )
        if not messagebox.askyesno(
            "确认处理样式",
            f"将对以下 {len(items)} 个文件执行：{actions}\n\n{names}\n\n{note}是否继续？",
        ):
            return

        files = [item.path for item in items]
        self.busy = True
        self.merge_button.configure(state="disabled")
        self.style_button.configure(state="disabled")
        self.format_button.configure(state="disabled")
        self.summary_button.configure(state="disabled")
        self.split_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.progress.start(12)
        self.status_var.set("正在处理样式…")
        self._append_log("=" * 60)

        worker = threading.Thread(
            target=self._style_worker,
            args=(files, template, import_styles, prune_styles, save_source),
            daemon=True,
        )
        worker.start()

    def _style_worker(
        self,
        files: list[Path],
        template: Path | None,
        import_styles: bool,
        prune_styles: bool,
        save_source: bool,
    ):
        try:
            folder = apply_styles_to_files(
                files,
                template,
                import_styles,
                prune_styles,
                self._log,
                save_source=save_source,
            )
            self.messages.put(("style_done", (folder, save_source)))
        except Exception as exc:
            self.messages.put(("error", (exc, traceback.format_exc())))

    def _on_style_done(self, folder: Path, save_source: bool = True):
        self._finish()
        self.status_var.set("样式处理完成")
        self._open_dir = folder
        if save_source:
            self._append_log(f"样式处理完成，原文件备份：{folder}")
            message = f"已完成样式处理。\n\n原文件备份在：\n{folder}\n\n是否打开备份文件夹？"
        else:
            self._append_log(f"样式处理完成，结果输出：{folder}")
            message = f"已完成样式处理（源文件未改动）。\n\n结果输出在：\n{folder}\n\n是否打开输出文件夹？"
        if messagebox.askyesno("样式处理完成", message):
            try:
                os.startfile(str(folder))
            except Exception as exc:
                messagebox.showerror("无法打开文件夹", str(exc))

    # ---------------- 删除标记区间 ----------------
    def start_delete_marked(self):
        if self.busy:
            return
        selection = list(self.listbox.curselection())
        items = [self.items[i] for i in selection] if selection else list(self.items)
        if not items:
            messagebox.showwarning("提示", "请先添加要处理的 Word 文档。")
            return

        start_marker = self.marker_start.get().strip()
        end_marker = self.marker_end.get().strip()
        if not start_marker or not end_marker:
            messagebox.showwarning("提示", "请填写起始标记和结束标记。")
            return
        if start_marker == end_marker:
            messagebox.showwarning("提示", "起始标记和结束标记不能相同。")
            return
        if any(char in start_marker + end_marker for char in "\r\n"):
            messagebox.showwarning("提示", "标记不能包含换行。")
            return

        names = "\n".join(f"  {item.path.name}" for item in items[:8])
        if len(items) > 8:
            names += f"\n  …等 {len(items)} 个文件"
        save_source = self.save_source.get()
        if save_source:
            note = "源文件会被直接修改（处理前自动备份）。"
        else:
            note = (
                f"源文件保持原样，结果输出到「{items[0].path.parent.name}"
                "」目录下新建的“标记删除结果_时间”文件夹。"
            )
        if not messagebox.askyesno(
            "确认删除标记区间",
            f"将删除下列 {len(items)} 个文件中\n"
            f"{start_marker} 到 {end_marker} 之间的全部内容（含两个标记）：\n\n"
            f"{names}\n\n{note}是否继续？",
        ):
            return

        files = [item.path for item in items]
        self.busy = True
        for button in (
            self.merge_button,
            self.style_button,
            self.format_button,
            self.summary_button,
            self.split_button,
            self.mark_button,
            self.open_button,
        ):
            button.configure(state="disabled")
        self.progress.start(12)
        self.status_var.set("正在删除标记区间…")
        self._append_log("=" * 60)
        self._append_log(
            f"删除标记区间：{start_marker} … {end_marker}，共 {len(files)} 个文档"
        )

        worker = threading.Thread(
            target=self._marked_worker,
            args=(files, start_marker, end_marker, save_source),
            daemon=True,
        )
        worker.start()

    def _marked_worker(
        self,
        files: list[Path],
        start_marker: str,
        end_marker: str,
        save_source: bool,
    ):
        try:
            folder, failures = delete_marked_sections_in_files(
                files,
                start_marker,
                end_marker,
                self._log,
                save_source=save_source,
            )
            self.messages.put(("marked_done", (folder, save_source, failures)))
        except Exception as exc:
            self.messages.put(("error", (exc, traceback.format_exc())))

    def _on_marked_done(
        self,
        folder: Path,
        save_source: bool = True,
        failures: list[tuple[str, str]] | None = None,
    ):
        self._finish()
        self._open_dir = folder
        self.open_button.configure(state="normal")
        failures = failures or []
        if save_source:
            self.status_var.set("标记区间已删除")
            self._append_log(f"标记区间已删除，原文件备份：{folder}")
            message = f"已删除标记区间。\n\n原文件备份在：\n{folder}\n\n是否打开备份文件夹？"
        else:
            self.status_var.set("标记区间已删除")
            self._append_log(f"标记区间已删除，结果输出：{folder}")
            message = (
                f"已删除标记区间（源文件未改动）。\n\n结果输出在：\n{folder}"
                "\n\n是否打开输出文件夹？"
            )
        if failures:
            detail = "\n".join(f"  {name}：{reason}" for name, reason in failures[:5])
            if len(failures) > 5:
                detail += f"\n  …等 {len(failures)} 个文件"
            message = (
                f"有 {len(failures)} 个文件未能处理：\n{detail}\n\n"
                + message
            )
        if messagebox.askyesno("标记区间已删除", message):
            try:
                os.startfile(str(folder))
            except Exception as exc:
                messagebox.showerror("无法打开文件夹", str(exc))

    # ---------------- 合并流程 ----------------
    def start_merge(self):
        if self.busy:
            return
        if not self.items:
            messagebox.showwarning("提示", "请先添加要合并的 Word 文档。")
            return

        output_text = self.output_var.get().strip().strip('"')
        if not output_text:
            messagebox.showwarning("提示", "请填写输出文件路径。")
            return
        output = Path(output_text)
        if output.suffix.lower() != ".docx":
            output = output.with_suffix(".docx")
            self.output_var.set(str(output))

        try:
            output.parent.mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            messagebox.showerror("输出路径无效", str(exc))
            return

        for item in self.items:
            try:
                if item.path.resolve() == output.resolve():
                    messagebox.showerror("路径冲突", "输出文件不能与输入文件相同。")
                    return
            except Exception:
                pass

        if output.exists():
            if not messagebox.askyesno("文件已存在", f"输出文件已存在，要覆盖吗？\n{output}"):
                return

        style_source: Path | None = None
        if self.style_source.get().strip():
            style_source = self._template_path()
            if style_source is None:
                return

        # 按章号排序后再交给合并流程
        ordered = list(self.items)
        if self.auto_sort.get():
            ordered.sort(key=lambda item: item.sort_key)
        files = [item.path for item in ordered]
        options = MergeOptions(
            make_toc=self.make_toc.get(),
            toc_title=self.toc_title.get().strip() or "目录",
            page_break=self.page_break.get(),
            style_source=style_source,
            prune_styles=self.prune_styles.get(),
            unify_styles=self.unify_styles.get() and style_source is not None,
            save_source=self.save_source.get(),
        )

        self.busy = True
        self.merge_button.configure(state="disabled")
        self.split_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.progress.start(12)
        self.status_var.set("正在合并…")
        self._append_log("=" * 60)
        self._append_log(f"开始合并 {len(files)} 个文档 -> {output}")

        worker = threading.Thread(
            target=self._merge_worker, args=(files, output, options), daemon=True
        )
        worker.start()

    def _merge_worker(self, files: list[Path], output: Path, options: MergeOptions):
        try:
            engine = merge_documents(files, output, options, self._log)
            self.messages.put(("done", (output, engine)))
        except Exception as exc:
            self.messages.put(("error", (exc, traceback.format_exc())))

    def _on_merge_done(self, output: Path, engine: str):
        self._finish()
        self.status_var.set(f"合并完成（{engine}）")
        self._append_log(f"合并完成：{output}")
        self._open_dir = output.parent
        self.open_button.configure(state="normal")
        if messagebox.askyesno(
            "合并完成", f"文档已生成：\n{output}\n\n是否立即打开？"
        ):
            try:
                os.startfile(str(output))
            except Exception as exc:
                messagebox.showerror("无法打开文档", str(exc))

    def _on_merge_error(self, exc: Exception, detail: str):
        self._finish()
        self.status_var.set("合并失败")
        self._append_log("合并失败：\n" + detail)
        messagebox.showerror("合并失败", str(exc))

    def _finish(self):
        self.busy = False
        self.merge_button.configure(state="normal")
        self.style_button.configure(state="normal")
        self.format_button.configure(state="normal")
        self.summary_button.configure(state="normal")
        self.split_button.configure(state="normal")
        self.mark_button.configure(state="normal")
        self.progress.stop()

    def _on_close(self):
        if self.busy:
            messagebox.showinfo("正在合并", "文档正在合并，请等待完成后再关闭窗口。")
            return
        self.root.destroy()

    # ---------------- 消息与日志 ----------------
    def _log(self, text: str):
        self.messages.put(("log", text))

    def _append_log(self, text: str):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", text + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _poll_messages(self):
        try:
            while True:
                kind, payload = self.messages.get_nowait()
                if kind == "log":
                    self._append_log(str(payload))
                elif kind == "done":
                    self._on_merge_done(*payload)
                elif kind == "style_done":
                    self._on_style_done(payload)
                elif kind == "marked_done":
                    self._on_marked_done(*payload)
                elif kind == "into_done":
                    self._on_into_done(*payload)
                elif kind == "format_done":
                    self._on_format_done(*payload)
                elif kind == "summary_done":
                    self._on_summary_done(payload)
                elif kind == "split_done":
                    self._on_split_done(*payload)
                elif kind == "error":
                    self._on_merge_error(*payload)
        except queue.Empty:
            pass
        self.root.after(120, self._poll_messages)


def _selftest() -> int:
    """打包后排障：检查依赖是否齐全，结果写入 selftest.log。

    用法：WordMergeTool.exe --selftest [日志路径]
    """
    import sys
    import tempfile
    from datetime import datetime

    lines = [f"Word 文档合并工具 自检 {datetime.now():%Y-%m-%d %H:%M:%S}"]
    ok = True

    checks = (
        ("tkinter", "tkinter"),
        ("python-docx", "docx"),
        ("lxml", "lxml.etree"),
        ("docxcompose", "docxcompose"),
        ("pywin32", "win32com.client"),
        ("pythoncom", "pythoncom"),
        ("tkinterdnd2", "tkinterdnd2"),
    )
    for label, module in checks:
        try:
            __import__(module)
            lines.append(f"[OK]   {label}")
        except Exception as exc:
            ok = False
            lines.append(f"[FAIL] {label}: {exc}")

    # 实际生成并读取一个 docx，确认打包后的 docx/lxml 组合可用
    try:
        from docx import Document

        temp_dir = Path(tempfile.mkdtemp(prefix="wm_selftest_"))
        sample = temp_dir / "selftest.docx"
        document = Document()
        document.add_heading("自检标题", level=1)
        document.add_paragraph("自检正文。")
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "A"
        document.save(sample)
        reopened = Document(str(sample))
        assert reopened.paragraphs[0].text == "自检标题"
        assert len(reopened.tables) == 1
        styles = len(_template_style_catalog(sample))
        lines.append(f"[OK]   docx 读写与样式解析（{styles} 个段落样式）")
        shutil.rmtree(temp_dir, ignore_errors=True)
    except Exception as exc:
        ok = False
        lines.append(f"[FAIL] docx 读写: {exc}")

    lines.append("结果：" + ("全部通过" if ok else "存在失败项"))
    text = "\n".join(lines)

    target = None
    for arg in sys.argv[1:]:
        if not arg.startswith("-"):
            target = Path(arg)
            break
    if target is None:
        base = Path(sys.executable if getattr(sys, "frozen", False) else __file__)
        target = base.with_name("selftest.log")
    try:
        target.write_text(text, encoding="utf-8")
    except Exception:
        pass
    print(text)
    return 0 if ok else 1


def main():
    import sys

    if "--selftest" in sys.argv:
        sys.exit(_selftest())
    root = _BaseTk()
    MergeApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
